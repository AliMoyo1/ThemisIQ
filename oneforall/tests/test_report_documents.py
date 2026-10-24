"""
The documents of the Reports page: the PDF, the workbook and the print view of a frozen run.

Every expectation is written by hand from the seed data. Three things are pinned: the figures (a chart series and a table row
are the figures that were frozen, not a recomputation), the shape every report shares (a cover with the logo and the tier,
and a footer on every page naming the module, who generated it and when), and who may open a document (the same people who may
open the run).
"""
import asyncio
import io
import json
import zipfile

import pytest
from pypdf import PdfReader

import modules.launcher.routes_reports as reports
from core.report_kit import theme
from core.report_kit.render_pdf import render_pdf
from modules.launcher import report_documents as documents
from modules.launcher import report_snapshot as snap
from tests.test_platform_ownership import _insert, _request, signed_in, world  # noqa: F401
from tests.test_report_runs import _quick


def _raw(coro):
    response = asyncio.run(coro)
    return response.status_code, response.body, response.headers


def _document(user, run_id, fmt="pdf"):
    return _raw(reports.api_report_run_document(_request(user, {"format": fmt}), run_id))


def _print(user, run_id):
    return _raw(reports.report_run_print(_request(user), run_id))


def _pages(pdf_bytes):
    return [page.extract_text() for page in PdfReader(io.BytesIO(pdf_bytes)).pages]


def _envelope(report_type, result, **run):
    """A frozen envelope as the Reports page stores it, with the figures written by hand."""
    return {"format": 1,
            "run": {"definition_id": 1, "definition_name": run.pop("name", "Test report"), "report_type": report_type,
                    "type_label": run.pop("type_label", "Test report"), "definition_version": "2026-10-22.1",
                    "generated_at": "2026-10-21 16:05:00", "generated_by": {"id": 9, "name": "Tatenda Chikomba", "username": "tchikomba"},
                    **run},
            "scope": {"org_name": "Econet", "label": "Whole organization", "fingerprint": "abc"},
            "period": {"basis": "as_of", "note": "Every figure is read at the moment the report is generated."},
            "result": result, "sources": [], "inputs": {}, "limitations": ["A limit stated for the reader."],
            "exclusions": ["Records of other organizations."], "complete": True, "decision_status": snap.DECISION_STATUS,
            "narrative": None, "changes": None}


def _chart(report, title):
    return next(c for c in report.charts() if c.title == title)


def _table(report, title):
    return next(t for t in report.tables() if t.title == title)


def _kpis(report):
    return {k.label: k for block in report.blocks if block.__class__.__name__ == "Kpis" for k in block.items}


COMPLIANCE = {
    "total_controls": 16, "compliant": 6, "non_compliant": 4, "partial": 3,
    "by_status": {"Implemented": 6, "In Progress": 3, "Not Started": 4, "Draft": 2, "Needs Update": 1},
    "by_framework": [{"name": "Alpha", "total": 8, "ok": 1}, {"name": "Beta", "total": 8, "ok": 5}, {"name": "Empty", "total": 0, "ok": None}],
}


# ── What kind of report each one is ─────────────────────────────────────────

def test_every_report_type_has_a_tier_and_only_the_board_styles_are_board():
    assert set(reports.REPORT_TIERS) == set(reports.REPORT_TYPES)
    board = {key for key, tier in reports.REPORT_TIERS.items() if tier == "board"}
    assert board == {"executive_brief", "executive_summary", "preset_board"}
    assert {key for key, tier in reports.REPORT_TIERS.items() if tier == "operational"} == set(reports.REPORT_TYPES) - board
    page_groups = {title: keys for title, keys in reports.REPORT_GROUPS}
    assert page_groups["Board reports"] == ["executive_brief", "executive_summary", "preset_board"]
    assert all(reports.REPORT_TIERS[k] == "operational" for t, keys in page_groups.items() if t != "Board reports" for k in keys)


def test_a_new_run_records_its_tier_and_the_username_of_who_generated_it(world):
    run = _quick(world.users["super"], "executive_brief")[1]["snapshot"]["run"]
    assert run["tier"] == "board" and run["generated_by"] == {"id": world.users["super"]["id"], "name": "super", "username": "super"}


# ── The figures: a chart and a table are the frozen figures ─────────────────

def test_the_compliance_charts_and_table_are_the_frozen_figures_rounded_half_up():
    report = documents.build_document(_envelope("compliance_summary", COMPLIANCE))
    kpis = _kpis(report)
    assert kpis["Controls in scope"].value == "16" and kpis["Implemented"].value == "6"
    assert kpis["Implemented"].note == "38% of controls", "6 of 16 is 37.5, which rounds to 38 as on the pages"
    assert kpis["In another state"].value == "3", "16 less 6, 4 and 3: draft and needs update"
    assert _chart(report, "Controls by status").data == [
        ("Implemented", 6), ("Not Started", 4), ("In Progress", 3), ("Draft", 2), ("Needs Update", 1)], \
        "implemented first, then the largest groups, so that what is folded into other states is the smallest"
    assert _chart(report, "Implemented share by framework").data == [("Beta", 63), ("Alpha", 13)], \
        "5 of 8 is 62.5 and 1 of 8 is 12.5: both round up, and a framework with no controls is not charted"
    stack = _chart(report, "Controls by framework: implemented and still to do")
    assert stack.data == [("Alpha", [("Implemented", 1), ("Not yet implemented", 7)]), ("Beta", [("Implemented", 5), ("Not yet implemented", 3)])]
    table = _table(report, "Frameworks with controls in scope")
    assert [(r["framework"], r["controls"], r["implemented"], r["share"]) for r in table.rows] == [
        ("Alpha", 8, 1, 13), ("Beta", 8, 5, 63)]
    assert table.note == "1 other active framework has no controls in scope and is not listed."
    assert report.meta.tier == "operational" and report.meta.module == "aria"


def test_a_missing_figure_is_n_a_and_is_never_drawn_as_a_zero():
    report = documents.build_document(_envelope("compliance_summary", {
        "total_controls": None, "compliant": None, "non_compliant": None, "partial": None, "by_status": None, "by_framework": None}))
    assert {k.value for k in _kpis(report).values()} == {"n/a"}
    donut = _chart(report, "Controls by status")
    assert donut.data == [] and donut.is_empty() and "no controls" in donut.empty_text.lower()


def test_the_risk_chart_reconciles_to_the_open_total_with_a_not_rated_bar():
    result = {"total_open": 7, "by_level": {"critical": 2, "high": 2, "medium": 1, "low": 1}, "unrated": 1,
              "by_module": {"erm": 4, "bcm": 2, "unassigned": 1},
              "top_risks": [{"title": "Loss of the data centre", "risk_level": "critical", "source_module": "erm",
                             "treatment": "mitigate", "status": "open"}]}
    report = documents.build_document(_envelope("risk_report", result))
    levels = _chart(report, "Open risks by level")
    assert levels.data == [("Critical", 2), ("High", 2), ("Medium", 1), ("Low", 1), ("Not rated", 1)]
    assert sum(v for _, v in levels.data) == 7
    assert "4 of 7 open risks (57%) are critical or high" in levels.takeaway
    assert _chart(report, "Open risks by where they are recorded").data == [
        ("Enterprise Risk", 4), ("Resilience", 2), ("No source module", 1)]
    assert _table(report, "Highest scoring open risks").rows[0]["module"] == "Enterprise Risk"


def test_a_board_report_is_short_leads_with_measures_and_names_what_needs_a_decision():
    result = {"controls_total": 8, "controls_compliant": 5, "compliance_pct": 63, "risks_open": 7, "risks_critical": 2,
              "audits_active": 1, "breaches_open": 1, "dsrs_open": 0, "incidents_active": 0, "bcm_plans": 3, "tasks_open": 10,
              "tasks_overdue": 4}
    report = documents.build_document(_envelope("executive_summary", result))
    assert report.meta.tier == "board" and report.meta.tier_label == "Board report"
    kpis = _kpis(report)
    assert kpis["Controls implemented"].value == "63%" and kpis["Controls implemented"].note == "5 of 8 controls"
    assert kpis["Critical risks"].tone == "bad" and kpis["Open data subject requests"].tone == "neutral", "a zero is not flagged"
    callouts = [b for b in report.blocks if b.__class__.__name__ == "Callout"]
    decisions = next(c for c in callouts if c.kind == "decision")
    assert decisions.items == [
        "2 critical risks: agree an owner and a treatment date for each.",
        "1 open breach: confirm the notification position and who leads.",
        "4 overdue tasks: ask the owners to clear or re-plan them."]
    assert not report.tables() or all(t.title != "Sources read" for t in report.tables()), "a board report carries no source table"
    attention = _chart(report, "Items that need attention")
    assert attention.data == [("Critical risks", 2), ("Open breaches", 1), ("Open data subject requests", 0),
                              ("Active incidents", 0), ("Overdue tasks", 4)]


def test_the_board_preset_states_a_decision_once_however_many_of_its_parts_raise_it():
    """The brief, the risk part and the privacy part each raised "critical risks" or "open breach", so the board read it twice."""
    result = {"sections": {
        "executive_brief": {"controls_total": 20, "controls_compliant": 10, "risks_critical": 2, "risks_high": 1, "audits_active": 1,
                            "breaches_open": 1, "sla_breaches": 0},
        "risk_report": {"total_open": 5, "by_level": {"critical": 2, "high": 1, "medium": 2}},
        "privacy_overview": {"ropa_count": 3, "dpia_count": 1, "breaches_open": 1, "dsr_open": 2}}}
    report = documents.build_document(_envelope("preset_board", result))
    decisions = next(b for b in report.blocks if b.__class__.__name__ == "Callout" and b.title == "For the board to decide or note").items
    assert len(decisions) == len(set(decisions)), decisions
    assert sum("open breach" in d for d in decisions) == 1 and sum("critical risk" in d for d in decisions) == 1
    assert any("data subject request" in d for d in decisions), "a decision only one part raises is kept"
    messages = next(b for b in report.blocks if b.__class__.__name__ == "Callout" and b.title == "Key messages").items
    assert len(messages) == len(set(messages))


def test_with_nothing_to_decide_the_board_report_says_so():
    result = {"controls_total": 8, "controls_compliant": 8, "risks_critical": 0, "risks_high": 0, "audits_active": 0,
              "breaches_open": 0, "sla_breaches": 0}
    report = documents.build_document(_envelope("executive_brief", result))
    decision = next(b for b in report.blocks if b.__class__.__name__ == "Callout" and b.kind == "decision")
    assert decision.items == ["No item in these figures needs a board decision. Read the movement and the limits below."]


def test_the_trend_chart_needs_two_saved_runs_and_otherwise_says_so():
    report = documents.build_document(_envelope("executive_brief", {"controls_total": 8, "controls_compliant": 5}))
    trend = next(c for c in report.charts() if c.kind == "line")
    assert trend.data == [] and "at least two saved runs" in trend.empty_text


# ── The narrative: grounded or left out ─────────────────────────────────────

def test_a_draft_narrative_with_a_figure_that_is_not_in_the_data_is_left_out_of_the_document():
    env = _envelope("executive_summary", {"controls_total": 8, "controls_compliant": 5, "compliance_pct": 63, "risks_critical": 2})
    env["narrative"] = {"state": "draft", "text": "Compliance stands at 63 percent with 2 critical risks and 47 open audits.", "model": "m"}
    report = documents.build_document(env)
    assert not [b for b in report.blocks if b.__class__.__name__ == "Paragraphs"]
    left_out = next(b for b in report.blocks if b.__class__.__name__ == "Callout" and b.title == "Draft narrative left out")
    assert "47" in left_out.items[0] and "not in the data" in left_out.items[0]


def test_a_draft_narrative_whose_figures_are_all_in_the_data_is_printed_as_a_draft():
    env = _envelope("executive_summary", {"controls_total": 8, "controls_compliant": 5, "compliance_pct": 63, "risks_critical": 2})
    env["narrative"] = {"state": "draft", "text": "5 of 8 controls (63%) are implemented. 2 risks are critical. As at 21 October 2026.", "model": "m"}
    report = documents.build_document(env)
    paragraphs = next(b for b in report.blocks if b.__class__.__name__ == "Paragraphs")
    assert paragraphs.paragraphs == ["5 of 8 controls (63%) are implemented. 2 risks are critical. As at 21 October 2026."]
    assert "not reviewed" in paragraphs.label and "found in the data" in paragraphs.label


# ── The same real figures from a database ───────────────────────────────────

@pytest.fixture
def seeded(world):
    db = world.db
    fa = _insert(db, "frameworks", name="Alpha")
    fb = _insert(db, "frameworks", name="=HYPERLINK(\"http://evil.example\",\"x\")")
    for i in range(8):
        _insert(db, "controls", framework_id=fa, ref=f"A.{i}", name=f"Control A{i}", status="Implemented" if i < 5 else "Not Started")
    _insert(db, "controls", framework_id=fb, ref="B.1", name="Control B1", status="Implemented")
    # Enterprise risks scored 25 (critical), 12 (high), 6 (medium), 4 (low) and a closed one; the shared fixture adds one open
    # enterprise risk at the default 3 by 3 (9, medium). Platform: critical, high, no level, and a closed one.
    for title, likelihood, impact, status in (("E crit", 5, 5, "open"), ("E high", 4, 3, "open"), ("E med", 3, 2, "open"),
                                                ("E low", 2, 2, "open"), ("E closed", 5, 5, "closed")):
        _insert(db, "erm_enterprise_risks", title=title, likelihood=likelihood, impact=impact, status=status)
    for title, level, status, module in (("P crit", "critical", "open", "bcm"), ("P high", "high", "open", "bcm"),
                                          ("P none", None, "open", None), ("P closed", "low", "closed", "bcm")):
        _insert(db, "risk_register", title=title, risk_level=level, status=status, source_module=module)
    return world


def test_the_executive_figures_count_both_risk_registers(seeded):
    _, run = _quick(seeded.users["super"], "executive_brief")
    figures = run["result"]
    assert figures["risks_critical"] == 2, "one enterprise risk and one platform risk (the old figure counted only the platform one)"
    assert figures["risks_high"] == 2
    _, summary = _quick(seeded.users["super"], "executive_summary")
    assert summary["result"]["risks_open"] == 8 and summary["result"]["risks_critical"] == 2, "five enterprise and three platform"
    assert summary["result"]["compliance_pct"] == 67, "6 of 9 controls is 66.7 percent"


def test_the_compliance_percentage_rounds_half_up(world):
    db = world.db
    fid = _insert(db, "frameworks", name="Half")
    for i in range(8):
        _insert(db, "controls", framework_id=fid, ref=f"H.{i}", name=f"c{i}", status="Implemented" if i < 5 else "Draft")
    result = _quick(world.users["super"], "executive_summary")[1]["result"]
    assert (result["controls_total"], result["controls_compliant"], result["compliance_pct"]) == (8, 5, 63), \
        "5 of 8 is 62.5 percent: a person and the page script say 63, Python round() says 62"


def test_the_risk_report_keeps_the_unrated_risk_in_the_total(seeded):
    result = _quick(seeded.users["super"], "risk_report")[1]["result"]
    assert result["total_open"] == 8 and result["by_level"] == {"critical": 2, "high": 2, "medium": 2, "low": 1}
    assert result["unrated"] == 1 and result["by_module"] == {"bcm": 2, "unassigned": 1, "erm": 5}
    assert sum(result["by_level"].values()) + result["unrated"] == result["total_open"]


def test_an_open_audit_finding_is_one_that_neither_its_status_nor_its_action_plan_calls_closed(world):
    audit = world.rec[("audit", "A")]
    for title, status, cap in (("Closed by plan", "open", "Closed"), ("In progress", "in_progress", "RCA"),
                               ("Being verified", "in_progress", "Verification"), ("Closed", "closed", "Closed")):
        _insert(world.db, "grid_non_conformances", audit_id=audit, title=title, status=status, cap_status=cap)
    result = _quick(world.users["super"], "audit_status")[1]["result"]
    assert result["open_ncs"] == 3, "the fixture's own open finding and the two in progress; the old status = 'open' test counted the one the plan had closed"


def test_the_audit_report_says_how_many_audits_are_in_every_stage(world):
    db = world.db
    for status in ("Planning", "Active", "Review", "Complete"):
        _insert(db, "grid_audits", name=f"Audit {status}", status=status)
    result = _quick(world.users["super"], "audit_status")[1]["result"]
    assert result["by_status"] == {"Planning": 2, "Active": 1, "Review": 1, "Complete": 1}, "the shared fixture holds one more in planning"
    assert result["total_audits"] == 5 and result["active"] == 3 and result["completed"] == 1


# ── The document files ──────────────────────────────────────────────────────

def _run(seeded, report_type):
    return _quick(seeded.users["super"], report_type)[1]["run_id"]


@pytest.mark.parametrize("report_type, tier, chip", [("compliance_summary", "operational", "OPERATIONAL REPORT"),
                                                     ("executive_brief", "board", "BOARD REPORT")])
def test_the_pdf_has_the_logo_and_tier_on_the_cover_and_the_footer_on_every_page(seeded, report_type, tier, chip):
    run_id = _run(seeded, report_type)
    status, body, headers = _document(seeded.users["super"], run_id)
    assert status == 200 and body[:5] == b"%PDF-" and headers["content-type"] == "application/pdf"
    assert headers["cache-control"] == "no-store" and headers["x-content-type-options"] == "nosniff"
    disposition = headers["content-disposition"]
    assert disposition.startswith('attachment; filename="ThemisIQ_') and disposition.endswith('.pdf"')
    assert all(c.isalnum() or c in '._-"; =' for c in disposition), "an ASCII file name with no path or markup"
    reader = PdfReader(io.BytesIO(body))
    pages = _pages(body)
    assert len(reader.pages[0].images) >= 1, "the ThemisIQ logo is on the cover"
    tier_label = "Board report" if tier == "board" else "Operational report"
    assert tier_label in pages[0] and "Whole organization" in pages[0] and "super" in pages[0]
    for number, text in enumerate(pages, start=1):
        assert "Module:" in text and "Generated by super" in text and "UTC" in text and "CAT" in text, f"page {number} footer"
        assert f"Page {number} of {len(pages)}" in text
        assert tier_label in text
    assert reader.metadata.title and reader.metadata.author == "super"
    event = seeded.db.execute("SELECT details FROM audit_log WHERE action='report_run_export' ORDER BY id DESC").fetchone()
    assert f"format=pdf; tier={tier}" in event["details"]


def test_the_module_in_the_footer_is_the_navigation_name_of_the_module_the_report_reads(seeded):
    expected = {"compliance_summary": "Governance", "audit_status": "Audit", "privacy_overview": "Privacy",
                "bcm_readiness": "Resilience", "risk_report": "Command Centre", "executive_brief": "Command Centre"}
    for report_type, name in expected.items():
        text = _pages(_document(seeded.users["super"], _run(seeded, report_type))[1])[1]
        assert f"Module: {name}" in text, report_type


def test_the_pdf_shows_the_cover_scope_and_who_exported_it(seeded):
    text = _pages(_document(seeded.users["super"], _run(seeded, "compliance_summary"))[1])[0]
    assert "ORGANISATION" in text and "org" in text and "GENERATED BY" in text and "EXPORTED BY" in text and "REFERENCE" in text


def test_the_workbook_has_a_cover_sheet_print_settings_and_no_formula(seeded):
    import openpyxl
    run_id = _run(seeded, "compliance_summary")
    status, body, headers = _document(seeded.users["super"], run_id, "xlsx")
    assert status == 200 and headers["content-type"].endswith("spreadsheetml.sheet")
    wb = openpyxl.load_workbook(io.BytesIO(body))
    assert wb.sheetnames[0] == "Cover" and "Summary" in wb.sheetnames and "Frameworks" in wb.sheetnames
    cover = wb["Cover"]
    assert cover["B9"].value == "OPERATIONAL REPORT" and len(cover._images) == 1
    for ws in wb.worksheets:
        assert ws.page_setup.fitToWidth == 1 and ws.sheet_properties.pageSetUpPr.fitToPage
        assert "Module: Governance" in ws.oddFooter.left.text and "Generated by super" in ws.oddFooter.left.text
        assert "UTC" in ws.oddFooter.center.text and "Page &P of &N" in ws.oddFooter.right.text
    frameworks = wb["Frameworks"]
    names = [frameworks.cell(row=r, column=1) for r in range(2, frameworks.max_row + 1)]
    formula = next(c for c in names if str(c.value).startswith("=HYPERLINK"))
    assert formula.data_type == "s", "a framework named like a formula is stored as text"
    with zipfile.ZipFile(io.BytesIO(body)) as z:
        assert not [n for n in z.namelist() if n.startswith("xl/worksheets/") and b"<f>" in z.read(n)], "no sheet holds a formula"
    assert frameworks.print_title_rows == "$1:$1" and frameworks.freeze_panes == "A2"


def test_the_print_view_carries_the_cover_the_footer_and_print_css_and_escapes_what_it_shows(seeded):
    run_id = _run(seeded, "executive_brief")
    status, body, headers = _print(seeded.users["super"], run_id)
    html = body.decode("utf-8")
    assert status == 200 and headers["content-type"].startswith("text/html") and headers["cache-control"] == "private, no-store"
    assert "BOARD REPORT" in html and "@page" in html and "counter(pages)" in html and "print-color-adjust" in html
    assert "Module: Command Centre" in html and "Generated by super" in html
    assert "/static/img/themisiq-logo-clear.png" in html
    assert "<script>alert" not in html


def test_the_print_view_escapes_markup_in_a_name_and_in_the_footer():
    from core.report_kit.render_html import render_html
    hostile = "</style><script>alert(2)</script>\"; } body{display:none"
    env = _envelope("compliance_summary", COMPLIANCE, name="<script>alert(1)</script> Board \"pack\"",
                    generated_by={"id": 9, "name": hostile, "username": "x\\y"})
    html = render_html(documents.build_document(env))
    assert "<script>alert(1)</script>" not in html and "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert html.count("<script>") == 1 and html.count("</style>") == 1, "only the page's own script and style element exist"
    assert "alert(2)" not in html.split("</style>", 1)[1].split("<script>", 1)[0], "the footer name did not reach the page as markup"
    style = html.split("<style>", 1)[1].split("</style>", 1)[0]
    assert "<script" not in style and "body{display:none" not in style, "the footer name stayed inside its CSS string"


# ── Who may open a document ─────────────────────────────────────────────────

def test_only_the_person_the_run_was_generated_for_may_open_its_documents(seeded):
    run_id = _run(seeded, "compliance_summary")
    colleague, other = seeded.users["owner_a"], seeded.users["dpo_b"]
    assert _document(seeded.users["super"], run_id)[0] == 200
    for user in (colleague, other):
        assert _document(user, run_id)[0] == 404 and _document(user, run_id, "xlsx")[0] == 404 and _print(user, run_id)[0] == 404
    assert _document(seeded.users["super"], run_id, "docx")[0] == 400, "only pdf and xlsx are offered"
    denied = seeded.db.execute("SELECT COUNT(*) FROM audit_log WHERE action='report_run_denied'").fetchone()[0]
    assert denied >= 6, "each refusal is written to the audit log"


def test_a_run_saved_before_reports_were_frozen_has_no_document(seeded):
    db = seeded.db
    did = _insert(db, "report_definitions", name="Old", report_type="compliance_summary", created_by=seeded.users["super"]["id"])
    rid = _insert(db, "report_runs", definition_id=did, status="completed", triggered_by="user:1:x", result_json=json.dumps({"x": 1}))
    status, body, _ = _document(seeded.users["super"], rid)
    assert status == 409 and b"saved before reports were frozen" in body


def test_a_tampered_run_is_refused_not_rendered(seeded):
    run_id = _run(seeded, "compliance_summary")
    seeded.db.execute("UPDATE report_runs SET result_json = REPLACE(result_json, 'Alpha', 'Forged') WHERE id = %s", (run_id,))
    seeded.db.commit()
    status, body, _ = _document(seeded.users["super"], run_id)
    assert status == 409 and b"integrity" in body


# ── The kit itself: a board cover is not an operational cover ───────────────

def test_the_two_tiers_have_different_covers():
    board = documents.build_document(_envelope("executive_brief", {"controls_total": 1, "controls_compliant": 1}))
    ops = documents.build_document(_envelope("compliance_summary", COMPLIANCE))
    assert board.meta.tier_chip == "BOARD REPORT" and ops.meta.tier_chip == "OPERATIONAL REPORT"
    assert render_pdf(board)[:5] == b"%PDF-" and render_pdf(ops)[:5] == b"%PDF-"
    assert theme.TIERS["board"]["label"] != theme.TIERS["operational"]["label"]
