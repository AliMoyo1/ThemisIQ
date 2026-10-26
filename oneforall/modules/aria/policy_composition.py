"""Compose an imported DOCX without rewriting its retained OOXML relationships.

Every output is a candidate in the existing named policy approval workflow.
The imported original and its external attestation are never modified.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
import base64
import binascii
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from lxml import etree

from core.timeutils import utcnow
from modules.aria import policy_intake as intake
from modules.aria import policy_storage as storage
from modules.aria import policy_workflow_service as workflow
from modules.aria.policy_access import policy_authoring_enabled_for

MAX_BLOCKS = 3000
MAX_REPLACEMENT = 100_000
MAX_DIAGRAMS = 4
MAX_NODES = 12
MAX_EDGES = 20
MAX_LOGO_BYTES = 2 * 1024 * 1024
MODES = {"replace_cover", "replace_body", "replace_both"}


def _cover_logo(encoded: str) -> bytes | None:
    if not encoded:
        return None
    if not isinstance(encoded, str) or len(encoded) > (MAX_LOGO_BYTES * 4 // 3) + 8:
        raise intake.IntakeError(422, "Cover logo must be a PNG or JPEG under 2 MB.")
    try:
        from PIL import Image, ImageOps, UnidentifiedImageError
        raw = base64.b64decode(encoded, validate=True)
        if len(raw) > MAX_LOGO_BYTES:
            raise ValueError()
        with Image.open(io.BytesIO(raw)) as image:
            if image.format not in ("PNG", "JPEG") or image.width * image.height > 2_000_000:
                raise ValueError()
            normalized = ImageOps.exif_transpose(image).convert("RGB")
            result = io.BytesIO()
            normalized.save(result, format="PNG")
            return result.getvalue()
    except (ValueError, binascii.Error, OSError, UnidentifiedImageError, Image.DecompressionBombError) as exc:
        raise intake.IntakeError(422, "Cover logo must be a valid PNG or JPEG under 2 MB and 2 megapixels.") from exc


def _blocks(document):
    blocks = list(document.iter_inner_content())
    if len(blocks) > MAX_BLOCKS:
        raise intake.IntakeError(422, "Document has too many blocks to compose safely.")
    return blocks


def outline(db, actor: dict, import_id: str) -> list[dict]:
    row = intake.get_import(db, actor, import_id)
    if row["state"] != "registered":
        raise intake.IntakeError(409, "Register the original before composing a revision.")
    if not row["file_name"].lower().endswith(".docx"):
        raise intake.IntakeError(422, "This action needs a DOCX. Keep the PDF original or supply a reviewed editable source.")
    source = intake._source_path(row)
    if not source.is_file() or storage.sha256_file(source) != row["file_hash"]:
        raise intake.IntakeError(409, "The original is missing or changed.")
    storage.validate_docx(source)
    doc = Document(source)
    result = []
    for i, block in enumerate(_blocks(doc)):
        if hasattr(block, "rows"):
            raw = " | ".join(" | ".join(cell.text for cell in row.cells) for row in block.rows)
        else:
            raw = block.text or ""
        result.append({"index": i, "kind": type(block).__name__.lower(),
                       "preview": re.sub(r"\s+", " ", raw)[:180]})
    return result


def _clean_text(value, name, maximum):
    if not isinstance(value, str) or len(value) > maximum:
        raise intake.IntakeError(422, f"Invalid {name}.")
    return value.strip()


def _diagram_image(spec: dict):
    """A deterministic diagram; input is data, never executable markup."""
    from PIL import Image, ImageDraw, ImageFont

    if not isinstance(spec, dict):
        raise intake.IntakeError(422, "Invalid diagram.")
    kind = spec.get("kind")
    if kind not in ("process_flow", "system_block", "data_flow"):
        raise intake.IntakeError(422, "Choose a supported diagram type.")
    caption = _clean_text(spec.get("caption", ""), "diagram caption", 180)
    alt = _clean_text(spec.get("alt_text", ""), "diagram description", 300)
    if not caption or not alt:
        raise intake.IntakeError(422, "Diagrams need a caption and description.")
    nodes, edges = spec.get("nodes"), spec.get("edges")
    if not isinstance(nodes, list) or not 1 <= len(nodes) <= MAX_NODES:
        raise intake.IntakeError(422, "Diagram needs 1 to 12 nodes.")
    if not isinstance(edges, list) or len(edges) > MAX_EDGES:
        raise intake.IntakeError(422, "Diagram has too many edges.")
    ids = set()
    clean_nodes = []
    for node in nodes:
        if not isinstance(node, dict):
            raise intake.IntakeError(422, "Invalid diagram node.")
        ident = _clean_text(node.get("id"), "node ID", 30)
        label = _clean_text(node.get("label"), "node label", 80)
        if not ident or not label or ident in ids:
            raise intake.IntakeError(422, "Diagram node IDs and labels must be unique and nonempty.")
        ids.add(ident)
        clean_nodes.append({"id": ident, "label": label})
    clean_edges = []
    for edge in edges:
        if not isinstance(edge, dict) or edge.get("from") not in ids or edge.get("to") not in ids:
            raise intake.IntakeError(422, "Diagram edge refers to an unknown node.")
        clean_edges.append({"from": edge["from"], "to": edge["to"]})
    width, height = 1240, max(210, 120 + len(clean_nodes) * 155)
    image = Image.new("RGB", (width, height), "#f4f6f4")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 24)
    except OSError:
        font = ImageFont.load_default()
    locations = {}
    for i, node in enumerate(clean_nodes):
        x, y = 110, 72 + i * 155
        locations[node["id"]] = (x, y)
    for edge in clean_edges:
        x1, y1 = locations[edge["from"]]
        x2, y2 = locations[edge["to"]]
        draw.line((x1 + 900, y1 + 55, x2 + 20, y2 + 55), fill="#4e6470", width=5)
        draw.polygon([(x2 + 20, y2 + 55), (x2 + 2, y2 + 45), (x2 + 2, y2 + 65)], fill="#4e6470")
    for node in clean_nodes:
        x, y = locations[node["id"]]
        draw.rounded_rectangle((x, y, x + 920, y + 108), radius=18, fill="#e7ece9", outline="#9baeb2", width=2)
        draw.text((x + 22, y + 34), node["label"], fill="#263633", font=font)
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue(), {"kind": kind, "caption": caption, "alt_text": alt,
                               "nodes": clean_nodes, "edges": clean_edges}


def _render(source: Path, target: Path, selection: dict, metadata: dict) -> dict:
    mode = selection.get("mode")
    if mode not in MODES:
        raise intake.IntakeError(422, "Choose a composition action.")
    doc = Document(source)  # editing this source package preserves media relationships in kept blocks
    blocks = _blocks(doc)
    if not blocks:
        raise intake.IntakeError(422, "The source has no document body.")
    boundary = selection.get("boundary")
    if not isinstance(boundary, int) or isinstance(boundary, bool) or not 0 < boundary < len(blocks):
        raise intake.IntakeError(422, "Select the exact block where the body begins.")
    exclude = selection.get("exclude_blocks", [])
    if not isinstance(exclude, list) or len(exclude) > 100 or any(
            not isinstance(i, int) or isinstance(i, bool) or i < 0 or i >= len(blocks) for i in exclude):
        raise intake.IntakeError(422, "Invalid excluded blocks.")
    excluded = set(exclude)
    replacing_body = mode in ("replace_body", "replace_both")
    cover_text = _clean_text(selection.get("cover_text", ""), "cover text", 4000)
    cover_logo = _cover_logo(selection.get("cover_logo_base64", ""))
    if mode == "replace_body" and (cover_text or cover_logo):
        raise intake.IntakeError(422, "A retained cover cannot also be replaced.")
    body_text = _clean_text(selection.get("replacement_body", ""), "replacement body", MAX_REPLACEMENT)
    if replacing_body and not body_text:
        raise intake.IntakeError(422, "Enter reviewed replacement body text.")
    diagram_specs = selection.get("diagrams", [])
    if not isinstance(diagram_specs, list) or len(diagram_specs) > MAX_DIAGRAMS:
        raise intake.IntakeError(422, "Add at most four diagrams.")
    if diagram_specs and not replacing_body:
        raise intake.IntakeError(422, "Add diagrams to a replacement body candidate.")
    figure = selection.get("figure") or None
    if figure is not None:
        if not replacing_body or not isinstance(figure, dict):
            raise intake.IntakeError(422, "Add an approved figure only to a replacement body candidate.")
        try:
            figure_bytes = _cover_logo(figure.get("image_base64", ""))
        except intake.IntakeError as exc:
            raise intake.IntakeError(422, "Figure must be a valid PNG or JPEG under 2 MB and 2 megapixels.") from exc
        if not figure_bytes:
            raise intake.IntakeError(422, "Choose a PNG or JPEG figure.")
        figure_caption = _clean_text(figure.get("caption", ""), "figure caption", 180)
        figure_alt = _clean_text(figure.get("alt_text", ""), "figure description", 300)
        if not figure_caption or not figure_alt:
            raise intake.IntakeError(422, "A figure needs a caption and an accessible description.")
    keep = [block for i, block in enumerate(blocks) if i not in excluded and (
        (i >= boundary if mode == "replace_cover" else i < boundary if mode == "replace_body" else False))]
    if mode == "replace_cover" and not keep:
        raise intake.IntakeError(422, "The selected body is empty.")
    retained_xml = b"".join(etree.tostring(block._element) for block in keep)
    retained_hash = hashlib.sha256(retained_xml).hexdigest()
    for block in blocks:
        if block not in keep:
            doc.element.body.remove(block._element)
    if mode in ("replace_cover", "replace_both"):
        # New cover goes before the retained body; no guessing about source headings.
        cover_lines = [metadata["title"], f"Version {metadata['version']}",
                       f"Effective date: {metadata['effective_date'] or 'To be set on approval'}"]
        if cover_text:
            cover_lines.extend(cover_text.splitlines())
        for line in reversed(cover_lines):
            paragraph = doc.add_paragraph(line)
            paragraph.style = "Title" if line == cover_lines[0] else "Subtitle"
            doc.element.body.remove(paragraph._element)
            doc.element.body.insert(0, paragraph._element)
        if cover_logo:
            paragraph = doc.add_paragraph()
            image = paragraph.add_run().add_picture(io.BytesIO(cover_logo), width=Inches(1.8))
            image._inline.docPr.set("descr", "Organization logo supplied for this policy cover")
            doc.element.body.remove(paragraph._element)
            doc.element.body.insert(0, paragraph._element)
    if replacing_body:
        for raw in body_text.splitlines():
            line = raw.strip()
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            else:
                doc.add_paragraph(line)
    figure_manifest = None
    if figure is not None:
        paragraph = doc.add_paragraph()
        image = paragraph.add_run().add_picture(io.BytesIO(figure_bytes), width=Inches(6.2))
        image._inline.docPr.set("descr", figure_alt)
        doc.add_paragraph(figure_caption)
        figure_manifest = {"image_sha256": hashlib.sha256(figure_bytes).hexdigest(),
                           "caption": figure_caption, "alt_text": figure_alt}
    diagrams = []
    for spec in diagram_specs:
        png, clean = _diagram_image(spec)
        paragraph = doc.add_paragraph()
        image = paragraph.add_run().add_picture(io.BytesIO(png), width=Inches(6.2))
        image._inline.docPr.set("descr", clean["alt_text"])
        caption = doc.add_paragraph(clean["caption"])
        for run in caption.runs:
            run.font.size = Pt(9)
        diagrams.append({**clean, "image_sha256": hashlib.sha256(png).hexdigest()})
    if selection.get("clear_headers"):
        for section in doc.sections:
            for part in (section.header, section.footer):
                part.is_linked_to_previous = False
                for block in list(part.paragraphs) + list(part.tables):
                    part._element.remove(block._element)
                part.add_paragraph("")
    # Fail if the retained content was modified as a side effect of composition.
    if retained_hash != hashlib.sha256(b"".join(etree.tostring(block._element) for block in keep)).hexdigest():
        raise intake.IntakeError(500, "Source body changed while composing.")
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    storage.validate_docx(target)
    return {"mode": mode, "boundary": boundary, "excluded_blocks": sorted(excluded),
            "source_block_count": len(blocks), "retained_body_sha256": retained_hash,
            "clear_headers": bool(selection.get("clear_headers")), "diagrams": diagrams,
            "figure": figure_manifest,
            "cover_text_sha256": hashlib.sha256(cover_text.encode()).hexdigest() if cover_text else None,
            "cover_logo_sha256": hashlib.sha256(cover_logo).hexdigest() if cover_logo else None}


def create_candidate(db, actor: dict, import_id: str, selection: dict) -> dict:
    row = intake.get_import(db, actor, import_id)
    if row["state"] != "registered":
        raise intake.IntakeError(409, "Register the original before composing a candidate.")
    if intake.detail(db, actor, import_id)["pending_replacement"]:
        raise intake.IntakeError(409, "Verify the replacement before composing another version.")
    if not row["file_name"].lower().endswith(".docx"):
        raise intake.IntakeError(422, "PDF body editing needs a separately reviewed editable DOCX source.")
    if not isinstance(selection, dict):
        raise intake.IntakeError(422, "Invalid composition selection.")
    if not policy_authoring_enabled_for(actor.get("org_id")):
        raise intake.IntakeError(403, "Policy authoring is not enabled for this organization.")
    doc = db.execute("SELECT * FROM aria_documents WHERE id=%s AND org_id=%s",
                     (row["document_id"], row["org_id"])).fetchone()
    if not doc or not workflow._draft_can_edit_document(actor, dict(doc)):
        raise intake.IntakeError(403, "Policy edit permission required.")
    source = intake._source_path(row)
    if not source.is_file() or storage.sha256_file(source) != row["file_hash"]:
        raise intake.IntakeError(409, "The original changed or is missing.")
    metadata = json.loads(row["reviewed_json"] or "{}")
    if not metadata:
        raise intake.IntakeError(409, "Review the document metadata first.")
    candidate_dir = storage._org_root(row["org_id"]) / "compositions"
    target = candidate_dir / f"{uuid.uuid4().hex}.docx"
    try:
        manifest = _render(source, target, selection, metadata)
        manifest.update({"source_import_id": import_id, "source_sha256": row["file_hash"],
                         "output_sha256": storage.sha256_file(target), "created_by": actor["id"],
                         "created_at": utcnow().isoformat(), "review_required": True})
        draft = workflow.start_revision_draft(db, actor, doc["doc_id"], commit=False)
        draft_metadata = json.loads(draft["metadata_json"] or "{}")
        draft_metadata.update({"title": metadata["title"], "doc_type": metadata["doc_type"],
                               "composition": {"source_import_id": import_id, "mode": manifest["mode"]}})
        serialized = json.dumps(manifest, sort_keys=True)
        db.execute("UPDATE aria_policy_drafts SET content_kind='uploaded_docx', body='',"
                   "uploaded_source_path=%s, uploaded_source_sha256=%s, metadata_json=%s,"
                   "renderer_manifest_json=%s,renderer_manifest_sha256=%s,updated_at=%s "
                   "WHERE id=%s AND org_id=%s AND state='editing'",
                   (storage.relative_path(target), manifest["output_sha256"], json.dumps(draft_metadata),
                    serialized, hashlib.sha256(serialized.encode()).hexdigest(), utcnow().isoformat(),
                    draft["id"], row["org_id"]))
        db.commit()
        return {"draft_id": draft["id"], "lock_version": draft["lock_version"],
                "source_sha256": row["file_hash"], "candidate_sha256": manifest["output_sha256"],
                "manifest": manifest}
    except Exception:
        target.unlink(missing_ok=True)
        raise
