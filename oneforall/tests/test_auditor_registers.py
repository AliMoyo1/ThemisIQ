"""Saved register builders retain every authorized detail row and its source scope."""
import pytest

from modules.launcher import report_builders, report_documents
from modules.launcher.register_reports import RegisterNotAvailable
from modules.launcher.scoped_metrics import table_scope


ADMIN = {"id": 1, "org_id": 1, "roles": ["super_admin"], "is_super_admin": True,
         "username": "admin", "full_name": "Admin"}


def test_framework_register_has_every_control_above_ui_page_size(test_db, monkeypatch):
    fid = test_db.execute("INSERT INTO frameworks (name) VALUES ('Register QA')").lastrowid
    test_db.executemany("INSERT INTO controls (framework_id,ref,name) VALUES (%s,%s,%s)",
                        [(fid, f"A.{n}", f"Control {n}") for n in range(1, 506)])
    test_db.commit()
    built = report_builders.build(test_db, ADMIN, "framework_auditor_register", {"framework_id": fid})
    assert built.result["counts"]["controls"] == 505
    assert len(built.result["sections"]["controls"]) == 505
    assert built.inputs["controls"]["truncated"] is False
    assert len(built.inputs["controls"]["ids"]) == 505
    env = {"run": {"report_type": "framework_auditor_register", "definition_name": "QA register",
                    "generated_by": {"id": 1, "name": "Admin"}},
           "result": built.result, "sources": built.sources, "inputs": built.inputs,
           "scope": {"org_name": "QA", "label": "Whole organization"}}
    report = report_documents.build_document(env)
    control_table = next(t for t in report.tables() if t.title == "Controls")
    assert len(control_table.shown()) == 505 and control_table.max_rows is None
    from core.report_kit.render_pdf import render_pdf
    from core.report_kit.render_xlsx import render_xlsx
    from pypdf import PdfReader
    from openpyxl import load_workbook
    from io import BytesIO
    pdf = PdfReader(BytesIO(render_pdf(report)))
    assert "A.505" in " ".join(page.extract_text() or "" for page in pdf.pages)
    workbook = load_workbook(BytesIO(render_xlsx(report)), read_only=True)
    assert workbook["Controls"].max_row >= 506
    import core.report_kit.render_xlsx as xlsx_module
    monkeypatch.setattr(xlsx_module, "SHEET_ROW_LIMIT", 200)
    split = load_workbook(BytesIO(render_xlsx(report)), read_only=True)
    assert sum(sum(str(row[0] or "").startswith("A.") for row in split[f"Controls {part}"].iter_rows(values_only=True))
               for part in (1, 2, 3)) == 505


def test_ai_registers_use_application_and_assessment_scope(test_db):
    unit_a = test_db.execute("INSERT INTO business_units (name,is_active) VALUES ('A',1)").lastrowid
    unit_b = test_db.execute("INSERT INTO business_units (name,is_active) VALUES ('B',1)").lastrowid
    app_a = test_db.execute("INSERT INTO applications (name,business_unit_id) VALUES ('Visible AI',%s)",
                            (unit_a,)).lastrowid
    app_b = test_db.execute("INSERT INTO applications (name,business_unit_id) VALUES ('Other AI',%s)",
                            (unit_b,)).lastrowid
    for app in (app_a, app_b):
        test_db.execute("INSERT INTO ai_use_profiles (application_id,purpose) VALUES (%s,'test')", (app,))
    ass_a = test_db.execute("INSERT INTO aims_assessments (ref,title,business_unit_id,application_id) "
                            "VALUES ('ASSESS-A','Assessment A',%s,%s)", (unit_a, app_a)).lastrowid
    ass_b = test_db.execute("INSERT INTO aims_assessments (ref,title,business_unit_id,application_id) "
                            "VALUES ('ASSESS-B','Assessment B',%s,%s)", (unit_b, app_b)).lastrowid
    for assessment, title in ((ass_a, "Visible risk"), (ass_b, "Other risk")):
        test_db.execute("INSERT INTO aims_risks (assessment_id,risk_description) VALUES (%s,%s)",
                        (assessment, title))
    test_db.commit()
    viewer = {"id": 14, "org_id": 1, "roles": ["compliance_manager"],
              "business_unit_id": unit_a, "licensed_modules": ["aria", "orm", "sentinel", "grid"]}
    systems = report_builders.build(test_db, viewer, "ai_systems_register")
    assert [r["name"] for r in systems.result["sections"]["systems"]] == ["Visible AI"]
    risks = report_builders.build(test_db, viewer, "ai_risk_register")
    assert [r["risk_description"] for r in risks.result["sections"]["risks"]] == ["Visible risk"]
    assert [r["system"] for r in risks.result["sections"]["systems"]] == ["Visible AI"]
    assert table_scope("aims_risks", viewer)[0] != "(1 = 1)"
    with pytest.raises(report_builders.RecordNotAvailable):
        report_builders.build(test_db, {**viewer, "roles": ["employee"]}, "ai_risk_register")


def test_framework_register_excludes_unadopted_and_foreign_documents(test_db):
    test_db.execute("INSERT INTO organizations (id,name,slug) VALUES (1,'Own','own')")
    fid = test_db.execute("INSERT INTO frameworks (name) VALUES ('Scoped framework')").lastrowid
    test_db.execute("INSERT INTO aria_documents (doc_id,framework,title,version,org_id) "
                    "VALUES ('GLOBAL-1','Scoped framework','Legacy cross-tenant title','1.0',NULL)")
    test_db.execute("INSERT INTO organizations (id,name,slug) VALUES (2,'Other','other')")
    test_db.execute("INSERT INTO aria_documents (doc_id,framework,title,version,org_id) "
                    "VALUES ('OTHER-1','Scoped framework','Foreign document title','1.0',2)")
    test_db.execute("INSERT INTO aria_documents (doc_id,framework,title,version,org_id) "
                    "VALUES ('LOCAL-1','Scoped framework','Owned document title','1.0',1)")
    built = report_builders.build(test_db, ADMIN, "framework_auditor_register", {"framework_id": fid})
    assert [row["title"] for row in built.result["sections"]["documents"]] == ["Owned document title"]
