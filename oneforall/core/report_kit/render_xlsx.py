"""
The workbook of a report: a cover sheet with the logo, a summary sheet with native Excel charts, one sheet per table,
and print settings (A4, fit to width, repeated header row, a header and footer naming the module, who and when).

openpyxl. Every string is stored as a string (`identity.set_cell`), so a value that begins with an equals sign is shown as
text and never evaluated. Missing figures are the text `n/a`, never a zero. The workbook holds every row of a table (the PDF,
DOCX and HTML may cut a long table; the workbook never does).
"""
import io
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from core.report_kit import charts, identity, theme
from core.report_kit.model import Chart, Heading, Kpis, Meta, Paragraphs, Callout, Report, Table, cell_text

FONT = theme.DOCX_FONT


def _fill(hex_colour) -> PatternFill:
    return PatternFill("solid", start_color=hex_colour.lstrip("#"), end_color=hex_colour.lstrip("#"))


def _font(size=10, bold=False, italic=False, colour=theme.INK, underline=False) -> Font:
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=colour.lstrip("#"), underline="single" if underline else None)


def _edge(colour=theme.SILVER_LIGHT) -> Side:
    return Side(style="thin", color=colour.lstrip("#"))


def _put(ws, ref, value, font=None, fill=None, align=None, border=None, number_format=None):
    cell = ws[ref] if isinstance(ref, str) else ws.cell(row=ref[0], column=ref[1])
    identity.set_cell(cell, value)
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    if number_format:
        cell.number_format = number_format
    return cell


def _hf(text) -> str:
    """Header and footer text: an ampersand is a code there, so a real one is doubled."""
    return str(text).replace("&", "&&")


def _print_setup(ws, meta: Meta, landscape: bool, title_rows=None):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.85, bottom=0.95, header=0.3, footer=0.4)
    ws.print_options.horizontalCentered = False
    if title_rows:
        ws.print_title_rows = title_rows
    ws.oddHeader.left.text = _hf(theme.PRODUCT_NAME)
    ws.oddHeader.left.size = 9
    ws.oddHeader.left.font = f"{FONT},Bold"
    ws.oddHeader.right.text = _hf(identity.ellipsize(meta.title, 90))
    ws.oddHeader.right.size = 8
    ws.oddFooter.left.text = _hf(f"{meta.footer_module()}  |  {meta.footer_who()}")
    ws.oddFooter.left.size = 8
    ws.oddFooter.center.text = _hf(meta.footer_when())
    ws.oddFooter.center.size = 8
    ws.oddFooter.right.text = _hf(meta.tier_label) + "  |  Page &P of &N"
    ws.oddFooter.right.size = 8


def _safe_sheet_name(name, used) -> str:
    base = re.sub(r"[\\/*?\[\]:]", " ", identity.clean_text(name, keep_newlines=False)).strip(" '")[:31] or "Sheet"
    candidate, n = base, 2
    while candidate.lower() in used:
        suffix = f" {n}"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


# ── Cover ───────────────────────────────────────────────────────────────────

def _cover(wb, report: Report, sheet_names):
    meta = report.meta
    board = meta.tier == theme.BOARD
    ws = wb.active
    ws.title = "Cover"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = theme.GRAPHITE.lstrip("#")
    for column, width in (("A", 3), ("B", 22), ("C", 52), ("D", 30), ("E", 30)):
        ws.column_dimensions[column].width = width
    image = XLImage(str(theme.LOGO_PATH))
    image.width, image.height = 120, 120
    ws.add_image(image, "B2")
    for row in range(2, 9):
        ws.row_dimensions[row].height = 21
    _put(ws, "B9", meta.tier_chip, _font(9, True, colour=theme.WHITE if board else theme.INK),
         _fill(theme.GRAPHITE if board else theme.SILVER), Alignment(horizontal="center", vertical="center"))
    ws.row_dimensions[9].height = 20
    ws.merge_cells("B11:E11")
    _put(ws, "B11", meta.title, _font(22, True, colour=theme.WHITE if board else theme.GRAPHITE),
         _fill(theme.GRAPHITE if board else theme.SILVER_LIGHT), Alignment(wrap_text=True, vertical="center", indent=1))
    ws.row_dimensions[11].height = 62
    row = 12
    if meta.subtitle:
        ws.merge_cells(f"B{row}:E{row}")
        _put(ws, f"B{row}", meta.subtitle, _font(11, colour=theme.SLATE), None, Alignment(wrap_text=True, vertical="center", indent=1))
        ws.row_dimensions[row].height = 22
        row += 1
    row += 1
    for label, value in meta.cover_rows():
        _put(ws, f"B{row}", label.upper(), _font(8, True, colour=theme.SLATE), None, Alignment(vertical="center"),
             Border(bottom=_edge()))
        ws.merge_cells(f"C{row}:E{row}")
        _put(ws, f"C{row}", value, _font(11), None, Alignment(wrap_text=True, vertical="center"), Border(bottom=_edge()))
        ws.row_dimensions[row].height = 22
        row += 1
    row += 1
    _put(ws, f"B{row}", "CONTENTS", _font(8, True, colour=theme.SLATE))
    row += 1
    for name in sheet_names:
        _put(ws, f"B{row}", name, _font(10, True, colour=theme.GRAPHITE))
        row += 1
    row += 1
    ws.merge_cells(f"B{row}:E{row + 1}")
    _put(ws, f"B{row}", "These are the figures as they stood when the report was generated. Every cell is stored as text or "
                        "as a number: nothing in this workbook is a formula, and a missing figure is shown as n/a, never as zero.",
         _font(9, italic=True, colour=theme.SLATE), None, Alignment(wrap_text=True, vertical="top"))
    row += 3
    _put(ws, f"B{row}", f"{theme.PRODUCT_NAME}  |  {theme.TAGLINE}  |  {meta.confidentiality}", _font(8, colour=theme.SLATE))
    _print_setup(ws, meta, landscape=False)
    return ws


# ── Summary: headings, measures, native charts ──────────────────────────────

def _chart_data(ws, chart: Chart, row: int):
    """Write the chart's own data as a small table at column B and return (header row, last row, series count)."""
    labels_header = {"stack": "Item", "line": chart.x_label or "Reading"}.get(chart.kind, chart.y_label or "Item")
    if chart.kind == "stack":
        names = []
        for _, segments in chart.data:
            for name, _ in segments:
                if name not in names:
                    names.append(name)
        header = [labels_header, *names]
        body = [[label, *[dict(segments).get(name) for name in names]] for label, segments in chart.data]
    else:
        header = [labels_header, chart.unit or "Value"]
        body = [[label, value] for label, value in chart.data]
    for offset, text in enumerate(header):
        _put(ws, (row, 2 + offset), text, _font(9, True), _fill(theme.SILVER), Alignment(horizontal="left" if offset == 0 else "right"),
             Border(bottom=_edge(theme.SILVER_DEEP)))
    for r, values in enumerate(body, start=1):
        for offset, value in enumerate(values):
            _put(ws, (row + r, 2 + offset), "n/a" if value is None else value, _font(9),
                 _fill(theme.SILVER_MIST) if r % 2 == 0 else None,
                 Alignment(horizontal="left" if offset == 0 else "right"), Border(bottom=_edge()),
                 number_format=None if offset == 0 or value is None else "#,##0.##")
    return row, row + len(body), len(header) - 1


def _native_chart(ws, chart: Chart, header_row: int, last_row: int, series: int, anchor: str):
    categories = Reference(ws, min_col=2, min_row=header_row + 1, max_row=last_row)
    if chart.kind == "donut":
        native = DoughnutChart()
        native.holeSize = 55
    elif chart.kind == "line":
        native = LineChart()
    else:
        native = BarChart()
        native.type = "bar"
        if chart.kind == "stack":
            native.grouping = "stacked"
            native.overlap = 100
    data = Reference(ws, min_col=3, max_col=2 + series, min_row=header_row, max_row=last_row)
    native.add_data(data, titles_from_data=True)
    native.set_categories(categories)
    native.title = None
    native.height, native.width = 7.6, 15.5
    palette = [(chart.colours[i % len(chart.colours)] if chart.colours else theme.series_colour(i)) for i in range(max(series, 1))]
    if chart.kind == "donut":
        series_obj = native.series[0]
        for index in range(len(chart.data)):
            point = DataPoint(idx=index)
            colour = (chart.colours[index % len(chart.colours)] if chart.colours else theme.series_colour(index)).lstrip("#")
            point.graphicalProperties.solidFill = colour
            point.graphicalProperties.line.solidFill = theme.WHITE.lstrip("#")
            series_obj.dPt.append(point)
        native.dataLabels = DataLabelList()
        native.dataLabels.showVal = True
        native.dataLabels.showPercent = False
        native.legend.position = "r"
        return native
    for index, series_obj in enumerate(native.series):
        colour = (theme.GRAPHITE if chart.kind == "line" else
                  (chart.colours[0] if chart.colours and chart.kind == "bar" else
                   (theme.SILVER_DEEP if chart.kind == "bar" else palette[index % len(palette)]))).lstrip("#")
        series_obj.graphicalProperties.solidFill = colour
        series_obj.graphicalProperties.line.solidFill = theme.GRAPHITE.lstrip("#")
        if chart.kind == "line":
            series_obj.smooth = False
            series_obj.marker.symbol = "circle"
            series_obj.marker.size = 6
            series_obj.marker.graphicalProperties.solidFill = theme.SILVER.lstrip("#")
            series_obj.marker.graphicalProperties.line.solidFill = theme.GRAPHITE.lstrip("#")
    native.dataLabels = DataLabelList()
    native.dataLabels.showVal = True
    native.x_axis.delete = False
    native.y_axis.delete = False
    if chart.kind == "line":
        native.x_axis.title = chart.x_label or None
        native.y_axis.title = chart.y_label or None
    else:
        native.x_axis.title = chart.y_label or None
        native.y_axis.title = chart.x_label or None
    native.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=theme.SILVER_MIST.lstrip("#")))
    if chart.unit == "%":
        native.y_axis.number_format = '0"%"'
        native.y_axis.scaling.max = chart.max_value or 100
        native.y_axis.scaling.min = 0
    if series == 1:
        native.legend = None
    else:
        native.legend.position = "b"
    return native


def _summary(wb, report: Report):
    meta = report.meta
    ws = wb.create_sheet("Summary")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = theme.SILVER_DEEP.lstrip("#")
    for column, width in (("A", 3), ("B", 34), ("C", 16), ("D", 16), ("E", 16), ("F", 16), ("G", 16)):
        ws.column_dimensions[column].width = width
    row = 2
    wrote_anything = False
    for block in report.blocks:
        if isinstance(block, Heading):
            _put(ws, f"B{row}", block.text, _font(15, True, colour=theme.GRAPHITE), None, None, Border(bottom=_edge(theme.SILVER)))
            for column in "CDEFG":
                ws[f"{column}{row}"].border = Border(bottom=_edge(theme.SILVER))
            row += 1
            if block.intro:
                ws.merge_cells(f"B{row}:G{row}")
                _put(ws, f"B{row}", block.intro, _font(9, colour=theme.SLATE), None, Alignment(wrap_text=True, vertical="top"))
                ws.row_dimensions[row].height = 26
                row += 1
            row += 1
            wrote_anything = True
        elif isinstance(block, Kpis):
            if block.title:
                _put(ws, f"B{row}", block.title, _font(10, True, colour=theme.GRAPHITE))
                row += 1
            for index, label in enumerate(("Measure", "Value", "Note")):
                _put(ws, (row, 2 + index), label, _font(9, True), _fill(theme.SILVER), Alignment(horizontal="left"),
                     Border(bottom=_edge(theme.SILVER_DEEP)))
            row += 1
            for kpi in block.items:
                _put(ws, (row, 2), kpi.label, _font(10), None, None, Border(bottom=_edge()))
                _put(ws, (row, 3), kpi.value, _font(11, True, colour=theme.TONES.get(kpi.tone, theme.GRAPHITE)), None,
                     Alignment(horizontal="right"), Border(bottom=_edge()))
                ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=7)
                _put(ws, (row, 4), kpi.note, _font(9, colour=theme.SLATE), None, None, Border(bottom=_edge()))
                row += 1
            row += 1
            wrote_anything = True
        elif isinstance(block, Chart):
            _put(ws, f"B{row}", block.title, _font(10, True, colour=theme.GRAPHITE))
            row += 1
            ws.merge_cells(f"B{row}:G{row}")
            _put(ws, f"B{row}", charts.takeaway(block), _font(9, italic=True, colour=theme.SLATE), None, Alignment(wrap_text=True))
            row += 1
            if not charts.usable(block):
                ws.merge_cells(f"B{row}:G{row}")
                _put(ws, f"B{row}", block.empty_text if block.is_empty() else
                     "A trend needs at least two readings. Run the report again later to see the movement.",
                     _font(9, italic=True, colour=theme.SLATE), _fill(theme.PEARL), Alignment(wrap_text=True, vertical="center"))
                ws.row_dimensions[row].height = 24
                row += 2
            else:
                header_row, last_row, series = _chart_data(ws, block, row)
                native = _native_chart(ws, block, header_row, last_row, series, f"I{row}")
                ws.add_chart(native, f"I{row}")
                row = max(last_row + 2, row + 17)
            wrote_anything = True
        elif isinstance(block, Callout):
            _put(ws, f"B{row}", block.title, _font(10, True, colour=theme.GRAPHITE))
            row += 1
            for item in block.items:
                ws.merge_cells(f"B{row}:G{row}")
                _put(ws, f"B{row}", "• " + item, _font(10), _fill(theme.PEARL), Alignment(wrap_text=True, vertical="top"))
                ws.row_dimensions[row].height = max(18, 15 * (1 + len(item) // 95))
                row += 1
            row += 1
            wrote_anything = True
        elif isinstance(block, Paragraphs):
            _put(ws, f"B{row}", block.title, _font(10, True, colour=theme.GRAPHITE))
            row += 1
            if block.label:
                _put(ws, f"B{row}", block.label, _font(8, italic=True, colour=theme.SLATE))
                row += 1
            for text in block.paragraphs:
                if identity.clean_text(text):
                    ws.merge_cells(f"B{row}:G{row}")
                    _put(ws, f"B{row}", identity.clean_text(text), _font(10), None, Alignment(wrap_text=True, vertical="top"))
                    ws.row_dimensions[row].height = max(18, 14 * (1 + len(str(text)) // 100))
                    row += 1
            row += 1
            wrote_anything = True
    if not wrote_anything:
        wb.remove(ws)
        return None
    _print_setup(ws, meta, landscape=True)
    return ws


# ── Data sheets ─────────────────────────────────────────────────────────────

def _typed(value, kind):
    """(cell value, number format) for one table value. A missing figure is the text n/a."""
    if value is None or value == "":
        return "n/a", None
    if kind in ("int", "number", "pct"):
        try:
            number = float(value)
        except (TypeError, ValueError):
            return str(value), None
        if kind == "int":
            return int(round(number)), "#,##0"
        if kind == "pct":
            return number, '0"%"' if number == int(number) else '0.0"%"'
        return number, "#,##0.0"
    if kind == "date":
        moment = identity.to_utc(value)
        return (moment.date() if moment and not (moment.hour or moment.minute) else moment) or str(value), "d mmm yyyy"
    return identity.clean_text(value, keep_newlines=True), None


def _data_sheet(wb, block: Table, meta: Meta, name):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = theme.SILVER.lstrip("#")
    ws.sheet_view.showGridLines = False
    columns = block.sheet_columns or block.columns
    for index, col in enumerate(columns, start=1):
        _put(ws, (1, index), col.header, _font(10, True), _fill(theme.SILVER),
             Alignment(horizontal=col.align if col.align in ("left", "right", "center") else "left", vertical="center", wrap_text=True),
             Border(bottom=_edge(theme.SILVER_DEEP)))
    ws.row_dimensions[1].height = 24
    for r, row in enumerate(block.rows, start=2):
        for index, col in enumerate(columns, start=1):
            value, number_format = _typed(row.get(col.key), col.kind)
            url = identity.safe_link(row.get(col.link)) if col.link else ""
            cell = _put(ws, (r, index), value, _font(10, colour=theme.GRAPHITE, underline=True) if url else _font(10),
                        _fill(theme.SILVER_MIST) if r % 2 == 1 else None,
                        Alignment(horizontal=col.align if col.align in ("left", "right", "center") else "left", vertical="top",
                                  wrap_text=col.kind == "text"), Border(bottom=_edge()), number_format)
            if url:
                cell.hyperlink = url
    if not block.rows:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(columns), 1))
        _put(ws, "A2", block.empty_text, _font(10, italic=True, colour=theme.SLATE), _fill(theme.PEARL))
    for index, col in enumerate(columns, start=1):
        sample = [len(cell_text(row.get(col.key), col.kind)) for row in block.rows[:200]]
        width = max([len(col.header)] + sample) * 1.05 + 3
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = min(max(width, 9), 70)
    ws.freeze_panes = "A2"
    if block.rows:
        ws.auto_filter.ref = f"A1:{ws.cell(row=len(block.rows) + 1, column=len(columns)).column_letter}{len(block.rows) + 1}"
    if block.note:
        note_row = len(block.rows) + 3
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=max(len(columns), 1))
        _put(ws, (note_row, 1), block.note, _font(9, italic=True, colour=theme.SLATE), None, Alignment(wrap_text=True, vertical="top"))
        ws.row_dimensions[note_row].height = 30
    _print_setup(ws, meta, landscape=True, title_rows="1:1")
    return ws


# ── Entry point ─────────────────────────────────────────────────────────────

def render_xlsx(report: Report) -> bytes:
    """The report as XLSX bytes."""
    meta = report.meta
    wb = Workbook()
    used = {"cover", "summary"}
    summary_present = any(isinstance(b, (Heading, Kpis, Chart, Callout, Paragraphs)) for b in report.blocks)
    table_blocks = [b for b in report.blocks if isinstance(b, Table)]
    # Sheet names are fixed first, so the cover can list them.
    plan = [("Cover", "the title, scope, who generated it and when")]
    if summary_present:
        plan.append(("Summary", "headline measures and charts, from the same figures as the tables"))
    names = [_safe_sheet_name(t.sheet_name or t.title or "Data", used) for t in table_blocks]
    for name, block in zip(names, table_blocks):
        plan.append((name, f"{block.title or name}, {len(block.rows)} row{'s' if len(block.rows) != 1 else ''}"))
    _cover(wb, report, [f"{name}: {what}" for name, what in plan])
    if summary_present:
        _summary(wb, report)
    for name, block in zip(names, table_blocks):
        _data_sheet(wb, block, meta, name)
    props = wb.properties
    props.title = identity.clean_text(meta.title, keep_newlines=False)[:200]
    props.creator = identity.clean_text(meta.who, keep_newlines=False)[:200]
    props.subject = f"{meta.module_label}: {meta.tier_label}"
    props.keywords = f"{theme.PRODUCT_NAME}, {meta.tier_label}"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
