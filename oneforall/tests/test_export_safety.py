"""
Data exports and workbooks: scope, formula safety, file names and the tables they read.

A spreadsheet reads a cell that starts with =, +, - or @ as a formula, so a title typed as `=HYPERLINK(...)` would run when
someone opens the export. Every export writes such a cell as text. The exports are also held to the scope of the page they
export: an operational event, an AIMS assessment or a risk of another business unit is neither listed nor named in a download.
Expectations are written out by hand from the seeded file, and each "cannot" has a positive twin.
"""
import csv
import io
import json
import re

import openpyxl
import pytest
from fastapi import FastAPI
from starlette.testclient import TestClient

import core.middleware as middleware
import modules.aria.routes as aria_routes
import modules.bcm.routes as bcm_routes
import modules.erm.routes as erm_routes
import modules.grid.routes as grid_routes
import modules.orm.routes as orm_routes
from core import rbac
from modules.launcher.routes import router as launcher_router
from tests.test_platform_ownership import _insert, world  # noqa: F401

FORMULA = '=HYPERLINK("http://example.test","click")'


@pytest.fixture
def web(world, monkeypatch):
    who = {"name": "super"}

    async def current_user(request):
        return world.users[who["name"]]

    monkeypatch.setattr(middleware, "get_current_user", current_user)
    app = FastAPI()
    for router in (launcher_router, aria_routes.router, erm_routes.router, grid_routes.router, orm_routes.router, bcm_routes.router):
        app.include_router(router)
    client = TestClient(app, base_url="http://testserver", raise_server_exceptions=False, follow_redirects=False)

    def call(path, name="super"):
        who["name"] = name
        return client.get(path)

    def post(path, name="super", **kwargs):
        who["name"] = name
        return client.post(path, **kwargs)

    call.post = post
    call.world = world
    for name, tag in (("owner_a1", "A1"), ("owner_none", None)):
        unit = world.unit.get(tag) if tag else None
        uid = 300 + len(world.users)
        world.users[name] = {"id": uid, "username": name, "org_id": 1, "business_unit_id": unit, "is_super_admin": 0, "roles": [rbac.RISK_OWNER]}
        _insert(world.db, "users", id=uid, username=name, email=f"{name}@example.test", full_name=name, password_hash="x", org_id=1,
                business_unit_id=unit)
    return call


def _rows(text):
    return list(csv.DictReader(io.StringIO(text)))


SAFE_NAME = r'attachment; filename="[A-Za-z0-9._-]+"'


# ── Operations Risk: the events export is scoped like the events list ───────

ORM_SEES = {"super": {"org", "A", "A1", "B"}, "owner_a": {"org", "A", "A1"}, "owner_a1": {"org", "A1"}, "owner_b": {"org", "B"},
            "owner_none": {"org"}}


@pytest.mark.parametrize("persona", list(ORM_SEES))
def test_the_events_export_holds_the_events_the_viewer_may_open_and_no_others(web, persona):
    """It selected every event of every unit for anyone with ORM access."""
    for tag, unit_id in web.world.unit.items():
        _insert(web.world.db, "orm_events", title=f"export event {tag}", business_unit_id=unit_id)
    answer = web("/orm/api/export/csv", name=persona)
    assert answer.status_code == 200
    titles = {r["title"] for r in _rows(answer.text) if r["title"].startswith("export event ")}
    assert titles == {f"export event {tag}" for tag in ORM_SEES[persona]}
    assert re.fullmatch(r'attachment; filename="ThemisIQ_operations_risk_events_\d{8}_\d{4}\.csv"', answer.headers["content-disposition"])
    assert answer.headers["x-content-type-options"] == "nosniff" and answer.headers["cache-control"] == "no-store"


def test_the_events_export_keeps_its_columns_and_is_refused_to_roles_without_orm(web):
    """The layout of a data export is left as it was: the same ten columns in the same order, with the reporter by name."""
    _insert(web.world.db, "orm_events", title="Event", business_unit_id=None)
    header = next(csv.reader(io.StringIO(web("/orm/api/export/csv").text)))
    assert header == ["title", "event_type", "severity", "status", "department", "financial_impact", "root_cause", "reporter_name",
                      "created_at", "resolved_at"]
    for persona in ("dpo_a", "employee_a"):
        assert web("/orm/api/export/csv", name=persona).status_code == 403, persona


def test_an_event_title_that_looks_like_a_formula_is_exported_as_text(web):
    _insert(web.world.db, "orm_events", title=FORMULA, root_cause="+1+1", business_unit_id=None)
    row = next(r for r in _rows(web("/orm/api/export/csv").text) if "HYPERLINK" in r["title"])
    assert row["title"] == "'" + FORMULA and row["root_cause"] == "'+1+1"


# ── Operations Risk: an AIMS assessment is read only by those who may open its unit ──

def test_an_aims_assessment_is_read_only_through_its_own_unit(web):
    """The list was scoped; the read of one assessment, its risks and its aggregation were not."""
    db, unit = web.world.db, web.world.unit
    mine = _insert(db, "aims_assessments", ref="AIMS-A", title="Assessment A", business_unit_id=unit["A"])
    theirs = _insert(db, "aims_assessments", ref="AIMS-B", title="Assessment B", business_unit_id=unit["B"])
    mine_risk = _insert(db, "aims_risks", assessment_id=mine, risk_description="A risk of A")
    their_risk = _insert(db, "aims_risks", assessment_id=theirs, risk_description="A risk of B")
    for path in (f"/orm/api/aims/{mine}", f"/orm/api/aims/{mine}/risks", f"/orm/api/aims/risks/{mine_risk}",
                 f"/orm/api/aims/{mine}/aggregation", f"/orm/api/aims/{mine}/aggregation/csv"):
        assert web(path, name="owner_a").status_code == 200, path
    for path in (f"/orm/api/aims/{theirs}", f"/orm/api/aims/{theirs}/risks", f"/orm/api/aims/risks/{their_risk}",
                 f"/orm/api/aims/{theirs}/aggregation", f"/orm/api/aims/{theirs}/aggregation/csv"):
        assert web(path, name="owner_a").status_code == 404, path
        assert web(path, name="super").status_code == 200, path
    csv_answer = web(f"/orm/api/aims/{mine}/aggregation/csv", name="owner_a")
    assert re.fullmatch(rf'attachment; filename="ThemisIQ_aims_aggregation_{mine}.csv"', csv_answer.headers["content-disposition"])
    assert _rows(csv_answer.text)[0]["risk_description"] == "A risk of A"


# ── The other data exports: formula safe, branded names, safe headers ───────

def test_the_enterprise_risk_and_continuity_exports_write_a_formula_as_text(web):
    db = web.world.db
    _insert(db, "erm_enterprise_risks", title=FORMULA, business_unit_id=None)
    _insert(db, "bcm_incidents", title="+SUM(1+1)", business_unit_id=None)
    risks = web("/erm/api/export/csv")
    assert next(r for r in _rows(risks.text) if "HYPERLINK" in r["name"])["name"] == "'" + FORMULA
    assert re.fullmatch(r'attachment; filename="ThemisIQ_enterprise_risk_register_\d{8}_\d{4}\.csv"', risks.headers["content-disposition"])
    incidents = web("/bcm/api/export/csv")
    assert next(r for r in _rows(incidents.text) if "SUM" in r["title"])["title"] == "'+SUM(1+1)"
    assert re.fullmatch(r'attachment; filename="ThemisIQ_continuity_incidents_\d{8}_\d{4}\.csv"', incidents.headers["content-disposition"])


def test_the_platform_audit_log_export_is_the_callers_organization_and_formula_safe(web):
    db = web.world.db
    db.execute("INSERT INTO organizations (id, name, slug) VALUES (2, 'other', 'other')")
    for org, details in ((1, FORMULA), (1, "plain"), (2, "another organization")):
        db.execute("INSERT INTO audit_log (username, module, action, entity_type, details, org_id) VALUES ('someone', 'platform', 'x', 'y', %s, %s)",
                   (details, org))
    db.commit()
    mine = web("/admin/api/logs/export", name="dpo_a")
    assert mine.status_code == 200 and re.fullmatch(SAFE_NAME, mine.headers["content-disposition"])
    text = mine.text
    assert "'" + FORMULA in text.replace('""', '"') or "'=HYPERLINK" in text
    assert "plain" in text and "another organization" not in text
    assert "another organization" in web("/admin/api/logs/export", name="super").text, "the positive twin: a super administrator reads all"


def test_the_bulk_export_of_controls_is_the_register_the_app_shows(web):
    """It read `aria_controls`, a legacy table nothing writes to, so the export held other rows than the page."""
    framework = web.world.db.execute("SELECT id FROM frameworks ORDER BY id LIMIT 1").fetchone()["id"]
    _insert(web.world.db, "controls", framework_id=framework, ref="EXPORT-1", name="A control the page shows")
    answer = web("/api/bulk/export/controls")
    assert answer.status_code == 200
    assert "EXPORT-1" in {row["ref"] for row in json.loads(answer.text)}
    assert re.fullmatch(r'attachment; filename="ThemisIQ_controls_export.json"', answer.headers["content-disposition"])
    assert web("/api/bulk/export/controls", name="dpo_a").status_code == 403, "whole tables are for the super administrator"


# ── The workbooks ───────────────────────────────────────────────────────────

def test_the_users_workbook_has_a_cover_and_goes_back_through_the_importer(web):
    from modules.launcher._user_import import parse_users_excel
    _insert(web.world.db, "users", username="formula.user", email="formula.user@example.test", full_name=FORMULA, password_hash="x", org_id=1)
    answer = web("/admin/api/users/export")
    assert answer.status_code == 200 and re.fullmatch(r'attachment; filename="ThemisIQ_user_accounts_\d{8}_\d{4}\.xlsx"', answer.headers["content-disposition"])
    book = openpyxl.load_workbook(io.BytesIO(answer.content))
    assert book.sheetnames[0] == "Cover" and "Users" in book.sheetnames
    users = book["Users"]
    assert [c.value for c in users[1]][:3] == ["Organization", "Full Name", "Email"]
    names = [c for row in users.iter_rows(min_row=2) for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert names and all(c.data_type == "s" for c in names), "typed as text, never a formula"
    parsed = parse_users_excel(answer.content, is_super=True, caller_org_id=1)
    assert "error" not in parsed, "the importer finds the Users sheet behind the cover"
    assert "formula.user@example.test" in {r["email"] for r in parsed["rows"]}
    assert len(parsed["rows"]) == users.max_row - 1


def test_the_controls_workbook_has_the_cover_the_summary_and_a_sheet_per_framework_and_never_evaluates_a_name(web):
    db = web.world.db
    framework = db.execute("SELECT id, name FROM frameworks WHERE is_active = 1 AND name = 'GDPR'").fetchone()
    for n, (status, name) in enumerate((("Implemented", FORMULA), ("Implemented", "Second"), ("Not Started", "Third"))):
        _insert(db, "controls", framework_id=framework["id"], ref=f"WB-{n}", name=name, status=status)
    answer = web("/aria/export/excel", name="dpo_a")
    assert answer.status_code == 200
    assert re.fullmatch(r'attachment; filename="ThemisIQ_Governance_controls_register_\d{8}_\d{4}\.xlsx"', answer.headers["content-disposition"])
    book = openpyxl.load_workbook(io.BytesIO(answer.content))
    assert book.sheetnames[:2] == ["Cover", "Summary"] and "GDPR" in book.sheetnames
    sheet = book["GDPR"]
    cells = [c for row in sheet.iter_rows(min_row=2) for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert cells and all(c.data_type == "s" for c in cells)
    assert {row[0] for row in sheet.iter_rows(min_row=2, values_only=True)} >= {"WB-0", "WB-1", "WB-2"}
    footer = sheet.oddFooter.left.text or ""
    assert "Module: Governance" in footer and "Generated by dpo_a" in footer


# ── Bundle and document downloads named from a person's text ─────────────────

HOSTILE = 'Q3 "audit" “smart” 安全 \r\nX-Injected: 1'


def test_an_evidence_bundle_is_named_in_ascii_whatever_the_audit_or_the_control_is_called(web):
    """The bundle name was the audit's or the control's own text: a smart quote or a Chinese letter made the header impossible
    to encode (HTTP 500) and a straight quote ended the file name early."""
    db, unit = web.world.db, web.world.unit
    fw = _insert(db, "grid_frameworks", name="ISO 27001:2022")
    audit = _insert(db, "grid_audits", name=HOSTILE, framework_id=fw, status="Active", business_unit_id=unit["A"])
    control = _insert(db, "grid_controls", audit_id=audit, framework_id=fw, control_id='A.5 "q" “s”', name="A control",
                      status="Not Started")
    item = _insert(db, "grid_evidence_items", control_id=control, name="Evidence")
    _insert(db, "grid_evidence_files", evidence_item_id=item, control_id=control, filename="x.pdf", original_name="x.pdf",
            file_path="missing/x.pdf", status="Uploaded")
    for path in (f"/grid/api/evidence/{control}/download-all", f"/grid/api/evidence/audit/{audit}/download-all"):
        answer = web(path, name="dpo_a")
        assert answer.status_code == 200, path
        assert re.fullmatch(r'attachment; filename="evidence_[A-Za-z0-9._-]+\.zip"', answer.headers["content-disposition"]), \
            answer.headers["content-disposition"]
        assert answer.headers["x-content-type-options"] == "nosniff" and answer.headers["cache-control"] == "no-store"
        assert "x-injected" not in answer.headers
        assert answer.content[:2] == b"PK"


def test_a_policy_document_download_is_named_in_ascii_from_whatever_heading_it_has(web):
    """The name was the first heading of the text with the word characters kept, and Python's word characters include Chinese letters."""
    heading = '# 安全政策 "quoted" Policy\n\nBody text.'
    answer = web.post("/aria/api/export-word", name="dpo_a", data={"content": heading, "org_name": "Org"})
    assert answer.status_code == 200, answer.text[:200]
    assert re.fullmatch(r'attachment; filename="[A-Za-z0-9._-]+\.docx"', answer.headers["content-disposition"]), \
        answer.headers["content-disposition"]
    assert answer.content[:2] == b"PK" and answer.headers["x-content-type-options"] == "nosniff"
    plain = web.post("/aria/api/export-word", name="dpo_a", data={"content": "# Information Security Policy\n\nBody.", "org_name": "Org"})
    assert plain.headers["content-disposition"] == 'attachment; filename="Information_Security_Policy.docx"'
