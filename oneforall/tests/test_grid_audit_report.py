"""
The audit report of GRID: PDF and Word, operational and board summary.

The figures are written by hand from one seeded audit. Pinned: the figures and chart series are the audit's own (counted as the
Audit pages count them), the evidence counts include files awaiting review (the old report left the default status out), a
narrative that states a number the audit does not hold is left out, the cover and footer are on every page, and a person who
may not open the audit may not download its report.
"""
import asyncio

import pytest
from fastapi import HTTPException
from pypdf import PdfReader

import modules.grid.routes as grid_routes
from core import rbac
from modules.grid import report_service as rpt
from tests.test_platform_ownership import ORG, _insert, _request, signed_in, world  # noqa: F401


@pytest.fixture
def audit(world, tmp_path, monkeypatch):
    """One audit with eight controls, three findings and five evidence files, all in a known state."""
    monkeypatch.setattr(rpt, "REPORTS_DIR", tmp_path / "reports")
    db = world.db
    fid = _insert(db, "grid_frameworks", name="ISO 27001")
    aid = _insert(db, "grid_audits", name="Q3 internal audit", framework_id=fid, status="Active", business_unit_id=world.unit["A"],
                  start_date="2026-07-01", end_date="2026-09-30", audit_date="2026-09-15")
    past, future = "2020-01-01", "2999-01-01"
    rows = [("A.1", "Access control policy", "Critical", "Complete", past), ("A.2", "Key management", "Critical", "In Progress", past),
            ("A.3", "Logging and monitoring", "High", "Not Started", future), ("A.4", "Backup", "High", "Complete", future),
            ("A.5", "Awareness training", "Low", "In Progress", future), ("A.6", "Supplier review", "Medium", "Not Started", past),
            ("A.7", "Asset register", "Medium", "Complete", None), ("A.8", "Change control", "Low", "Not Started", None)]
    ids = {}
    for ref, name, risk, status, due in rows:
        ids[ref] = _insert(db, "grid_controls", audit_id=aid, framework_id=fid, control_id=ref, name=name, risk_level=risk,
                           status=status, due_date=due)
    # A.1 has two evidence items: one holds an approved file, one a file awaiting review (status Uploaded, the default).
    e1 = _insert(db, "grid_evidence_items", control_id=ids["A.1"], name="Policy")
    e2 = _insert(db, "grid_evidence_items", control_id=ids["A.1"], name="Approval record")
    _insert(db, "grid_evidence_items", control_id=ids["A.2"], name="Key register")
    for item, control, status in ((e1, "A.1", "Approved"), (e2, "A.1", "Uploaded"), (None, "A.4", "Rejected"),
                                  (None, "A.4", "Uploaded"), (None, "A.7", "approved")):
        _insert(db, "grid_evidence_files", evidence_item_id=item, control_id=ids[control], filename="f.pdf", original_name="f.pdf",
                file_path="x", status=status)
    _insert(db, "grid_non_conformances", audit_id=aid, title="Keys not rotated", severity="critical", status="open", cap_status="Open")
    _insert(db, "grid_non_conformances", audit_id=aid, title="Old finding", severity="major", status="open", cap_status="Closed")
    _insert(db, "grid_non_conformances", audit_id=aid, title="Closed finding", severity="minor", status="closed", cap_status="Closed")
    world.audit_id = aid
    for uid, name, full, unit in ((50, "lead", "Rudo Lead", "A"), (51, "lead_b", "Other Lead", "B")):
        world.users[name] = {"id": uid, "username": name, "full_name": full, "org_id": ORG, "business_unit_id": world.unit[unit],
                             "is_super_admin": 0, "roles": [rbac.AUDIT_LEAD]}
        _insert(db, "users", id=uid, username=name, email=f"{name}@example.test", full_name=full, password_hash="x", org_id=ORG,
                business_unit_id=world.unit[unit])
    return world


def _data(world):
    data = rpt.gather(world.audit_id)
    return data, rpt.figures(data)


def test_the_figures_are_counted_as_the_audit_pages_count_them(audit):
    data, fig = _data(audit)
    assert (fig["total"], fig["complete"], fig["not_complete"]) == (8, 3, 5)
    assert fig["overdue"] == 2, "A.2 (in progress) and A.6 (not started) are past their date; A.1 is complete, so it is not overdue"
    assert fig["in_progress"] == 1 and fig["not_started"] == 2 and fig["completion_pct"] == 38, "3 of 8 is 37.5, which rounds to 38"
    assert fig["critical_high_open"] == 2 and {c["control_id"] for c in fig["gaps"]} == {"A.2", "A.3"}
    assert fig["by_risk"]["Critical"] == {"complete": 1, "not_complete": 1} and fig["by_risk"]["Low"] == {"complete": 0, "not_complete": 2}
    assert (fig["nc_total"], fig["nc_open"], fig["nc_closed"], fig["nc_open_critical"]) == (3, 1, 2, 1), \
        "a finding is closed when its status or its action plan says closed"
    assert fig["nc_by_severity"] == {"critical": {"open": 1, "closed": 0}, "major": {"open": 0, "closed": 1}, "minor": {"open": 0, "closed": 1}}


def test_the_evidence_counts_include_files_awaiting_review_and_agree_with_the_evidence_tab(audit):
    _, fig = _data(audit)
    assert (fig["files_total"], fig["files_approved"], fig["files_pending"], fig["files_rejected"], fig["files_other"]) == (5, 2, 2, 1, 0), \
        "two files are approved (one written in lower case), two are Uploaded (the default, awaiting review) and one is rejected"
    assert fig["items_total"] == 3 and fig["items_with_evidence"] == 2
    from modules.grid import data_service as ds
    tab = ds.get_evidence_completeness(audit.audit_id)
    assert (tab["totalFiles"], tab["approvedFiles"], tab["pendingFiles"], tab["rejectedFiles"]) == (5, 2, 2, 1)


def test_the_report_blocks_carry_those_figures_and_the_control_rows_say_what_evidence_covers(audit):
    data, fig = _data(audit)
    report = rpt.build_report(data, actor=audit.users["lead"], tier="operational")
    kpis = {k.label: k for b in report.blocks if b.__class__.__name__ == "Kpis" for k in b.items}
    assert kpis["Complete"].value == "3" and kpis["Complete"].note == "38% of controls" and kpis["Overdue"].value == "2"
    assert kpis["Open findings"].value == "1" and kpis["Evidence approved"].note == "of 5 files (40%)"
    donut = next(c for c in report.charts() if c.title == "Controls by status")
    assert donut.data == [("Complete", 3), ("In progress", 1), ("Not started", 2), ("Overdue", 2)] and donut.centre == "38%"
    assert sum(v for _, v in donut.data) == fig["total"], "every control is in exactly one state"
    stack = next(c for c in report.charts() if c.title.startswith("Controls by risk level"))
    assert stack.data[0] == ("Critical", [("Complete", 1), ("Not complete", 1)])
    controls = next(t for t in report.tables() if t.title == "Every control in this audit")
    by_ref = {r["ref"]: r for r in controls.rows}
    assert by_ref["A.1"]["evidence"] == "2 of 2 items" and by_ref["A.2"]["evidence"] == "0 of 1 items"
    assert by_ref["A.8"]["evidence"] == "No checklist", "a control with no evidence items is not shown as 0 of 1"
    assert by_ref["A.2"]["status"] == "Overdue" and by_ref["A.3"]["status"] == "Not Started" and by_ref["A.1"]["status"] == "Complete"
    assert len(controls.rows) == 8, "every control is listed (the old report cut the list at 60 without saying so)"


def test_the_board_summary_is_short_and_leads_with_the_decisions(audit):
    data, _ = _data(audit)
    board = rpt.build_report(data, actor=audit.users["lead"], tier="board")
    ops = rpt.build_report(data, actor=audit.users["lead"], tier="operational")
    assert board.meta.tier == "board" and board.meta.tier_label == "Board report"
    decision = next(b for b in board.blocks if b.__class__.__name__ == "Callout" and b.kind == "decision")
    assert decision.items == [
        "1 critical finding still open: confirm the owner and the due date of each.",
        "2 overdue controls: agree a recovery date for each.",
        "2 critical or high risk controls not complete: confirm who completes them and by when.",
        "1 evidence file rejected: confirm who resubmits."]
    assert len(board.tables()) == 1 and board.tables()[0].title.startswith("Critical and high risk controls")
    assert [r["ref"] for r in board.tables()[0].rows] == ["A.2", "A.3"], "earliest due date first"
    assert len(ops.tables()) >= 2 and len(board.blocks) < len(ops.blocks)


def test_a_narrative_with_a_number_the_audit_does_not_hold_is_left_out(audit):
    data, fig = _data(audit)
    good = {"executive_summary": "3 of 8 controls (38%) are complete and 2 are overdue.", "key_findings": ["1 critical finding is open."],
            "conclusion": "The audit is in progress."}
    kept, note = rpt.checked_narrative(data, fig, good)
    assert kept is good and "found in the data" in note
    bad = {**good, "executive_summary": "3 of 8 controls are complete and 41 are overdue."}
    dropped, note = rpt.checked_narrative(data, fig, bad)
    assert dropped is None and "41" in note
    report = rpt.build_report(data, actor=audit.users["lead"], narrative=dropped, narrative_note=note)
    assert not [b for b in report.blocks if b.__class__.__name__ == "Paragraphs"]
    assert any(b.__class__.__name__ == "Callout" and b.title == "Draft narrative left out" for b in report.blocks)


def test_the_model_is_given_the_report_figures_and_not_a_different_pending_count(audit):
    data, fig = _data(audit)
    given = rpt.narrative_inputs(data, fig)
    assert (given["totalControls"], given["complete"], given["pending"], given["overdue"], given["completionPct"]) == (8, 3, 5, 2, 38)
    assert given["criticalGaps"] == ["Key management", "Logging and monitoring"], "critical and high risk controls that are not complete"


def _pdf_pages(path):
    return [p.extract_text() for p in PdfReader(path).pages]


@pytest.mark.parametrize("tier, label", [("operational", "Operational report"), ("board", "Board report")])
def test_the_downloaded_pdf_has_a_cover_with_the_logo_and_a_footer_on_every_page(audit, tier, label):
    req = _request(audit.users["lead"], {"tier": tier})
    response = asyncio.run(grid_routes.api_report_pdf(req, audit.audit_id))
    assert response.status_code == 200 and response.media_type == "application/pdf"
    reader = PdfReader(response.path)
    pages = [p.extract_text() for p in reader.pages]
    assert len(reader.pages[0].images) >= 1, "the ThemisIQ logo is on the cover"
    assert "Q3 internal audit" in pages[0] and "Business unit: Unit A" in pages[0] and label in pages[0]
    for number, text in enumerate(pages, start=1):
        assert "Module: Audit" in text and "Generated by Rudo Lead (lead)" in text and "UTC" in text, f"page {number}"
        assert f"Page {number} of {len(pages)}" in text and label in text
    assert "Complete" in "".join(pages) and "38%" in "".join(pages)
    saved = audit.db.execute("SELECT title, report_type FROM grid_reports ORDER BY id DESC").fetchone()
    assert saved["report_type"] == "pdf" and saved["title"] == ("Q3 internal audit: Board summary (PDF)" if tier == "board" else "Q3 internal audit: Audit report (PDF)")


def test_the_word_report_has_a_real_footer_with_page_fields(audit):
    from docx import Document
    response = asyncio.run(grid_routes.api_report_docx(_request(audit.users["lead"], {"tier": "board"}), audit.audit_id))
    assert response.status_code == 200 and response.filename.endswith(".docx")
    doc = Document(response.path)
    footer_xml = "".join(s.footer._element.xml for s in doc.sections)
    assert "Module: Audit" in footer_xml and "Generated by Rudo Lead (lead)" in footer_xml and "Board report" in footer_xml
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml


def test_an_unknown_tier_is_refused_and_a_person_outside_the_audits_unit_cannot_download_it(audit):
    with pytest.raises(HTTPException) as bad:
        asyncio.run(grid_routes.api_report_pdf(_request(audit.users["lead"], {"tier": "secret"}), audit.audit_id))
    assert bad.value.status_code == 400
    with pytest.raises(HTTPException) as hidden:
        asyncio.run(grid_routes.api_report_pdf(_request(audit.users["lead_b"]), audit.audit_id))
    assert hidden.value.status_code == 404, "a person who may not open the audit gets the answer a missing audit gets"
