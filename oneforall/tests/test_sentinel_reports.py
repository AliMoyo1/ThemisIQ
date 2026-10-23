"""
Privacy (Sentinel) reports: the record of processing activities and the breach register (operational), the privacy
compliance score (the board report), the evidence pack and the audit trail export.

What a report states is what the same person sees on screen. The expectations are written out by hand from a seeded file;
the report, the lists, the score and the dashboard are asked the same question and must give the same answer. Each "cannot"
(a record of another unit, a part the role may not open) has a positive twin, so an empty document cannot pass for isolation.
A test of a fixed bug is written so that the old behaviour fails it: the old score read an empty part as 100, rounded half to
even, and counted an emptied data category list as a category; the old pack handed every part to anyone who could open the
module, stopped at 500 rows and put the organization name into the download header as it was typed.
"""
import csv
import io
import json
import re
import zipfile

import openpyxl
import pytest
from pypdf import PdfReader

from core import rbac
from core.report_kit import theme
from core.report_kit.model import Kpis
from modules.sentinel import dashboard_metrics, data_service as ds, export_reports as exp
from tests.test_platform_ownership import _insert, world  # noqa: F401
from tests.test_sentinel_write_scope import http  # noqa: F401

PERSONAS = ("super", "dpo_a", "dpo_a1", "dpo_b", "dpo_no_unit")


# ── Helpers ─────────────────────────────────────────────────────────────────

def _kpis(report):
    return {k.label: k for block in report.blocks if isinstance(block, Kpis) for k in block.items}


def _table(report, starts_with):
    """The one table whose title begins with `starts_with` (the RoPA report has a register per role and a readiness list)."""
    found = [t for t in report.tables() if str(t.title).startswith(starts_with)]
    assert len(found) == 1, (starts_with, [t.title for t in report.tables()])
    return found[0]


def _chart(report, title):
    return next(c for c in report.charts() if c.title == title)


def _pages(pdf_bytes):
    return [page.extract_text() for page in PdfReader(io.BytesIO(pdf_bytes)).pages]


def _persona(world, name, role, unit_tag=None):
    unit = world.unit.get(unit_tag) if unit_tag else None
    uid = 100 + len(world.users)
    world.users[name] = {"id": uid, "username": name, "org_id": 1, "business_unit_id": unit, "is_super_admin": 0, "roles": [role]}
    _insert(world.db, "users", id=uid, username=name, email=f"{name}@example.test", full_name=name, password_hash="x", org_id=1,
            business_unit_id=unit)
    return world.users[name]


def _ropa(world, tag, n, **cols):
    return _insert(world.db, "sentinel_ropa", ref_number=f"R-{tag}-{n}", processing_name=f"{tag} activity {n}",
                   business_unit_id=world.unit[tag], **cols)


COMPLETE = {"purpose": "To run it", "data_categories": '["Payroll"]'}


def _file(world):
    """RoPA entries. A: A1 high, needs a DPIA and has none; A2 medium, special category data, review long overdue.
    B: B1 critical with a DPIA on file; B2 with no legal basis. Organization wide: one whose category list was emptied.
    DSRs: A open past its deadline, A1 closed, B open in the future. The fixture's four breaches are all open."""
    db, unit = world.db, world.unit
    _ropa(world, "A", 1, legal_basis="Contract", department="Human Resources", risk_level="high", dpia_required=1, **COMPLETE)
    _ropa(world, "A", 2, legal_basis="Consent", department="Marketing", risk_level="medium", special_categories='["Health"]',
          review_date="2001-01-01", **COMPLETE)
    dpia = _insert(db, "sentinel_dpias", ref_number="D-B", title="DPIA for B1", status="approved", business_unit_id=unit["B"])
    _ropa(world, "B", 1, legal_basis="Legal obligation", department="Customer Operations", risk_level="critical", dpia_required=1,
          dpia_id=dpia, **COMPLETE)
    _ropa(world, "B", 2, department="Customer Operations", risk_level="medium", **COMPLETE)
    _insert(db, "sentinel_ropa", ref_number="R-org-1", processing_name="Visitor register", legal_basis="Consent",
            purpose="To keep the building safe", data_categories="[]", risk_level="low", business_unit_id=None)
    for tag, status, deadline in (("A", "open", "2001-01-01"), ("A1", "closed", "2001-01-01"), ("B", "open", "2999-01-01")):
        _insert(db, "sentinel_dsr", ref_number=f"S-{tag}", requester_name=f"Requester {tag}", status=status,
                deadline_date=deadline, business_unit_id=unit[tag])


# ── The compliance score ────────────────────────────────────────────────────

def test_a_part_with_nothing_to_measure_is_not_scored_and_an_empty_module_has_no_grade(test_db):
    score = ds.get_compliance_score()
    assert [p["score"] for p in score["breakdown"].values()] == [None] * 5
    assert score["overall"] is None and score["grade"] == "N/A"


def test_the_overall_is_worked_out_over_the_parts_that_were_measured(test_db):
    """One breach, closed: the only part measured is at 100, so the overall is 100. The old score filled the empty parts
    with 100 (and the RoPA part with 0), so a module holding one closed breach read as 75."""
    _insert(test_db, "sentinel_breaches", ref_number="B-1", title="One", status="closed")
    score = ds.get_compliance_score()
    assert score["breakdown"]["breach_response"]["score"] == 100
    assert score["breakdown"]["ropa_completeness"]["score"] is None
    assert score["overall"] == 100 and score["grade"] == "A"


def test_the_weights_of_the_measured_parts_are_scaled_to_one_hundred(test_db):
    _insert(test_db, "sentinel_ropa", ref_number="R-1", processing_name="Whole", purpose="p", legal_basis="Contract",
            data_categories='["Name"]')
    _insert(test_db, "sentinel_ropa", ref_number="R-2", processing_name="Half")
    _insert(test_db, "sentinel_breaches", ref_number="B-1", title="One", status="resolved")
    score = ds.get_compliance_score()
    assert score["breakdown"]["ropa_completeness"]["score"] == 50 and score["breakdown"]["breach_response"]["score"] == 100
    # (50 * 25 + 100 * 20) / 45 = 72.2
    assert score["overall"] == 72 and score["grade"] == "C"
    assert score["breakdown"]["dpia_coverage"]["score"] is None


def test_percentages_round_half_up(test_db):
    """5 of 8 breaches done is 62.5%: a person, the page and the reports say 63. round() says 62."""
    for n in range(8):
        _insert(test_db, "sentinel_breaches", ref_number=f"B-{n}", title=f"B {n}", status="closed" if n < 5 else "open")
    assert ds.get_compliance_score()["breakdown"]["breach_response"]["score"] == 63


def test_an_emptied_data_category_list_is_not_a_category(test_db):
    """The list is stored as JSON, so an entry saved with it emptied holds '[]'. It used to count as complete."""
    _insert(test_db, "sentinel_ropa", ref_number="R-1", processing_name="Named", purpose="p", legal_basis="Contract",
            data_categories='["Name"]')
    _insert(test_db, "sentinel_ropa", ref_number="R-2", processing_name="Emptied", purpose="p", legal_basis="Contract",
            data_categories="[]")
    part = ds.get_compliance_score()["breakdown"]["ropa_completeness"]
    assert (part["done"], part["total"], part["score"]) == (1, 2, 50)


# ── What each report states ─────────────────────────────────────────────────

def test_the_ropa_report_counts_the_records_the_caller_can_open(http):
    world = http.world
    _file(world)
    # persona: (records, need a DPIA, have one, incomplete, high or critical)
    expected = {"super": (5, 2, 1, 2, 2), "dpo_a": (3, 1, 0, 1, 1), "dpo_a1": (1, 0, 0, 1, 0), "dpo_b": (3, 1, 1, 2, 1),
                "dpo_no_unit": (1, 0, 0, 1, 0)}
    for persona, (total, needs, has, incomplete, high) in expected.items():
        user = world.users[persona]
        listed = http("get", "/sentinel/api/ropa", name=persona).json()
        report = exp.ropa_report(user)
        kpis = _kpis(report)
        assert len(listed) == total == len(_table(report, "Record readiness").rows), persona
        assert len(_table(report, "Controller activities").rows) + len(_table(report, "Processor activities").rows) == total, persona
        assert kpis["Processing activities"].value == str(total), persona
        assert kpis["DPIA still to do"].value == str(needs - has), persona
        assert kpis["Incomplete records"].value == str(incomplete), persona
        assert kpis["High or critical risk"].value == str(high), persona
        # the same counts the compliance score of that person holds
        score = http("get", "/sentinel/api/compliance-score", name=persona).json()["breakdown"]
        assert (score["ropa_completeness"]["total"], total - score["ropa_completeness"]["done"]) == (total, incomplete), persona
        assert (score["dpia_coverage"]["total"], score["dpia_coverage"]["done"]) == (needs, has), persona


def test_the_ropa_charts_are_built_from_the_records_of_the_table(http):
    world = http.world
    _file(world)
    report = exp.ropa_report(world.users["super"])
    assert _chart(report, "Processing activities by legal basis").data == [
        ("Consent", 2), ("Contract", 1), ("Legal obligation", 1), ("Not recorded", 1)]
    assert _chart(report, "Processing activities by risk level").data == [("Critical", 1), ("High", 1), ("Medium", 2), ("Low", 1)]
    dpia = _chart(report, "DPIA position of the activities that need one")
    assert dpia.data == [("DPIA on file", 1), ("DPIA still to do", 1)] and dpia.centre == "50%"
    assert sum(v for _, v in _chart(report, "Processing activities by department").data) == 5
    assert _kpis(report)["Special category data"].value == "1" and _kpis(report)["Review overdue"].value == "1"
    rows = {r["name"]: r for r in _table(report, "Controller activities").rows}
    assert rows["A activity 1"]["dpia"] == "Still to do" and rows["B activity 1"]["dpia"] == "On file"
    assert rows["A activity 2"]["dpia"] == "Not required"
    assert rows["Visitor register"]["data_categories"] is None, "an emptied list is no category, in the table too"


def test_the_breach_report_agrees_with_the_list_and_the_dashboard(http):
    world = http.world
    for tag, severity, status, found, affected in (("A", "critical", "investigating", "2026-10-02", 40), ("A1", "high", "resolved", "2026-09-10", 7),
                                                    ("B", "low", "closed", None, None), ("org", "medium", "contained", "2026-10-15", 3)):
        world.db.execute("UPDATE sentinel_breaches SET severity=%s, status=%s, discovery_date=%s, affected_count=%s WHERE id=%s",
                         (severity, status, found, affected, world.rec[("breach", tag)]))
    world.db.commit()
    for persona in PERSONAS:
        user = world.users[persona]
        listed = http("get", "/sentinel/api/breaches", name=persona).json()
        report = exp.breach_report(user)
        kpis = _kpis(report)
        shown = dashboard_metrics.metrics(world.db, user)["sentinel.breaches_active"]["value"]
        assert len(listed) == len(report.tables()[0].rows) == int(kpis["Breaches recorded"].value), persona
        assert int(kpis["Active"].value) == shown == sum(b["status"] not in ("closed", "resolved") for b in listed), persona
        assert sum(v for _, v in _chart(report, "Breaches by severity").data) == len(listed), persona
        assert sum(v for _, v in _chart(report, "Breaches by status").data) == len(listed), persona
        counted = [b["affected_count"] for b in listed if b["affected_count"] is not None]
        assert kpis["People affected"].value == (f"{sum(counted):,}" if counted else "n/a"), persona


def test_the_breach_report_states_facts_and_calls_every_deadline_indicative(http):
    world = http.world
    for tag, notified in (("A", 0), ("B", 1)):
        world.db.execute("UPDATE sentinel_breaches SET notification_required=1, authority_notified=%s, discovery_date='2001-01-01', "
                         "notify_deadline='2001-01-04 00:00:00' WHERE id=%s", (notified, world.rec[("breach", tag)]))
    world.db.commit()
    report = exp.breach_report(world.users["super"])
    assert _kpis(report)["Authority not yet notified"].value == "1"
    notes = [c for c in report.blocks if c.__class__.__name__ == "Callout"]
    attention = next(c for c in notes if c.title == "Needs attention")
    assert any("past the indicative deadline" in item for item in attention.items)
    assert any("indicative" in item and "privacy specialist" in item for c in notes if c.kind == "limit" for item in c.items)
    rows = {r["ref"]: r for r in report.tables()[0].rows}
    assert rows["BR-A"]["authority"] == "Not yet" and rows["BR-B"]["authority"] == "Yes"
    assert rows["BR-org"]["authority"] == "Not marked as required"


def test_the_board_report_is_the_score_the_page_shows(http):
    world = http.world
    _file(world)
    for persona in PERSONAS:
        user = world.users[persona]
        shown = http("get", "/sentinel/api/compliance-score", name=persona).json()
        report = exp.score_report(user)
        assert report.meta.tier == theme.BOARD, "the score is the board report"
        want = f"{shown['overall']}%" if shown["overall"] is not None else "Not measured"
        assert _kpis(report)["Overall score"].value == want, persona
        table = {r["part"]: (r["done"], r["total"], r["score"]) for r in report.tables()[0].rows}
        assert table == {p["label"]: (p["done"], p["total"], p["score"]) for p in shown["breakdown"].values()}, persona
        bars = dict(_chart(report, "Score by part").data)
        assert bars == {p["label"]: p["score"] for p in shown["breakdown"].values() if p["score"] is not None}, persona
        dsrs = http("get", "/sentinel/api/dsrs", name=persona).json()
        open_dsr = [d for d in dsrs if d["status"] not in ("completed", "closed")]
        assert _kpis(report)["Open data subject requests"].value == str(len(open_dsr)), persona


def test_the_two_registers_are_operational_reports(http):
    world = http.world
    assert exp.ropa_report(world.users["super"]).meta.tier == theme.OPERATIONAL
    assert exp.breach_report(world.users["super"]).meta.tier == theme.OPERATIONAL


def test_an_empty_module_says_so_in_every_report(test_db):
    user = {"id": 1, "username": "nobody", "org_id": 1, "business_unit_id": None, "is_super_admin": 1, "roles": [rbac.SUPER_ADMIN]}
    board = exp.score_report(user)
    assert _kpis(board)["Overall score"].value == "Not measured"
    assert any("No part of the score could be measured" in item for b in board.blocks if b.__class__.__name__ == "Callout"
               for item in b.items)
    ropa = exp.ropa_report(user)
    register = _table(ropa, "Controller activities")
    assert register.rows == [] and "No processing activity" in register.empty_text
    assert _chart(ropa, "Processing activities by legal basis").is_empty()


# ── The documents ───────────────────────────────────────────────────────────

@pytest.mark.parametrize("kind, stem, tier_label", [("ropa", "ThemisIQ_privacy_ropa", "Operational report"),
                                                    ("breach", "ThemisIQ_privacy_breach_register", "Operational report"),
                                                    ("compliance", "ThemisIQ_privacy_compliance_score", "Board report")])
def test_each_report_downloads_as_a_pdf_and_a_workbook_with_a_cover_and_a_footer(http, kind, stem, tier_label):
    _file(http.world)
    pdf = http("get", f"/sentinel/api/reports/{kind}", name="dpo_a")
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf" and pdf.content.startswith(b"%PDF")
    assert re.fullmatch(rf'attachment; filename="{stem}_\d{{8}}_\d{{4}}\.pdf"', pdf.headers["content-disposition"])
    assert pdf.headers["x-content-type-options"] == "nosniff" and pdf.headers["cache-control"] == "no-store"
    reader = PdfReader(io.BytesIO(pdf.content))
    pages = _pages(pdf.content)
    assert len(reader.pages[0].images) >= 1, "the ThemisIQ logo is on the cover"
    assert tier_label in pages[0] and "Whole organization" not in pages[0], "the scope is the caller's own units"
    for number, text in enumerate(pages, start=1):
        assert "Module: Privacy" in text and "Generated by dpo_a" in text and "UTC" in text and "CAT" in text, number
        assert f"Page {number} of {len(pages)}" in text and tier_label in text, number
    book = http("get", f"/sentinel/api/reports/{kind}?format=xlsx", name="dpo_a")
    assert book.status_code == 200 and book.content.startswith(b"PK")
    assert openpyxl.load_workbook(io.BytesIO(book.content)).sheetnames[0] == "Cover"


def test_a_document_names_no_record_of_another_unit(http):
    world = http.world
    for persona, absent, present in (("dpo_a", "needle breach B", "needle breach A1"), ("dpo_b", "needle breach A", "needle breach B")):
        pdf = http("get", "/sentinel/api/reports/breach", name=persona)
        text = " ".join(_pages(pdf.content)).replace("\n", " ")
        assert present in text and absent not in text, persona
        book = openpyxl.load_workbook(io.BytesIO(http("get", "/sentinel/api/reports/breach?format=xlsx", name=persona).content))
        titles = {row[1] for row in book["Breaches"].iter_rows(min_row=2, values_only=True)}
        assert present in titles and absent not in titles, persona


def test_an_unknown_report_or_format_is_refused(http):
    assert http("get", "/sentinel/api/reports/nothing", name="dpo_a").status_code == 404
    assert http("get", "/sentinel/api/reports/ropa?format=docx", name="dpo_a").status_code == 400


def test_a_report_needs_the_permission_that_opens_the_same_records(http):
    world = http.world
    _persona(world, "analyst", rbac.PRIVACY_ANALYST, "A")
    # The breach list is closed to this role, so the breach register is too; the RoPA and the score are open to it.
    assert http("get", "/sentinel/api/breaches", name="analyst").status_code == 403
    assert http("get", "/sentinel/api/reports/breach", name="analyst").status_code == 403
    assert http("get", "/sentinel/api/reports/breach?format=xlsx", name="analyst").status_code == 403
    for kind in ("ropa", "compliance"):
        assert http("get", f"/sentinel/api/reports/{kind}", name="analyst").status_code == 200, kind
    assert http("get", "/sentinel/api/reports/breach", name="dpo_a").status_code == 200


def test_the_register_shows_what_is_typed_and_the_workbook_never_evaluates_it(http):
    world = http.world
    _insert(world.db, "sentinel_ropa", ref_number="R-X", processing_name='=HYPERLINK("http://example.test","click")', legal_basis="+1+1",
            business_unit_id=None)
    book = openpyxl.load_workbook(io.BytesIO(http("get", "/sentinel/api/reports/ropa?format=xlsx", name="super").content))
    cells = [c for row in book["Controller register"].iter_rows(min_row=2) for c in row if isinstance(c.value, str) and c.value.startswith(("=", "+"))]
    assert cells and all(c.data_type == "s" for c in cells)


def test_a_register_longer_than_the_document_limit_says_so_in_the_document_only(http):
    world = http.world
    for n in range(exp.ROWS_IN_DOCUMENT + 10):
        world.db.execute("INSERT INTO sentinel_ropa (ref_number, processing_name) VALUES (%s, %s)", (f"LONG-{n}", f"Long {n}"))
    world.db.commit()
    total = len(http("get", "/sentinel/api/ropa", name="super").json())
    text = " ".join(_pages(http("get", "/sentinel/api/reports/ropa", name="super").content)).replace("\n", " ")
    assert f"Showing the first {exp.ROWS_IN_DOCUMENT} of {total} rows. The Excel download holds every row." in text
    sheet = openpyxl.load_workbook(io.BytesIO(http("get", "/sentinel/api/reports/ropa?format=xlsx", name="super").content))["Controller register"]
    assert sheet.max_row == total + 1, "the workbook holds every row"
    assert not any("Excel download" in str(c.value) for row in sheet.iter_rows() for c in row if c.value)


# ── The evidence pack and the audit trail ───────────────────────────────────

def _pack(answer):
    assert answer.status_code == 200
    with zipfile.ZipFile(io.BytesIO(answer.content)) as zf:
        return {n: zf.read(n) for n in zf.namelist()}


def _stems(pack):
    return {re.sub(r"_\d{8}\.json$", "", n) for n in pack if n.endswith(".json")}


FULL = {"RoPA_Register", "DPIA_Summary", "Breach_Log", "DSR_Log", "Vendor_Register", "Audit_Trail", "Policy_Register", "Compliance_Score"}


def test_the_evidence_pack_holds_every_part_for_a_data_protection_officer_with_a_readme_that_says_who_and_when(http):
    _file(http.world)
    pack = _pack(http("get", "/sentinel/api/audit-export", name="dpo_a"))
    assert _stems(pack) == FULL, "the pack promised the RoPA and never held it"
    readme = pack["README.txt"].decode()
    assert "Module: Privacy" in readme and "Generated by: dpo_a" in readme and re.search(r"Generated: \d+ \w+ \d{4}, \d\d:\d\d UTC", readme)
    ropa = json.loads(next(v for n, v in pack.items() if n.startswith("RoPA_Register")))
    assert len(ropa) == 3, "the caller's records: A1, A2 and the organization wide one"
    assert "RoPA_Register" in readme and "3 records" in readme


def test_the_evidence_pack_leaves_out_what_the_role_may_not_open(http):
    """The pack used to hand the breach log, the request log (names and addresses), the impact assessments and the vendors to
    anyone who could open the module. Each part now needs the permission its own list needs."""
    world = http.world
    _file(world)
    _persona(world, "analyst", rbac.PRIVACY_ANALYST, "A")
    _persona(world, "manager", rbac.COMPLIANCE_MGR)
    analyst = _pack(http("get", "/sentinel/api/audit-export", name="analyst"))
    assert _stems(analyst) == FULL - {"Breach_Log", "Vendor_Register"}
    manager = _pack(http("get", "/sentinel/api/audit-export", name="manager"))
    assert _stems(manager) == {"RoPA_Register", "Audit_Trail", "Policy_Register", "Compliance_Score"}
    readme = manager["README.txt"].decode()
    assert "Left out" in readme and "sentinel.breach.manage" in readme and "sentinel.dsr.manage" in readme
    assert http("get", "/sentinel/api/audit-export?type=breach", name="analyst").status_code == 403
    assert http("get", "/sentinel/api/audit-export?type=breach", name="dpo_a").status_code == 200


def test_the_evidence_pack_is_complete_past_the_five_hundred_rows_the_screens_show(http):
    world = http.world
    for n in range(520):
        world.db.execute("INSERT INTO sentinel_ropa (ref_number, processing_name) VALUES (%s, %s)", (f"BULK-{n}", f"Bulk {n}"))
    world.db.commit()
    pack = _pack(http("get", "/sentinel/api/audit-export", name="super"))
    assert len(json.loads(next(v for n, v in pack.items() if n.startswith("RoPA_Register")))) == 520
    assert "520 records" in pack["README.txt"].decode()


def test_the_pack_file_name_is_safe_whatever_the_organization_is_called(http):
    ds.set_setting("org_name", 'Acme "Quote"\r\nInc/../Zimb\u00e1bwe')
    answer = http("get", "/sentinel/api/audit-export", name="super")
    assert answer.status_code == 200
    assert re.fullmatch(r'attachment; filename="[A-Za-z0-9._-]+\.zip"', answer.headers["content-disposition"])
    assert answer.headers["x-content-type-options"] == "nosniff"


def test_the_audit_trail_downloads_as_a_formula_safe_csv_of_the_callers_organization(http):
    world = http.world
    db = world.db
    db.execute("INSERT INTO organizations (id, name, slug) VALUES (2, 'other', 'other')")
    for org, action, details in ((1, "POST ropa", '=HYPERLINK("http://example.test","x")'), (1, "PUT breaches", "plain text"),
                                 (2, "POST ropa", "another organization")):
        db.execute("INSERT INTO audit_log (username, module, action, entity_type, details, org_id) VALUES (%s,'sentinel',%s,'ropa',%s,%s)",
                   ("someone", action, details, org))
    db.commit()
    answer = http("get", "/sentinel/api/audit/export", name="dpo_a")
    assert answer.status_code == 200 and answer.headers["content-type"].startswith("text/csv")
    assert re.fullmatch(r'attachment; filename="ThemisIQ_privacy_audit_trail_\d{8}_\d{4}\.csv"', answer.headers["content-disposition"])
    rows = list(csv.reader(io.StringIO(answer.text)))
    assert rows[0] == ["created_at_utc", "username", "action", "entity_type", "entity_id", "details", "ip_address"]
    details = {r[5] for r in rows[1:]}
    assert "'=HYPERLINK(\"http://example.test\",\"x\")" in details and "plain text" in details
    assert "another organization" not in details, "another organization's trail is not exported"
    everyone = {r[5] for r in list(csv.reader(io.StringIO(http("get", "/sentinel/api/audit/export", name="super").text)))[1:]}
    assert "another organization" in everyone, "the positive twin: a super administrator reads every organization"


def test_an_export_is_itself_recorded_in_the_audit_trail(http):
    http.world.db.execute("DELETE FROM audit_log")
    http.world.db.commit()
    assert http("get", "/sentinel/api/reports/ropa", name="dpo_a").status_code == 200
    assert http("get", "/sentinel/api/audit-export", name="dpo_a").status_code == 200
    assert http("get", "/sentinel/api/audit/export", name="dpo_a").status_code == 200
    actions = [r["action"] for r in http.world.db.execute("SELECT action FROM audit_log WHERE module = 'sentinel'").fetchall()]
    assert sorted(actions) == ["audit_trail_export", "evidence_pack_export", "report_export"]


def test_the_widest_reference_a_record_can_have_stays_on_one_line_in_the_register_tables(http):
    """A reference is made as PREFIX-YYYYMMDD-XXXXX from capitals and digits, so the widest one has W for every letter (19 characters for a
    RoPA entry). The Ref column of the two register tables held 12 percent of the page and broke such a reference in the middle of a number."""
    widest = {"ropa": "ROPA-20261022-WWWWW", "breach": "BRE-20261022-WWWWW"}
    db = http.world.db
    _insert(db, "sentinel_ropa", ref_number=widest["ropa"], processing_name="Widest reference", legal_basis="Contract", purpose="To test",
            data_categories='["Name"]', department="Human Resources", business_unit_id=None)
    _insert(db, "sentinel_breaches", ref_number=widest["breach"], title="Widest reference", status="open", severity="critical",
            discovery_date="2026-10-01", business_unit_id=None)
    for kind in ("ropa", "breach"):
        answer = http("get", f"/sentinel/api/reports/{kind}?format=pdf", name="super")
        assert answer.status_code == 200
        text = " ".join(" ".join(_pages(answer.content)).split())
        assert widest[kind] in text, f"the {kind} reference was split across lines: {widest[kind]}"
