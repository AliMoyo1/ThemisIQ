"""
The shared report kit (core/report_kit): the helpers every report uses and the four renderers that draw it.

No database and no request: a report is data, so a small hand built `Report` is enough to check what each renderer must do.
The rules under test are the ones the owner asked for: the module, who and when in the footer of every page, a cover with the
ThemisIQ logo, a visible tier, charts that are built from the same numbers as the tables, a figure that is missing shown as
n/a and never as zero, and a download that cannot carry a formula, a path or markup out of a name.
"""
import io
import re
import types
import zipfile
from datetime import datetime

import openpyxl
import pytest
from docx import Document
from lxml import etree
from pypdf import PdfReader

from core.report_kit import charts, grounding, helpers, identity, respond, theme
from core.report_kit.model import (Actor, Callout, Chart, Col, Heading, Kpi, Kpis, Meta, PageBreak, Paragraphs, Report, Table)
from core.report_kit.render_docx import render_docx
from core.report_kit.render_html import css_string, render_html
from core.report_kit.render_pdf import render_pdf
from core.report_kit.render_xlsx import render_xlsx

WHEN = datetime(2026, 10, 21, 16, 5)
FORMULA = '=HYPERLINK("http://example.test","click")'


# ── Identity: time, numbers, names, headers, cells ──────────────────────────

def test_a_stored_utc_time_is_written_in_utc_and_in_cat():
    assert identity.format_when("2026-10-21 16:05:00") == "21 Oct 2026, 16:05 UTC (18:05 CAT)"
    assert identity.format_when(datetime(2026, 10, 21, 23, 30)) == "21 Oct 2026, 23:30 UTC (22 Oct 01:30 CAT)"
    assert identity.format_when(None) == "Not recorded" and identity.format_day("2026-03-04") == "4 Mar 2026"


def test_a_percentage_rounds_half_up_and_a_missing_one_is_none():
    assert [identity.pct(5, 8), identity.pct(1, 8), identity.pct(1, 3), identity.pct(2, 3), identity.pct(0, 5)] == [63, 13, 33, 67, 0]
    assert identity.pct(0, 0) is None and identity.pct(None, 5) is None
    assert identity.pct(1, 8, digits=1) == 12.5
    assert round(5 / 8 * 100) == 62, "the reason the kit does not call round()"


def test_a_missing_figure_is_n_a_never_a_zero():
    assert identity.fmt_int(None) == identity.fmt_int("") == identity.fmt_pct(None) == identity.fmt_number(None) == "n/a"
    assert identity.fmt_int(0) == "0" and identity.fmt_int(1234567) == "1,234,567" and identity.fmt_int(2.5) == "3"
    assert identity.plural(1, "risk") == "1 risk" and identity.plural(3, "risk") == "3 risks" and identity.plural(2, "analysis", "analyses") == "2 analyses"


def test_module_names_are_the_navigation_names():
    assert [identity.module_name(k) for k in ("sentinel", "grid", "aria", "bcm", "erm", "orm", "platform")] == [
        "Privacy", "Audit", "Governance", "Resilience", "Enterprise Risk", "Operations Risk", "Command Centre"]
    assert identity.module_name("platform", "Reports") == "Command Centre, Reports" and identity.module_name("odd") == "odd"


def test_a_file_name_is_ascii_letters_digits_dot_dash_underscore_and_cannot_name_a_path():
    name = identity.safe_filename('Acme "Quote"\r\nInc/../Zimbábwe', "PDF", when=WHEN)
    assert re.fullmatch(r"[A-Za-z0-9._-]+", name) and "/" not in name and name.endswith("_20261021_1605.pdf")
    assert identity.safe_filename("", "csv") == "report.csv" and identity.safe_filename("a" * 200, "csv").startswith("a" * 80 + ".")


def test_download_headers_clean_a_name_that_was_built_by_hand():
    headers = identity.download_headers('evil"\r\nX-Injected: 1.pdf')
    assert re.fullmatch(r'attachment; filename="[A-Za-z0-9._-]+\.pdf"', headers["Content-Disposition"])
    assert "\r" not in headers["Content-Disposition"] and "\n" not in headers["Content-Disposition"]
    assert headers["X-Content-Type-Options"] == "nosniff" and headers["Cache-Control"] == "no-store"
    assert identity.download_headers("a.csv", private=False)["Cache-Control"] == "no-cache"


def test_a_cell_that_a_spreadsheet_would_read_as_a_formula_is_written_as_text():
    assert identity.csv_cell(FORMULA) == "'" + FORMULA
    assert [identity.csv_cell(v) for v in ("+1+1", "-cmd", "@SUM(A1)", "\tx", "\nx")] == ["'+1+1", "'-cmd", "'@SUM(A1)", "'\tx", "'\nx"]
    assert [identity.csv_cell(v) for v in ("-5", "+3.5", "-12,5%", "plain", "", 7, None)] == ["-5", "+3.5", "-12,5%", "plain", "", 7, None]
    assert identity.csv_row(["a", "=b", 3]) == ["a", "'=b", 3]


def test_text_for_a_document_loses_markup_and_control_characters_and_folds_to_what_the_fonts_draw():
    assert identity.clean_text("<b>Hi</b>\x00 there \n") == "Hi there" and identity.clean_text(None) == ""
    assert identity.clean_text("a\n b\t c", keep_newlines=False) == "a b c"
    assert identity.pdf_text("a‑b ✓ ā") == "a-b yes a"
    assert identity.markup_safe("a & b < c") == "a &amp; b &lt; c"
    assert identity.ellipsize("one two three four five", 12) == "one two…" and identity.ellipsize("short", 12) == "short"


def test_a_cell_set_as_text_stays_text_in_a_workbook():
    book = openpyxl.Workbook()
    cell = identity.set_cell(book.active["A1"], FORMULA)
    assert cell.data_type == "s" and cell.value == FORMULA


# ── Grounding ───────────────────────────────────────────────────────────────

def test_a_narrative_is_grounded_only_when_every_number_is_a_figure_of_the_data():
    data = {"total": 16, "done": 12, "name": "ISO 27001 controls"}
    assert grounding.check("12 of 16 controls (75%) are done, and 4 are not.", data).ok
    assert grounding.check("ISO 27001 is covered; 25% remain.", data).ok
    bad = grounding.check("There are 20 controls, 12 of them done.", data)
    assert not bad.ok and bad.unmatched == ["20"] and "20" in bad.sentence()
    assert grounding.check("Generated in 2026.", data).ok, "a year is not a figure"
    assert grounding.check("42 items.", data, extra=["42"]).ok
    assert grounding.tokens_in("1,250 and 3.50 and 7") == ["1250", "3.5", "7"]


# ── Helpers ─────────────────────────────────────────────────────────────────

def test_chart_rows_are_built_from_counts_in_one_way():
    assert helpers.by_count({"A": 1, None: 3, "B": 2}) == [("Not recorded", 3), ("B", 2), ("A", 1)]
    assert helpers.by_count({k: n for n, k in enumerate("abcdefgh", 1)}, limit=3, other="Rest") == [("h", 8), ("g", 7), ("Rest", 21)]
    rows, colours = helpers.severity_rows({"low": 4, "Critical": 1, None: 2, "high": 2, "odd": 1})
    assert rows == [("Critical", 1), ("High", 2), ("Low", 4), ("Not rated", 2), ("Odd", 1)]
    assert colours[:3] == [theme.GRAPHITE, theme.SLATE, theme.SILVER] and len(colours) == 5
    assert helpers.slices([("a", 5), ("b", 4), ("c", 3), ("d", 2), ("e", 1), ("f", 1)], limit=3) == [("a", 5), ("b", 4), ("Other states", 7)]
    assert helpers.slices([("a", 0), ("b", 2)]) == [("b", 2)]
    assert helpers.label("under_review") == "Under review" and helpers.label(None) == "No status"
    assert helpers.empty("No risk is recorded.") == "No risk is recorded. There is nothing to chart."


def test_a_tie_for_the_largest_group_is_never_reported_as_a_leader():
    assert helpers.leader([("A", 2), ("B", 2), ("C", 1)], 5, "type", "event") == "2 types are tied for the largest share: 2 events each."
    assert helpers.leader([("A", 3), ("B", 1)], 4, "category", "open risk", group_plural="categories") == \
        "A is the largest category: 3 of 4 open risks (75%)."
    assert helpers.leader([], 0, "type", "event") == ""
    assert helpers.tidy_number(60.0) == "60" and helpers.tidy_number(11.25) == "11.3" and helpers.tidy_number(0.04) == "0"


def test_a_metric_becomes_a_headline_measure_without_ever_becoming_a_zero():
    def metric(state, value=None, note=None, **extra):
        return {"label": "Plans", "state": state, "value": value, "note": note, **extra}
    assert helpers.metric_kpi(None) is None, "a metric the viewer may not open is left out"
    measured = helpers.metric_kpi(metric("measured", 3), note="2 approved", tone="bad")
    assert (measured.value, measured.note, measured.tone) == ("3", "2 approved", "bad")
    assert helpers.metric_kpi(metric("measured", 0), tone="bad").tone == "neutral", "a measured zero is a number, not an alarm"
    nothing = helpers.metric_kpi(metric("not_measured", note="No exercise is recorded yet."))
    assert (nothing.value, nothing.note) == ("Not measured", "No exercise is recorded yet.")
    assert helpers.metric_kpi(metric("unavailable", note="could not be read")).value == "Unavailable"
    assert helpers.metric_kpi(metric("measured", 63), text="63%").value == "63%"
    assert helpers.metric_kpi(metric("stale", 5, as_of="2026-10-01 00:00:00")).note.startswith("As of 1 Oct 2026")
    assert helpers.metric_kpi(metric("partial", 5)).note.startswith("Some sources could not be read")


def test_the_format_a_request_asks_for_is_pdf_or_xlsx_and_nothing_else():
    ask = lambda q: respond.format_of(types.SimpleNamespace(query_params=q))      # noqa: E731
    assert ask({}) == "pdf" and ask({"format": "XLSX"}) == "xlsx" and ask({"format": "docx"}) is None and ask({"format": ""}) == "pdf"
    assert respond.client_ip(types.SimpleNamespace(client=None)) == "" and respond.client_ip(types.SimpleNamespace(client=types.SimpleNamespace(host="1.2.3.4"))) == "1.2.3.4"


# ── The model ───────────────────────────────────────────────────────────────

def test_a_report_has_a_tier_a_module_and_a_footer_that_says_who_and_when():
    meta = Meta(title="T", module="sentinel", generated_at=WHEN, generated_by=Actor(name="Jane Doe", username="jdoe"))
    assert meta.tier == theme.OPERATIONAL and meta.tier_label == "Operational report"
    assert meta.footer_line() == ("Module: Privacy  |  Generated by Jane Doe (jdoe)  |  Generated 21 Oct 2026, 16:05 UTC (18:05 CAT)  |  Operational report")
    with pytest.raises(ValueError):
        Meta(title="T", module="sentinel", tier="executive")
    assert Actor(name="jdoe", username="jdoe").describe() == "jdoe"
    assert Actor(name="Run", kind="scheduled", owner_name="Jane Doe", owner_username="jdoe").describe() == "Scheduled run for Jane Doe (jdoe)"
    assert Actor(name="x", kind="system", owner_name="Jane Doe").describe() == "System run for Jane Doe"
    assert Actor.of_user({"full_name": "Jane Doe", "username": "jdoe"}).describe() == "Jane Doe (jdoe)" and Actor.of_user(None).name == "Unknown user"


def test_a_table_cut_for_the_document_says_so_and_says_where_every_row_is_only_when_an_excel_file_exists():
    rows = [{"n": i} for i in range(7)]
    cut = Table("T", [Col("N", "n")], rows, max_rows=5)
    assert len(cut.shown()) == 5 and cut.truncation_note() == "Showing the first 5 of 7 rows."
    cut.more_in_excel = True
    assert cut.truncation_note() == "Showing the first 5 of 7 rows. The Excel download holds every row."
    assert Table("T", [Col("N", "n")], rows[:3], max_rows=5, more_in_excel=True).truncation_note() == ""


def test_a_chart_with_nothing_to_draw_is_empty_and_a_missing_value_is_not_a_zero():
    assert Chart("bar", "T", []).is_empty() and Chart("bar", "T", [("a", 0)]).is_empty() and not Chart("bar", "T", [("a", 1)]).is_empty()
    assert Chart("stack", "T", [("a", [("x", 0), ("y", 0)])]).is_empty()
    drawing = charts.build_drawing(Chart("bar", "T", [("known", 4), ("missing", None)]), 480)
    assert "n/a" in [p["s"] for p in charts.primitives(drawing) if p["t"] == "text"]


# ── Charts ──────────────────────────────────────────────────────────────────

CHARTS = [Chart("bar", "Bars", [("Critical", 3), ("High", 5), ("Medium", 8)], x_label="Number of risks", y_label="Level", unit="risks",
                takeaway="Medium is the largest."),
          Chart("stack", "Stack", [("A", [("Done", 4), ("To do", 6)]), ("B", [("Done", 9), ("To do", 1)])], x_label="Controls", y_label="Framework"),
          Chart("donut", "Donut", [("Open", 3), ("Closed", 9)], centre="75%", unit="items"),
          Chart("line", "Line", [("Jan", 1), ("Feb", 3), ("Mar", 2)], x_label="Month", y_label="Events", unit="events")]


@pytest.mark.parametrize("chart", CHARTS, ids=[c.kind for c in CHARTS])
def test_every_chart_is_one_scene_drawn_as_vector_svg_and_png_with_its_numbers_on_the_marks(chart):
    drawing = charts.build_drawing(chart, 480)
    assert drawing.height > 40 and drawing.width == 480
    texts = [p["s"] for p in charts.primitives(drawing) if p["t"] == "text"]
    figures = [str(v) for _, v in chart.data if not isinstance(v, list)]
    assert all(any(f in t for t in texts) for f in figures) or chart.kind == "stack", "a value is written on its mark"
    svg = charts.drawing_to_svg(drawing, label=charts.alt_text(chart))
    assert svg.startswith("<svg") and "</svg>" in svg and "<script" not in svg and 'role="img"' in svg and "rgb(31,35,40)" in svg
    assert charts.drawing_to_png(drawing).startswith(b"\x89PNG")
    assert charts.alt_text(chart).startswith(chart.title) and str(chart.data[0][0]) in charts.alt_text(chart)


def test_no_chart_uses_a_blue_and_every_series_runs_from_dark_to_light():
    for chart in CHARTS:
        svg = charts.drawing_to_svg(charts.build_drawing(chart, 480))
        colours = re.findall(r"rgba?\((\d+),(\d+),(\d+)", svg)
        assert colours, chart.kind
        for r, g, b in ((int(x), int(y), int(z)) for x, y, z in colours):
            assert not (b > r + 24 and b > g + 8), f"rgb({r},{g},{b}) in a {chart.kind} chart reads as blue"
    assert theme.SERIES == sorted(theme.SERIES, key=lambda c: sum(theme.rgb(c)))


def test_axis_labels_are_spaced_so_they_never_print_over_each_other():
    assert charts.nice_ticks(310000, count=4, whole=True)[-1] == 400000
    fitted = charts._fit_ticks(310000, True, "USD", plot_w=100)
    assert len(fitted) <= 3 and fitted[-1] >= 310000, "fewer ticks when the labels are wide"
    assert len(charts._fit_ticks(8, True, "", plot_w=300)) >= 4, "and the usual number when there is room"
    bar = Chart("bar", "Wide numbers", [("a", 310000), ("b", 182000)], unit="USD", x_label="USD")
    ticks = sorted((p for p in charts.primitives(charts.build_drawing(bar, 234))
                    if p["t"] == "text" and p["y"] < 20 and re.fullmatch(r"[\d,]+", p["s"])), key=lambda p: p["x"])
    assert len(ticks) >= 2
    for a, b in zip(ticks, ticks[1:]):
        gap = b["x"] - a["x"] - (charts.text_width(a["s"], a["size"]) + charts.text_width(b["s"], b["size"])) / 2
        assert gap > 0, f"{a['s']} and {b['s']} touch"


def test_the_first_point_of_a_line_sits_clear_of_the_axis_labels():
    drawing = charts.build_drawing(Chart("line", "L", [("Jan", 14.3), ("Feb", 10.8)], unit=""), 234)
    circles = [p for p in charts.primitives(drawing) if p["t"] == "circle"]
    assert len(circles) == 2 and min(c["cx"] for c in circles) >= 38 + 9, "the plot starts at 38 points; the first point is inset from it"


# ── The four renderers ──────────────────────────────────────────────────────

def _report(tier=theme.OPERATIONAL, module="sentinel") -> Report:
    meta = Meta(title="Quarterly privacy report", module=module, tier=tier, org_name="Econet", scope="Whole organization",
                period="As of 21 Oct 2026", subtitle="A short subtitle", generated_at=WHEN, generated_by=Actor(name="Jane Doe", username="jdoe"),
                reference="Run 7")
    rows = [{"name": f"Item {i}", "status": "Open" if i % 2 else "Closed", "score": i * 3, "when": "2026-10-0%d" % (i % 9 + 1)} for i in range(70)]
    rows[0]["name"] = FORMULA
    report = Report(meta)
    report.add(Heading("Headline measures", "What the numbers say."),
               Kpis([Kpi("Items", "70"), Kpi("Open", "35", "half of them", "bad"), Kpi("Not measured", "n/a", "nothing recorded")]),
               Heading("At a glance"), *CHARTS,
               Callout("Decisions for the board", ["Decide one thing.", "Decide another."], kind="decision"),
               Callout("Limits", ["The figures are as recorded."], kind="limit"),
               Paragraphs("Draft", ["A written paragraph."], label="Draft written by AI, not reviewed"),
               PageBreak(),
               Table("Items", [Col("Name", "name", weight=3), Col("Status", "status"), Col("Score", "score", "right", 1, "int"), Col("When", "when", kind="date")],
                     rows, max_rows=40, more_in_excel=True, sheet_name="Items",
                     sheet_columns=[Col("Name", "name"), Col("Status", "status"), Col("Score", "score", "right", kind="int"), Col("When", "when", kind="date"),
                                    Col("Workbook only", "extra")]))
    return report


def test_the_pdf_has_a_cover_with_the_logo_and_the_tier_and_the_same_footer_on_every_page():
    for tier, label in ((theme.OPERATIONAL, "Operational report"), (theme.BOARD, "Board report")):
        body = render_pdf(_report(tier))
        reader = PdfReader(io.BytesIO(body))
        pages = [p.extract_text() for p in reader.pages]
        assert len(pages) >= 5 and len(reader.pages[0].images) >= 1
        assert "Quarterly privacy report" in pages[0] and label in pages[0] and "Whole organization" in pages[0] and "Run 7" in pages[0]
        for number, text in enumerate(pages, start=1):
            assert "Module: Privacy" in text and "Generated by Jane Doe (jdoe)" in text and "21 Oct 2026, 16:05 UTC (18:05 CAT)" in text, number
            assert f"Page {number} of {len(pages)}" in text and label in text, number
        text = " ".join(pages).replace("\n", " ")
        assert "Showing the first 40 of 70 rows. The Excel download holds every row." in text
        assert "Draft written by AI, not reviewed" in text and "Decisions for the board" in text and "n/a" in text
        assert reader.metadata.title == "Quarterly privacy report" and reader.metadata.author == "Jane Doe (jdoe)"


@pytest.mark.parametrize("filler", range(4, 64, 3))
def test_a_heading_is_never_left_alone_at_the_foot_of_a_page_with_its_chart_on_the_next(filler):
    """reportlab cannot tie a heading to a chart that is kept whole, so a heading used to be stranded under the last line of a
    page whenever the chart after it did not fit. The heading asks for the room of the chart itself."""
    report = Report(Meta(title="T", module="erm", generated_at=WHEN, generated_by=Actor(name="Jane Doe", username="jdoe")))
    rows = [(f"Row {n}", n + 1) for n in range(12)]            # tall enough that a chart often does not fit under the heading
    pair = [Chart("bar", "Left chart", rows, half=True, unit="x", takeaway="Left takeaway."),
            Chart("bar", "Right chart", rows, half=True, unit="x", takeaway="Right takeaway.")]
    report.add(Heading("First"), Paragraphs("Filler", ["A line of filler text that takes a line of the page."] * filler),
               Heading("Strand check"), *pair)
    pages = [p.extract_text() for p in PdfReader(io.BytesIO(render_pdf(report))).pages]
    holders = [text for text in pages if "Strand check" in text]
    assert len(holders) == 1 and "Left chart" in holders[0] and "Right chart" in holders[0], f"filler {filler}: heading and charts apart"


def test_the_two_tiers_look_different_on_the_cover():
    covers = [PdfReader(io.BytesIO(render_pdf(_report(t)))).pages[0] for t in (theme.OPERATIONAL, theme.BOARD)]
    assert covers[0].get_contents().get_data() != covers[1].get_contents().get_data()


def test_a_name_is_drawn_as_text_and_markup_in_it_is_not_markup():
    report = _report()
    report.meta.title = "A <b>bold</b> & “quoted” title"
    reader = PdfReader(io.BytesIO(render_pdf(report)))
    assert "bold" in reader.pages[0].extract_text() and "<b>" not in reader.pages[0].extract_text()


EXPECTED_ORDER = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku "
           "wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents "
           "suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange".split(),
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color "
           "spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath".split(),
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark".split(),
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd "
             "tblLayout tblCellMar tblLook tblCaption tblDescription".split(),
    "sectPr": "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType pgNumType cols formProt "
              "vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings".split(),
}


def test_the_docx_orders_every_property_element_as_word_requires():
    """Word refuses a file whose property children are out of schema order; this checks every part of the package."""
    with zipfile.ZipFile(io.BytesIO(render_docx(_report()))) as package:
        parts = [n for n in package.namelist() if re.fullmatch(r"word/(document|header\d*|footer\d*)\.xml", n)]
        assert any("footer" in n for n in parts) and any("header" in n for n in parts)
        for name in parts:
            root = etree.fromstring(package.read(name))
            for element in root.iter():
                local = etree.QName(element).localname
                if local in EXPECTED_ORDER and isinstance(element.tag, str):
                    order = EXPECTED_ORDER[local]
                    seen = [etree.QName(child).localname for child in element]
                    indexes = [order.index(s) for s in seen]
                    assert indexes == sorted(indexes), f"{name}: <{local}> children out of order: {seen}"


def test_the_docx_has_a_cover_a_footer_in_every_section_with_page_fields_and_the_charts_as_pictures():
    body = render_docx(_report(theme.BOARD))
    doc = Document(io.BytesIO(body))
    assert len(doc.inline_shapes) >= len(CHARTS) + 1, "the charts and the logo are pictures"
    assert doc.sections, "at least one section"
    for section in doc.sections:
        footer = " ".join(p.text for p in section.footer.paragraphs)
        assert "Module: Privacy" in footer and "Generated by Jane Doe (jdoe)" in footer and "Board report" in footer
    headers = [" ".join(p.text for p in s.header.paragraphs) for s in doc.sections]
    assert any("ThemisIQ" in h and "Quarterly privacy report" in h for h in headers), "the body pages carry a running header"
    with zipfile.ZipFile(io.BytesIO(body)) as package:
        footers = " ".join(package.read(n).decode() for n in package.namelist() if re.fullmatch(r"word/footer\d*\.xml", n))
    assert "PAGE" in footers and "NUMPAGES" in footers
    text = "\n".join(p.text for p in doc.paragraphs)
    cover = " ".join(c.text for t in doc.tables[:2] for r in t.rows for c in r.cells)
    assert "BOARD REPORT" in text, "the tier chip is on the cover"
    assert "Quarterly privacy report" in cover and "Jane Doe (jdoe)" in cover and "Board report" in cover, "title, who and the tier on the cover"


def test_the_workbook_has_a_cover_a_summary_with_native_charts_print_setup_and_no_formula():
    book = openpyxl.load_workbook(io.BytesIO(render_xlsx(_report())))
    assert book.sheetnames[:2] == ["Cover", "Summary"] and "Items" in book.sheetnames
    assert len(book["Cover"]._images) == 1, "the ThemisIQ logo"
    assert len(book["Summary"]._charts) >= 3, "native Excel charts"
    items = book["Items"]
    assert [c.value for c in items[1]] == ["Name", "Status", "Score", "When", "Workbook only"], "the workbook's own, fuller columns"
    assert items.max_row == 71, "every row, though the document is cut at 40"
    assert items.freeze_panes == "A2" and items.auto_filter.ref and items.page_setup.orientation == "landscape"
    assert "Module: Privacy" in (items.oddFooter.left.text or "") and "Generated by Jane Doe (jdoe)" in (items.oddFooter.left.text or "")
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f", f"{sheet.title}!{cell.coordinate} is a formula"
    typed = items["A2"]
    assert typed.data_type == "s" and typed.value == FORMULA
    assert book.properties.title == "Quarterly privacy report" and "Jane Doe" in book.properties.creator


def test_the_print_view_has_print_css_the_footer_and_escapes_every_dynamic_value():
    report = _report(theme.BOARD)
    report.meta.title = '<script>alert(1)</script> "x"'
    report.meta.generated_by = Actor(name='Jane "</style><b>', username="jdoe")
    html = render_html(report, logo_url="/static/img/themisiq-logo-clear.png")
    assert "@page" in html and "tier-board" in html and "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html and "</style><b>" not in html
    assert "counter(page)" in html and "Module: Privacy" in html
    assert 'src="/static/img/themisiq-logo-clear.png"' in html
    assert css_string('x"; } </style><script>') == "x\\22 ; \\7d ", "a tag is dropped, a quote and a brace are escaped"
    assert css_string("a<b\\c") == "a\\3c b\\5c c" and css_string("plain, safe text (1)") == "plain, safe text (1)"


def test_a_report_with_nothing_in_it_still_renders_in_every_format_and_says_so():
    report = Report(Meta(title="Empty", module="erm", generated_at=WHEN, generated_by=Actor(name="Jane Doe", username="jdoe")))
    report.add(Heading("Nothing"), Kpis([Kpi("Risks", "n/a")]), Chart("bar", "Empty chart", [], empty_text="No risk is recorded. There is nothing to chart."),
               Table("Risks", [Col("Name", "name")], [], empty_text="No risk is recorded."))
    pdf = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(render_pdf(report))).pages).replace("\n", " ")
    assert "There is nothing to chart" in pdf and "No risk is recorded." in pdf and "Module: Enterprise Risk" in pdf
    assert openpyxl.load_workbook(io.BytesIO(render_xlsx(report)))["Risks"]["A2"].value == "No risk is recorded."
    assert "No risk is recorded." in " ".join(p.text for p in Document(io.BytesIO(render_docx(report))).paragraphs) or \
        "No risk is recorded." in "".join(c.text for t in Document(io.BytesIO(render_docx(report))).tables for r in t.rows for c in r.cells)
    assert "There is nothing to chart" in render_html(report)
