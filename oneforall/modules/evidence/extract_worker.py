"""Bounded document-to-text helper. Invoked only as a child process by indexer.py.

The input is untrusted. This process never opens the application database or
returns raw parser errors. It is resource limited on Unix; the caller imposes a
wall-time limit on every platform. Image OCR is intentionally unsupported.
"""
from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

MAX_SOURCE_BYTES = 12 * 1024 * 1024
MAX_EXPANDED_BYTES = 48 * 1024 * 1024
MAX_ZIP_ENTRIES = 2000
MAX_SEGMENTS = 100
MAX_TEXT_CHARS = 200_000
MAX_SHEET_CELLS = 50_000
SUPPORTED = frozenset({".pdf", ".docx", ".xlsx", ".txt", ".csv", ".json"})


class Unsupported(Exception):
    def __init__(self, code: str):
        self.code = code


def _limits():
    if sys.platform == "win32":
        return
    import resource
    resource.setrlimit(resource.RLIMIT_CPU, (15, 15))
    resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_FSIZE, (1_000_000, 1_000_000))
    resource.setrlimit(resource.RLIMIT_CORE, (0, 0))


def _check_zip(path: Path):
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > MAX_ZIP_ENTRIES or sum(i.file_size for i in members) > MAX_EXPANDED_BYTES:
            raise Unsupported("archive_too_large")
        if any(i.file_size > MAX_EXPANDED_BYTES for i in members):
            raise Unsupported("archive_too_large")


def _clean(text) -> str:
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", str(text or ""))
    return re.sub(r"\s+", " ", value).strip()


def _parts(path: Path, extension: str):
    if extension in (".txt", ".csv", ".json"):
        try:
            yield "Document", path.read_text(encoding="utf-8-sig", errors="strict")
        except UnicodeError as exc:
            raise Unsupported("encoding_not_supported") from exc
    elif extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path), strict=True)
        if reader.is_encrypted:
            raise Unsupported("encrypted_file")
        if len(reader.pages) > MAX_SEGMENTS:
            raise Unsupported("too_many_pages")
        for number, page in enumerate(reader.pages, 1):
            yield f"Page {number}", page.extract_text() or ""
    elif extension == ".docx":
        _check_zip(path)
        from docx import Document
        document = Document(str(path))
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        yield "Document", "\n".join(paragraphs)
    elif extension == ".xlsx":
        _check_zip(path)
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if len(book.worksheets) > MAX_SEGMENTS:
                raise Unsupported("too_many_sheets")
            count = 0
            for sheet in book.worksheets:
                lines = []
                for row in sheet.iter_rows(values_only=True):
                    count += len(row)
                    if count > MAX_SHEET_CELLS:
                        raise Unsupported("too_many_cells")
                    lines.append(" | ".join(str(v) for v in row if v is not None))
                yield f"Sheet {sheet.title[:60]}", "\n".join(lines)
        finally:
            book.close()


def extract(path: Path) -> dict:
    extension = path.suffix.lower()
    if extension not in SUPPORTED:
        return {"status": "unsupported", "error_code": "format_not_supported", "text": "", "segments": []}
    size = path.stat().st_size
    if size > MAX_SOURCE_BYTES:
        return {"status": "unsupported", "error_code": "file_too_large", "text": "", "segments": []}
    text = ""
    segments = []
    for label, part in _parts(path, extension):
        if len(segments) >= MAX_SEGMENTS:
            raise Unsupported("too_many_segments")
        part = _clean(part)
        if not part:
            continue
        if len(text) + len(part) + 1 > MAX_TEXT_CHARS:
            raise Unsupported("text_too_large")
        start = len(text) + (1 if text else 0)
        text = (text + "\n" if text else "") + part
        segments.append({"label": label, "start": start, "end": len(text)})
    return {"status": "indexed" if text else "empty", "error_code": None, "text": text, "segments": segments}


def main() -> int:
    _limits()
    try:
        if len(sys.argv) != 2:
            raise Unsupported("invalid_input")
        result = extract(Path(sys.argv[1]))
    except Unsupported as exc:
        result = {"status": "unsupported", "error_code": exc.code, "text": "", "segments": []}
    except FileNotFoundError:
        result = {"status": "failed", "error_code": "file_missing", "text": "", "segments": []}
    except Exception:
        result = {"status": "failed", "error_code": "parse_failed", "text": "", "segments": []}
    sys.stdout.write(json.dumps(result, ensure_ascii=False, separators=(",", ":")))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
