"""
PLAN-35 T09: publication job processing and retention tests.

Covers claim/lease semantics (including a real concurrency proof that two
workers never claim the same job), idempotent Evidence Vault and GRID
evidence copies on replay, the business-unit compatibility gate before
GRID auto-attachment, bounded-backoff retry then permanent failure, the
explicit scoped-manager retry action, draft expiry, and the storage
retention sweep (orphan-to-trash, trash purge, and that a referenced path
is never touched), and that the pre-existing ARIA_POLICY_PUBLISHED handler
does not run its legacy mutable-document Vault copy a second time for a
managed-version publication that policy_publication.py already copied.
"""
import concurrent.futures
import io

import pytest
from docx import Document as DocxDocument

import database
from modules.aria import policy_publication as pub
from modules.aria import policy_storage as storage
from modules.aria import policy_preview as preview
from modules.aria import policy_workflow_service as svc


def _org(db, org_id=1):
    db.execute("INSERT INTO organizations (id, name, slug) VALUES (%s,%s,%s)",
               (org_id, f"org{org_id}", f"org{org_id}"))


def _bu(db, bu_id, name, parent_id=None):
    db.execute("INSERT INTO business_units (id, name, parent_id, is_active) VALUES (%s,%s,%s,1)",
               (bu_id, name, parent_id))


def _user(db, uid, org_id=1, bu_id=None, username=None, is_active=1):
    username = username or f"user{uid}"
    db.execute(
        "INSERT INTO users (id, username, email, full_name, password_hash, org_id, "
        "business_unit_id, is_active) VALUES (%s,%s,%s,%s,'x',%s,%s,%s)",
        (uid, username, f"{username}@x.com", username, org_id, bu_id, is_active),
    )


def _role(db, uid, role_key):
    db.execute("INSERT INTO user_roles (user_id, role_key) VALUES (%s,%s)", (uid, role_key))


def _actor(db, uid):
    row = db.execute(
        "SELECT id, username, full_name, org_id, business_unit_id, "
        "COALESCE(is_super_admin,0) AS is_super_admin FROM users WHERE id=%s", (uid,),
    ).fetchone()
    d = dict(row)
    d["roles"] = [r[0] for r in db.execute(
        "SELECT role_key FROM user_roles WHERE user_id=%s", (uid,)
    ).fetchall()]
    return d


def _real_template_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("TEMPLATE COVER")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture(autouse=True)
def _isolated_storage(tmp_path, monkeypatch):
    root = tmp_path / "aria_uploads"
    monkeypatch.setattr(storage, "ARIA_UPLOAD_DIR", root)
    monkeypatch.setattr(storage, "WORKFLOW_ROOT", root / "policy_workflow")
    monkeypatch.setattr(preview, "submit_conversion_job", lambda b, timeout_seconds=None: "job")
    monkeypatch.setattr(preview, "poll_conversion_result",
                         lambda j, timeout_seconds=None, poll_interval=0.5: b"%PDF-X")
    monkeypatch.setattr(preview, "cleanup_job", lambda j: None)
    yield root


def _approve_a_policy(db, monkeypatch, *, org_id=1, bu_id=100, control_ref="A.1", framework="ISO 27001"):
    """Full pipeline to a real approved version + real pending publication
    job: draft -> build -> confirm -> submit -> decide. Returns
    (job_row_dict, version_row_dict, doc_row_dict, author_actor, approver_actor)."""
    _org(db, org_id)
    _bu(db, bu_id, "Finance")
    _user(db, 1, org_id=org_id, bu_id=bu_id, username="author")
    _user(db, 2, org_id=org_id, bu_id=bu_id, username="approver")
    _role(db, 1, "policy_author")
    _role(db, 2, "compliance_manager")

    template_dir = storage.ARIA_UPLOAD_DIR.parent / "aria_templates"
    template_dir.mkdir(parents=True, exist_ok=True)
    (template_dir / "t1.docx").write_bytes(_real_template_bytes())
    monkeypatch.setattr("modules.aria.routes.ARIA_TEMPLATE_DIR", template_dir)

    db.execute(
        "INSERT INTO aria_doc_templates (name, file_path, file_name, org_id, is_active) "
        "VALUES ('T1', 't1.docx', 't1.docx', %s, 1)", (org_id,)
    )
    db.commit()
    template_id = db.execute("SELECT id FROM aria_doc_templates WHERE file_path='t1.docx'").fetchone()["id"]

    author = _actor(db, 1)
    draft = svc.create_draft_from_generation(
        db, author,
        control={"id": 1, "ref": control_ref, "name": "Test", "framework_id": 1, "fw_name": framework},
        generated_content="# Policy", org_name="Econet", doc_type="Policy", framework_label=framework,
        integrated_controls=None, request_id=None, target_business_unit_id=bu_id,
    )
    built = svc.build_draft(db, author, draft["id"], template_id, draft["lock_version"])
    confirmed = svc.confirm_draft(db, author, built["id"], built["build_id"], built["lock_version"])
    approval = svc.submit_for_approval(db, author, confirmed["version_id"], 2, "note", "req-1", 1)

    approver = _actor(db, 2)
    svc.decide_approval(db, approver, approval["id"], "approve", "ok", approval["lock_version"])

    job = dict(db.execute(
        "SELECT * FROM aria_policy_publication_jobs WHERE policy_version_id=%s",
        (confirmed["version_id"],),
    ).fetchone())
    version = dict(db.execute(
        "SELECT * FROM aria_policy_versions WHERE id=%s", (confirmed["version_id"],)
    ).fetchone())
    doc = dict(db.execute(
        "SELECT * FROM aria_documents WHERE id=%s", (job["document_id"],)
    ).fetchone())
    return job, version, doc, author, approver


# ─────────────────────────────────────────────────────────────────────────
# Claim / lease
# ─────────────────────────────────────────────────────────────────────────

def test_claim_next_job_returns_none_when_nothing_due(test_db):
    assert pub.claim_next_job(test_db, is_postgres=False) is None


def test_claim_next_job_marks_running_with_a_lease(test_db, monkeypatch):
    job, *_ = _approve_a_policy(test_db, monkeypatch)
    claimed = pub.claim_next_job(test_db, is_postgres=False)
    assert claimed is not None
    assert claimed["id"] == job["id"]
    assert claimed["state"] == "running"
    assert claimed["lease_token"]
    assert claimed["attempts"] == 1


def test_claim_next_job_reclaims_an_expired_lease(test_db, monkeypatch):
    job, *_ = _approve_a_policy(test_db, monkeypatch)
    test_db.execute(
        "UPDATE aria_policy_publication_jobs SET state='running', "
        "lease_until='2000-01-01T00:00:00', lease_token='stale' WHERE id=%s",
        (job["id"],),
    )
    test_db.commit()
    claimed = pub.claim_next_job(test_db, is_postgres=False)
    assert claimed is not None
    assert claimed["id"] == job["id"]
    assert claimed["lease_token"] != "stale"


def test_two_concurrent_claimants_exactly_one_wins(tmp_path, monkeypatch):
    """Real file-based DB and real separate threads/connections -- a
    genuine race, not a simulated one, mirroring the same proof already
    used for decide_approval and reserve_document_number."""
    monkeypatch.setattr(database, "_DB_PATH", str(tmp_path / "concurrency_pub.db"))
    database.init_db()
    root = tmp_path / "aria_uploads"
    monkeypatch.setattr(storage, "ARIA_UPLOAD_DIR", root)
    monkeypatch.setattr(storage, "WORKFLOW_ROOT", root / "policy_workflow")
    monkeypatch.setattr(preview, "submit_conversion_job", lambda b, timeout_seconds=None: "job")
    monkeypatch.setattr(preview, "poll_conversion_result",
                         lambda j, timeout_seconds=None, poll_interval=0.5: b"%PDF-X")
    monkeypatch.setattr(preview, "cleanup_job", lambda j: None)

    setup_db = database.get_db()
    job, *_ = _approve_a_policy(setup_db, monkeypatch)
    setup_db.close()

    results = []

    def claim_worker():
        db = database.get_db()
        try:
            claimed = pub.claim_next_job(db, is_postgres=False)
            if claimed:
                results.append(claimed)
        finally:
            db.close()

    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(claim_worker) for _ in range(2)]
        for f in futures:
            f.result()

    assert len(results) == 1, f"expected exactly one claimant, got {len(results)}"
    assert results[0]["id"] == job["id"]


# ─────────────────────────────────────────────────────────────────────────
# Processing and idempotency
# ─────────────────────────────────────────────────────────────────────────

def test_process_job_copies_to_evidence_vault_and_completes(test_db, monkeypatch):
    job, version, doc, *_ = _approve_a_policy(test_db, monkeypatch)
    claimed = pub.claim_next_job(test_db, is_postgres=False)
    outcome = pub.process_job(test_db, claimed, is_postgres=False)
    assert outcome == "complete"

    row = test_db.execute(
        "SELECT state, event_id FROM aria_policy_publication_jobs WHERE id=%s", (job["id"],)
    ).fetchone()
    assert row["state"] == "complete"
    assert row["event_id"] is not None

    vault = test_db.execute(
        "SELECT * FROM evidence_items WHERE aria_policy_version_id=%s", (version["id"],)
    ).fetchone()
    assert vault is not None
    assert vault["file_hash"] == version["branded_sha256"]

    link = test_db.execute(
        "SELECT * FROM evidence_links WHERE evidence_id=%s AND module='aria' AND entity_type='document'",
        (vault["id"],),
    ).fetchone()
    assert link is not None
    assert link["entity_id"] == doc["id"]


def test_replaying_a_completed_job_does_not_duplicate_vault_evidence(test_db, monkeypatch):
    job, version, doc, *_ = _approve_a_policy(test_db, monkeypatch)
    claimed = pub.claim_next_job(test_db, is_postgres=False)
    pub.process_job(test_db, claimed, is_postgres=False)

    # Simulate a crash-and-replay: process the SAME job data again directly
    # (bypassing claim, as a retried worker with stale job data might).
    pub.process_job(test_db, claimed, is_postgres=False)

    count = test_db.execute(
        "SELECT COUNT(*) AS c FROM evidence_items WHERE aria_policy_version_id=%s", (version["id"],)
    ).fetchone()["c"]
    assert count == 1


def test_replaying_a_completed_job_does_not_duplicate_event_side_effects(test_db, monkeypatch):
    """The vault/GRID copies are separately idempotent (proved above), but
    emit()'s own side effects (task creation, cross-module links,
    notifications, webhooks) are not, on their own -- this is what
    dedup_key on emit() actually guards. Checking only evidence_items here
    would miss a real replay duplicating a review task or re-notifying a
    GRID policy requester a second time."""
    import core.event_handlers  # noqa: F401 -- @on registration is a side effect of import;
    # without this, whether policy_published_handler is even registered
    # depends on test execution order (whatever else already imported it).
    job, version, doc, *_ = _approve_a_policy(test_db, monkeypatch)
    claimed = pub.claim_next_job(test_db, is_postgres=False)
    pub.process_job(test_db, claimed, is_postgres=False)
    pub.process_job(test_db, claimed, is_postgres=False)

    events = test_db.execute(
        "SELECT COUNT(*) AS c FROM events WHERE dedup_key=%s", (job["publication_key"],)
    ).fetchone()["c"]
    assert events == 1, "a second emit() for the same publication_key must not create a second event row"

    tasks = test_db.execute(
        "SELECT COUNT(*) AS c FROM task_board WHERE module='aria' AND entity_type='document' AND entity_id=%s",
        (doc["id"],),
    ).fetchone()["c"]
    assert tasks == 1, "the ARIA_POLICY_PUBLISHED handler's review task must not be created twice on replay"


def test_grid_evidence_attached_when_control_matches(test_db, monkeypatch):
    job, version, doc, *_ = _approve_a_policy(test_db, monkeypatch, control_ref="A.1", framework="ISO 27001")
    fw_id = database.insert_returning_id(test_db,
        "INSERT INTO grid_frameworks (name) VALUES ('ISO 27001')", ())
    audit_id = database.insert_returning_id(test_db,
        "INSERT INTO grid_audits (name, status, business_unit_id) VALUES ('Audit 1','Active',100)", ())
    ctrl_id = database.insert_returning_id(test_db,
        "INSERT INTO grid_controls (audit_id, framework_id, control_id, name) VALUES (%s,%s,'A.1','Ctrl')",
        (audit_id, fw_id))
    test_db.commit()

    claimed = pub.claim_next_job(test_db, is_postgres=False)
    pub.process_job(test_db, claimed, is_postgres=False)

    ev = test_db.execute(
        "SELECT * FROM grid_evidence_files WHERE control_id=%s AND aria_policy_version_id=%s",
        (ctrl_id, version["id"]),
    ).fetchone()
    assert ev is not None


def test_grid_evidence_skipped_for_incompatible_business_unit(test_db, monkeypatch):
    job, version, doc, *_ = _approve_a_policy(test_db, monkeypatch, control_ref="A.1", framework="ISO 27001")
    test_db.execute("INSERT INTO business_units (id, name, is_active) VALUES (200, 'OtherBU', 1)")
    fw_id = database.insert_returning_id(test_db,
        "INSERT INTO grid_frameworks (name) VALUES ('ISO 27001')", ())
    audit_id = database.insert_returning_id(test_db,
        # A different, unrelated subsidiary's audit (BU 200, not the policy's BU 100)
        "INSERT INTO grid_audits (name, status, business_unit_id) VALUES ('Audit 2','Active',200)", ())
    database.insert_returning_id(test_db,
        "INSERT INTO grid_controls (audit_id, framework_id, control_id, name) VALUES (%s,%s,'A.1','Ctrl')",
        (audit_id, fw_id))
    test_db.commit()

    claimed = pub.claim_next_job(test_db, is_postgres=False)
    pub.process_job(test_db, claimed, is_postgres=False)

    count = test_db.execute(
        "SELECT COUNT(*) AS c FROM grid_evidence_files WHERE aria_policy_version_id=%s", (version["id"],)
    ).fetchone()["c"]
    assert count == 0, "must never attach a subsidiary-private policy to an unrelated subsidiary's audit"


def test_grid_evidence_not_attached_to_an_org_wide_audit_for_a_bu_private_policy(test_db, monkeypatch):
    """An org-wide (no-BU) audit's evidence is visible to every BU in the
    org -- GRID's own evidence listing does not itself re-check ARIA scope
    before showing title/notes/document-id metadata, so attaching a
    BU-private policy here would leak that metadata org-wide. Only an
    org-wide policy may attach to an org-wide audit; a BU-scoped one may
    only attach to an audit scoped to that exact same BU."""
    job, version, doc, *_ = _approve_a_policy(test_db, monkeypatch, control_ref="A.1", framework="ISO 27001")
    fw_id = database.insert_returning_id(test_db,
        "INSERT INTO grid_frameworks (name) VALUES ('ISO 27001')", ())
    audit_id = database.insert_returning_id(test_db,
        "INSERT INTO grid_audits (name, status) VALUES ('Audit 3','Active')", ())
    database.insert_returning_id(test_db,
        "INSERT INTO grid_controls (audit_id, framework_id, control_id, name) VALUES (%s,%s,'A.1','Ctrl')",
        (audit_id, fw_id))
    test_db.commit()

    claimed = pub.claim_next_job(test_db, is_postgres=False)
    pub.process_job(test_db, claimed, is_postgres=False)

    count = test_db.execute(
        "SELECT COUNT(*) AS c FROM grid_evidence_files WHERE aria_policy_version_id=%s", (version["id"],)
    ).fetchone()["c"]
    assert count == 0, "a BU-private policy must never attach to an org-wide audit"


def test_grid_evidence_attached_for_an_org_wide_policy_regardless_of_audit_bu(test_db, monkeypatch):
    """An org-wide policy (no BU) is visible to everyone, so it may attach
    to any audit -- org-wide or BU-specific -- without leaking anything
    beyond what is already visible."""
    job, version, doc, *_ = _approve_a_policy(test_db, monkeypatch, control_ref="A.1", framework="ISO 27001")
    test_db.execute("UPDATE aria_documents SET business_unit_id=NULL WHERE id=%s", (doc["id"],))
    test_db.execute("UPDATE aria_policy_versions SET business_unit_id=NULL WHERE id=%s", (version["id"],))
    test_db.commit()

    fw_id = database.insert_returning_id(test_db,
        "INSERT INTO grid_frameworks (name) VALUES ('ISO 27001')", ())
    org_wide_audit = database.insert_returning_id(test_db,
        "INSERT INTO grid_audits (name, status) VALUES ('Audit 4','Active')", ())
    bu_specific_audit = database.insert_returning_id(test_db,
        "INSERT INTO grid_audits (name, status, business_unit_id) VALUES ('Audit 5','Active',100)", ())
    for audit_id in (org_wide_audit, bu_specific_audit):
        database.insert_returning_id(test_db,
            "INSERT INTO grid_controls (audit_id, framework_id, control_id, name) VALUES (%s,%s,'A.1','Ctrl')",
            (audit_id, fw_id))
    test_db.commit()

    claimed = pub.claim_next_job(test_db, is_postgres=False)
    pub.process_job(test_db, claimed, is_postgres=False)

    count = test_db.execute(
        "SELECT COUNT(*) AS c FROM grid_evidence_files WHERE aria_policy_version_id=%s", (version["id"],)
    ).fetchone()["c"]
    assert count == 2, "an org-wide policy may attach to both an org-wide and a BU-specific audit"


# ─────────────────────────────────────────────────────────────────────────
# Failure, backoff, and permanent failure
# ─────────────────────────────────────────────────────────────────────────

def test_a_failed_copy_schedules_a_bounded_retry_without_touching_approval(test_db, monkeypatch):
    job, version, doc, *_ = _approve_a_policy(test_db, monkeypatch)
    monkeypatch.setattr(pub, "_copy_to_evidence_vault",
                         lambda *a, **kw: (_ for _ in ()).throw(RuntimeError("disk full")))
    claimed = pub.claim_next_job(test_db, is_postgres=False)
    outcome = pub.process_job(test_db, claimed, is_postgres=False)
    assert outcome == "retry_scheduled"

    row = test_db.execute(
        "SELECT state, attempts, next_attempt_at, last_error FROM aria_policy_publication_jobs WHERE id=%s",
        (job["id"],),
    ).fetchone()
    assert row["state"] == "pending"
    assert row["attempts"] == 1
    assert row["next_attempt_at"] is not None
    assert "disk full" in row["last_error"]

    # Approval itself must be completely unaffected by a downstream copy failure.
    doc_after = test_db.execute("SELECT status FROM aria_documents WHERE id=%s", (doc["id"],)).fetchone()
    assert doc_after["status"] == "Approved"
    version_after = test_db.execute("SELECT state FROM aria_policy_versions WHERE id=%s", (version["id"],)).fetchone()
    assert version_after["state"] == "approved"


def test_job_permanently_fails_after_max_attempts_and_notifies_managers(test_db, monkeypatch):
    job, version, doc, author, approver = _approve_a_policy(test_db, monkeypatch)
    monkeypatch.setattr(pub, "_copy_to_evidence_vault",
                         lambda *a, **kw: (_ for _ in ()).throw(RuntimeError("persistent failure")))

    for _ in range(pub.MAX_ATTEMPTS):
        # Each real failed attempt schedules next_attempt_at minutes in the
        # future (bounded backoff) -- clear it here to simulate that much
        # real time having actually passed, rather than slowing the test
        # down or mocking utcnow just to prove the attempt-counting logic.
        test_db.execute(
            "UPDATE aria_policy_publication_jobs SET next_attempt_at=NULL WHERE id=%s", (job["id"],)
        )
        test_db.commit()
        claimed = pub.claim_next_job(test_db, is_postgres=False)
        assert claimed is not None, "job must remain claimable until permanently failed"
        pub.process_job(test_db, claimed, is_postgres=False)

    row = test_db.execute(
        "SELECT state FROM aria_policy_publication_jobs WHERE id=%s", (job["id"],)
    ).fetchone()
    assert row["state"] == "failed"

    notif = test_db.execute(
        "SELECT * FROM notifications WHERE user_id=%s AND title LIKE %s",
        (approver["id"], "%synchronization needs attention%"),
    ).fetchone()
    assert notif is not None, "a scoped manager (compliance_manager, same org) must be notified"


def test_failure_notification_excludes_a_same_org_different_bu_manager(test_db, monkeypatch):
    """A compliance_manager in a different business unit of the same
    organization holds the role, but not read access to this BU-private
    document -- notifying them would leak its title and doc_id."""
    job, version, doc, author, approver = _approve_a_policy(test_db, monkeypatch)
    test_db.execute("INSERT INTO business_units (id, name, is_active) VALUES (200, 'OtherBU', 1)")
    _user(test_db, 50, org_id=1, bu_id=200, username="other_bu_manager")
    _role(test_db, 50, "compliance_manager")
    test_db.commit()

    monkeypatch.setattr(pub, "_copy_to_evidence_vault",
                         lambda *a, **kw: (_ for _ in ()).throw(RuntimeError("persistent failure")))
    for _ in range(pub.MAX_ATTEMPTS):
        test_db.execute(
            "UPDATE aria_policy_publication_jobs SET next_attempt_at=NULL WHERE id=%s", (job["id"],)
        )
        test_db.commit()
        claimed = pub.claim_next_job(test_db, is_postgres=False)
        pub.process_job(test_db, claimed, is_postgres=False)

    notif = test_db.execute(
        "SELECT * FROM notifications WHERE user_id=50 AND title LIKE %s",
        ("%synchronization needs attention%",),
    ).fetchone()
    assert notif is None, "a manager outside the document's BU must not be notified of it"


# ─────────────────────────────────────────────────────────────────────────
# Explicit retry action
# ─────────────────────────────────────────────────────────────────────────

def test_retry_now_resets_a_failed_job_for_a_scoped_manager(test_db, monkeypatch):
    job, version, doc, author, approver = _approve_a_policy(test_db, monkeypatch)
    test_db.execute("UPDATE aria_policy_publication_jobs SET state='failed' WHERE id=%s", (job["id"],))
    test_db.commit()

    result = pub.retry_now(test_db, approver, job["id"])
    assert result["state"] == "pending"
    assert result["next_attempt_at"] is None


def test_retry_now_refuses_for_an_out_of_scope_actor(test_db, monkeypatch):
    job, version, doc, *_ = _approve_a_policy(test_db, monkeypatch, org_id=1, bu_id=100)
    test_db.execute("UPDATE aria_policy_publication_jobs SET state='failed' WHERE id=%s", (job["id"],))
    test_db.commit()

    _org(test_db, 2)
    _user(test_db, 99, org_id=2, username="outsider")
    _role(test_db, 99, "compliance_manager")
    outsider = _actor(test_db, 99)

    with pytest.raises(svc.NotFoundError):
        pub.retry_now(test_db, outsider, job["id"])


def test_retry_now_refuses_a_job_that_is_not_failed(test_db, monkeypatch):
    job, version, doc, author, approver = _approve_a_policy(test_db, monkeypatch)
    with pytest.raises(svc.NotFoundError):
        pub.retry_now(test_db, approver, job["id"])


# ─────────────────────────────────────────────────────────────────────────
# Draft expiry (section 7.5)
# ─────────────────────────────────────────────────────────────────────────

def test_expire_stale_drafts_moves_a_past_due_draft_and_keeps_body(test_db):
    _org(test_db)
    _bu(test_db, 100, "Finance")
    _user(test_db, 1, bu_id=100, username="author")
    _role(test_db, 1, "policy_author")
    test_db.commit()
    author = _actor(test_db, 1)
    draft = svc.create_draft_from_generation(
        test_db, author, control={"id": 1, "ref": "A.1", "name": "T", "framework_id": 1, "fw_name": "ISO"},
        generated_content="# Body text", org_name="Econet", doc_type="Policy", framework_label="ISO",
        integrated_controls=None, request_id=None, target_business_unit_id=100,
    )
    test_db.execute(
        "UPDATE aria_policy_drafts SET expires_at='2000-01-01T00:00:00', "
        "source_path='x', branded_path='y' WHERE id=%s", (draft["id"],)
    )
    test_db.commit()

    count = svc.expire_stale_drafts(test_db, 1)
    assert count == 1

    row = test_db.execute("SELECT * FROM aria_policy_drafts WHERE id=%s", (draft["id"],)).fetchone()
    assert row["state"] == "expired"
    assert row["body"] == "# Body text", "section 7.5: keep text/metadata for recovery"
    assert row["source_path"] is None
    assert row["branded_path"] is None


def test_expire_stale_drafts_leaves_a_fresh_draft_alone(test_db):
    _org(test_db)
    _bu(test_db, 100, "Finance")
    _user(test_db, 1, bu_id=100, username="author")
    _role(test_db, 1, "policy_author")
    test_db.commit()
    author = _actor(test_db, 1)
    draft = svc.create_draft_from_generation(
        test_db, author, control={"id": 1, "ref": "A.1", "name": "T", "framework_id": 1, "fw_name": "ISO"},
        generated_content="# Body", org_name="Econet", doc_type="Policy", framework_label="ISO",
        integrated_controls=None, request_id=None, target_business_unit_id=100,
    )
    count = svc.expire_stale_drafts(test_db, 1)
    assert count == 0
    row = test_db.execute("SELECT state FROM aria_policy_drafts WHERE id=%s", (draft["id"],)).fetchone()
    assert row["state"] == "editing"


def test_a_save_refreshes_expires_at_to_a_real_future_timestamp(test_db, monkeypatch):
    """Two deterministic, clearly-separated timestamps rather than two
    live utcnow() calls a few lines apart: on Windows in particular, two
    back-to-back calls can read the exact same system clock tick, which
    would make a strict '>' comparison between them flaky, not a real
    assertion about the refresh behavior."""
    import datetime as _dt
    from modules.aria import policy_workflow_service as svc_mod

    _org(test_db)
    _bu(test_db, 100, "Finance")
    _user(test_db, 1, bu_id=100, username="author")
    _role(test_db, 1, "policy_author")
    test_db.commit()
    author = _actor(test_db, 1)

    t1 = _dt.datetime(2026, 1, 1, 0, 0, 0)
    monkeypatch.setattr(svc_mod, "utcnow", lambda: t1)
    draft = svc.create_draft_from_generation(
        test_db, author, control={"id": 1, "ref": "A.1", "name": "T", "framework_id": 1, "fw_name": "ISO"},
        generated_content="# Body", org_name="Econet", doc_type="Policy", framework_label="ISO",
        integrated_controls=None, request_id=None, target_business_unit_id=100,
    )
    row = test_db.execute("SELECT expires_at FROM aria_policy_drafts WHERE id=%s", (draft["id"],)).fetchone()
    assert row["expires_at"] == (t1 + _dt.timedelta(days=30)).isoformat()

    t2 = t1 + _dt.timedelta(days=5)
    monkeypatch.setattr(svc_mod, "utcnow", lambda: t2)
    saved = svc.save_draft_body(test_db, author, draft["id"], body="# Body v2",
                                 expected_lock_version=draft["lock_version"])
    assert saved["expires_at"] == (t2 + _dt.timedelta(days=30)).isoformat()
    assert saved["expires_at"] > row["expires_at"]


# ─────────────────────────────────────────────────────────────────────────
# Storage retention (section 7.5)
# ─────────────────────────────────────────────────────────────────────────

def test_move_orphans_to_trash_moves_only_old_enough_unreferenced_dirs():
    old_build, old_dir = storage.new_staging_dir(org_id=1)
    (old_dir / "f.txt").write_bytes(b"x")
    import os as _os
    old_ts = __import__("time").time() - 100 * 3600
    _os.utime(old_dir, (old_ts, old_ts))

    fresh_build, fresh_dir = storage.new_staging_dir(org_id=1)
    (fresh_dir / "f.txt").write_bytes(b"x")

    referenced_build, referenced_dir = storage.new_staging_dir(org_id=1)
    (referenced_dir / "f.txt").write_bytes(b"x")
    referenced_rel = storage.relative_path(referenced_dir / "f.txt")
    ref_dir_rel = storage.relative_path(referenced_dir)
    _os.utime(referenced_dir, (old_ts, old_ts))  # old, but referenced -- must never move

    trashed = storage.move_orphans_to_trash(org_id=1, referenced_relative_paths={referenced_rel}, grace_hours=24)

    trashed_names = {t.split("/")[-1] for t in trashed}
    assert old_build in trashed_names
    assert fresh_build not in trashed_names, "not old enough yet"
    assert not any(ref_dir_rel.split("/")[-1] == n for n in trashed_names), "referenced dir must never be trashed"
    assert not old_dir.exists()
    assert fresh_dir.exists()
    assert referenced_dir.exists()


def test_purge_expired_trash_deletes_only_old_enough_entries():
    build_id, staging = storage.new_staging_dir(org_id=1)
    (staging / "f.txt").write_bytes(b"x")
    trashed = storage.move_orphans_to_trash(org_id=1, referenced_relative_paths=set(), grace_hours=0)
    assert len(trashed) == 1
    trash_path = storage.resolve_stored_path(trashed[0])

    # Backdate the .trashed_at sidecar so it looks like it has been sitting
    # in trash long enough to purge.
    old_ts = __import__("time").time() - 10 * 86400
    (trash_path / ".trashed_at").write_text(str(old_ts), encoding="utf-8")

    purged = storage.purge_expired_trash(org_id=1, retention_days=7)
    assert trashed[0] in purged
    assert not trash_path.exists()


def test_purge_expired_trash_leaves_recent_entries():
    build_id, staging = storage.new_staging_dir(org_id=1)
    (staging / "f.txt").write_bytes(b"x")
    trashed = storage.move_orphans_to_trash(org_id=1, referenced_relative_paths=set(), grace_hours=0)
    purged = storage.purge_expired_trash(org_id=1, retention_days=7)
    assert purged == []
    assert storage.resolve_stored_path(trashed[0]).exists()


# ─────────────────────────────────────────────────────────────────────────
# policy_published_handler: no double vault-copy (section 10.3)
# ─────────────────────────────────────────────────────────────────────────

def test_published_handler_skips_legacy_vault_copy_for_a_managed_publication(test_db, monkeypatch):
    from core import event_handlers
    monkeypatch.setattr(event_handlers, "get_db", lambda: test_db)
    monkeypatch.setattr(test_db, "close", lambda: None)  # handler closes its db; keep it open for assertions

    doc_pk = database.insert_returning_id(test_db,
        "INSERT INTO aria_documents (doc_id, framework, control_ref, title, version, status, "
        "policy_workflow_managed) VALUES ('DOC-0050','ISO 27001','A.1','Managed Policy','1.0','Approved',1)", ())
    test_db.commit()

    event_handlers.policy_published_handler(
        event_type="aria.policy.published", source_module="aria", entity_type="document",
        entity_id=doc_pk, payload={"title": "Managed Policy", "framework": "ISO 27001",
                                    "control_ref": "A.1", "version": "1.0", "version_id": 999},
        user_id=1,
    )

    # The legacy path dedups/creates by tag=f"aria_doc_id={entity_id}" -- it
    # must not have run at all for a managed-version payload.
    count = test_db.execute(
        "SELECT COUNT(*) AS c FROM evidence_items WHERE tags LIKE %s", (f"%aria_doc_id={doc_pk}%",)
    ).fetchone()["c"]
    assert count == 0


def test_published_handler_still_runs_legacy_vault_copy_without_a_version_id(test_db, monkeypatch):
    from core import event_handlers
    monkeypatch.setattr(event_handlers, "get_db", lambda: test_db)
    monkeypatch.setattr(test_db, "close", lambda: None)

    _org(test_db)
    _user(test_db, 1, username="author")
    doc_pk = database.insert_returning_id(test_db,
        "INSERT INTO aria_documents (doc_id, framework, control_ref, title, version, status, "
        "policy_workflow_managed) VALUES ('DOC-0051','ISO 27001','A.1','Legacy Policy','1.0','Approved',0)", ())
    test_db.commit()

    event_handlers.policy_published_handler(
        event_type="aria.policy.published", source_module="aria", entity_type="document",
        entity_id=doc_pk, payload={"title": "Legacy Policy", "framework": "ISO 27001", "control_ref": "A.1"},
        user_id=1,
    )

    count = test_db.execute(
        "SELECT COUNT(*) AS c FROM evidence_items WHERE tags LIKE %s", (f"%aria_doc_id={doc_pk}%",)
    ).fetchone()["c"]
    assert count == 1, "a plain legacy publication (no version_id) must keep working exactly as before"


# ─────────────────────────────────────────────────────────────────────────
# Current-version-aware search indexing (section 10.2)
# ─────────────────────────────────────────────────────────────────────────

def test_reindex_document_uses_the_approved_version_body_not_stale_document_body(test_db, monkeypatch):
    from modules.aria import ask_service
    ask_service.init_index()
    job, version, doc, *_ = _approve_a_policy(test_db, monkeypatch)

    # aria_documents.body reflects whatever the draft held at confirm time,
    # not necessarily the approved version's exact text -- simulate that
    # divergence directly to prove the index reads the version, not the
    # (potentially stale) document row.
    test_db.execute("UPDATE aria_documents SET body='STALE DOCUMENT BODY' WHERE id=%s", (doc["id"],))
    test_db.commit()

    ask_service.reindex_document(doc["doc_id"])

    rows = test_db.execute(
        "SELECT body FROM aria_ask_index WHERE content_type='document' AND content_id=%s",
        (doc["doc_id"],),
    ).fetchall()
    combined = " ".join(r["body"] for r in rows)
    assert "STALE DOCUMENT BODY" not in combined
    assert version["body"] in combined or "Policy" in combined


def test_process_job_refreshes_the_search_index_on_success(test_db, monkeypatch):
    from modules.aria import ask_service
    ask_service.init_index()
    job, version, doc, *_ = _approve_a_policy(test_db, monkeypatch)
    test_db.execute(
        "DELETE FROM aria_ask_index WHERE content_type='document' AND content_id=%s", (doc["doc_id"],)
    )
    test_db.commit()

    claimed = pub.claim_next_job(test_db, is_postgres=False)
    outcome = pub.process_job(test_db, claimed, is_postgres=False)
    assert outcome == "complete"

    count = test_db.execute(
        "SELECT COUNT(*) AS c FROM aria_ask_index WHERE content_type='document' AND content_id=%s",
        (doc["doc_id"],),
    ).fetchone()["c"]
    assert count > 0


def test_reindex_document_never_falls_back_to_document_body_for_an_unapproved_managed_document(test_db):
    """confirm_draft sets current_policy_version_id immediately for a
    brand-new document's very first version, which is inserted with
    state='draft' -- reachable directly through the admin 'rebuild index'
    action (rebuild_all calls reindex_document for every document
    unconditionally), not just a theoretical race."""
    from modules.aria import ask_service
    ask_service.init_index()
    _org(test_db)
    _bu(test_db, 100, "Finance")
    _user(test_db, 1, bu_id=100, username="author")
    _role(test_db, 1, "policy_author")
    test_db.commit()
    author = _actor(test_db, 1)

    draft = svc.create_draft_from_generation(
        test_db, author, control={"id": 1, "ref": "A.1", "name": "T", "framework_id": 1, "fw_name": "ISO"},
        generated_content="# UNAPPROVED CANDIDATE TEXT", org_name="Econet", doc_type="Policy",
        framework_label="ISO", integrated_controls=None, request_id=None, target_business_unit_id=100,
    )
    built = svc.build_draft(test_db, author, draft["id"], _template_id(test_db, 1), draft["lock_version"])
    confirmed = svc.confirm_draft(test_db, author, built["id"], built["build_id"], built["lock_version"])

    doc_row = test_db.execute(
        "SELECT current_policy_version_id, policy_workflow_managed FROM aria_documents WHERE id=%s",
        (confirmed["document_id"],),
    ).fetchone()
    assert doc_row["current_policy_version_id"] == confirmed["version_id"], (
        "precondition for this test to mean anything: current_policy_version_id must "
        "already point at the still-unapproved version"
    )
    version_state = test_db.execute(
        "SELECT state FROM aria_policy_versions WHERE id=%s", (confirmed["version_id"],)
    ).fetchone()["state"]
    assert version_state != "approved"

    ask_service.reindex_document(confirmed["doc_id"])

    rows = test_db.execute(
        "SELECT body FROM aria_ask_index WHERE content_type='document' AND content_id=%s",
        (confirmed["doc_id"],),
    ).fetchall()
    combined = " ".join(r["body"] for r in rows)
    assert "UNAPPROVED CANDIDATE TEXT" not in combined
    assert rows == [], "an unapproved managed document has nothing effective to search yet"


def _template_id(db, org_id):
    from docx import Document as DocxDocument
    doc = DocxDocument()
    doc.add_paragraph("TEMPLATE COVER")
    buf = io.BytesIO()
    doc.save(buf)
    db.execute(
        "INSERT INTO aria_doc_templates (name, file_path, file_name, org_id, is_active) "
        "VALUES ('T1', 't1.docx', 't1.docx', %s, 1)", (org_id,)
    )
    db.commit()
    tid = db.execute("SELECT id FROM aria_doc_templates WHERE file_path='t1.docx'").fetchone()["id"]
    from modules.aria import policy_storage as storage
    template_dir = storage.ARIA_UPLOAD_DIR.parent / "aria_templates"
    template_dir.mkdir(parents=True, exist_ok=True)
    (template_dir / "t1.docx").write_bytes(buf.getvalue())
    import modules.aria.routes as _routes
    _routes.ARIA_TEMPLATE_DIR = template_dir
    return tid


# ─────────────────────────────────────────────────────────────────────────
# Tenant context (org/schema binding for background jobs)
# ─────────────────────────────────────────────────────────────────────────

def test_tenant_context_binds_then_restores_previous_state(test_db):
    import database
    assert database._current_tenant.get() is None
    assert database._current_org_id.get() is None

    with database.tenant_context(org_id=7, slug="ecocash"):
        assert database._current_tenant.get() == "ecocash"
        assert database._current_org_id.get() == 7
        assert database._current_is_super.get() is True

    assert database._current_tenant.get() is None
    assert database._current_org_id.get() is None
    assert database._current_is_super.get() is False


def test_tenant_context_restores_even_on_exception(test_db):
    import database
    with pytest.raises(RuntimeError):
        with database.tenant_context(org_id=7, slug="ecocash"):
            raise RuntimeError("boom")
    assert database._current_tenant.get() is None
    assert database._current_org_id.get() is None


def test_list_active_tenants_excludes_inactive_orgs(test_db):
    import database
    test_db.execute("INSERT INTO organizations (id, name, slug, status) VALUES (1,'A','a','active')")
    test_db.execute("INSERT INTO organizations (id, name, slug, status) VALUES (2,'B','b','suspended')")
    test_db.commit()
    tenants = database.list_active_tenants()
    assert (1, "a") in tenants
    assert not any(t[0] == 2 for t in tenants)


def test_drain_publication_queue_binds_tenant_context_per_org(test_db, monkeypatch):
    """A scheduler tick has no request to inherit tenant context from --
    prove _drain_publication_queue actually binds each active org's
    context around its own slice of work, rather than running everything
    unbound against whatever the default/public schema is."""
    import database
    from modules.aria import scheduler as aria_scheduler

    # scheduler.py does `from database import ... list_active_tenants`, so
    # the name to patch is the one bound in scheduler's own namespace, not
    # database's -- monkeypatching database.list_active_tenants would not
    # reach a reference scheduler.py already captured at import time.
    monkeypatch.setattr(aria_scheduler, "list_active_tenants", lambda: [(1, "ecocash"), (2, "omni")])
    bound_during_call = []

    def fake_run_due_jobs(limit=10):
        bound_during_call.append((database._current_org_id.get(), database._current_tenant.get()))
        return {"complete": 0, "retry_scheduled": 0, "failed": 0}

    monkeypatch.setattr(aria_scheduler.policy_publication, "run_due_jobs", fake_run_due_jobs)
    aria_scheduler._drain_publication_queue()

    assert bound_during_call == [(1, "ecocash"), (2, "omni")]
    # And context is restored afterward, not leaked to whatever runs next
    # on a reused scheduler thread.
    assert database._current_org_id.get() is None
    assert database._current_tenant.get() is None
