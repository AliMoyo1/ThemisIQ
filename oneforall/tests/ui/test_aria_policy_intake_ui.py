"""The guided import is usable in a real browser, not just as service calls."""
import io
import json
import uuid

from docx import Document


def test_table_metadata_prefills_and_existing_import_can_refresh(login_as, live_app, synthetic_tenant,
                                                                 monkeypatch, tmp_path):
    import database
    from config import settings
    from modules.aria import routes as aria_routes
    monkeypatch.setattr(aria_routes, "ARIA_UPLOAD_DIR", tmp_path / "aria")
    monkeypatch.setattr(settings, "ARIA_POLICY_AUTHORING_ENABLED", True)
    monkeypatch.setattr(settings, "ARIA_POLICY_AUTHORING_ORG_IDS", [synthetic_tenant["org_id"]])
    doc = Document()
    doc.add_paragraph("DATA RETENTION POLICY")
    control = doc.add_table(rows=2, cols=4)
    for row, values in zip(control.rows, (
        ("Version Number", "3.0", "Implementation Date", "2026/03/19"),
        ("Approval Date", "19/03/2026", "Review Date", "19/03/2027"),
    )):
        for cell, value in zip(row.cells, values): cell.text = value
    approval = doc.add_table(rows=2, cols=4)
    for row, values in zip(approval.rows, (
        ("Authorized by", "Full Names", "Signature", "Date"),
        ("COO", "External Approver", "", "19-03-2026"),
    )):
        for cell, value in zip(row.cells, values): cell.text = value
    output = io.BytesIO(); doc.save(output)
    page = login_as("compliance_manager")
    page.goto(f"{live_app}/aria/policy-imports")
    page.set_input_files("#intake-files", files=[{"name": "retention_v3.0.docx", "mimeType":
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "buffer": output.getvalue()}])
    page.click('#intake-upload button[type="submit"]')
    page.wait_for_selector("#intake-review:not([hidden])", timeout=15000)
    assert page.locator("#intake-field-version").input_value() == "3.0"
    assert page.locator("#intake-field-approver").input_value() == "External Approver"
    assert page.locator("#intake-field-effective_date").input_value() == "2026-03-19"
    assert "2026-03-19" in page.locator("#intake-external-approval-suggestion").inner_text()
    import_id = page.locator("[data-import-id]").first.get_attribute("data-import-id")
    db = database.get_db()
    try:
        old = {"parser": "local-docx-1", "metadata": {}, "spans": [], "text_status": "extracted"}
        db.execute("UPDATE aria_policy_imports SET extracted_json=%s WHERE id=%s", (json.dumps(old), import_id))
        db.commit()
    finally:
        db.close()
    page.reload()
    page.click(f'[data-import-id="{import_id}"]')
    page.wait_for_selector("#intake-review:not([hidden])")
    assert page.locator("#intake-field-version").input_value() == ""
    page.fill("#intake-field-title", "My reviewed title")
    page.click("#intake-refresh")
    page.wait_for_function("() => document.getElementById('intake-field-version').value === '3.0'")
    assert page.locator("#intake-field-title").input_value() == "My reviewed title"
    assert page.locator("#intake-field-approver").input_value() == "External Approver"
    assert not page.console_errors, page.console_errors


def test_policy_intake_upload_obeys_organization_rollout(login_as, live_app, synthetic_tenant,
                                                         monkeypatch):
    from config import settings
    monkeypatch.setattr(settings, "ARIA_POLICY_AUTHORING_ENABLED", True)
    monkeypatch.setattr(settings, "ARIA_POLICY_AUTHORING_ORG_IDS", [])
    page = login_as("compliance_manager")
    document = Document()
    document.add_paragraph("Import must stay off for this organization.")
    output = io.BytesIO()
    document.save(output)
    response = page.context.request.post(f"{live_app}/aria/api/policy-imports", multipart={
        "file": {"name": "policy.docx", "mimeType":
                 "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                 "buffer": output.getvalue()},
    })
    assert response.status == 403
    assert "not enabled" in response.json()["error"]


def test_existing_policy_import_and_composition_review_queue(login_as, live_app, synthetic_tenant,
                                                             monkeypatch, tmp_path):
    from modules.aria import routes as aria_routes
    from modules.aria import policy_storage, policy_composition
    from modules.evidence import routes as evidence_routes

    monkeypatch.setattr(aria_routes, "ARIA_UPLOAD_DIR", tmp_path / "aria")
    monkeypatch.setattr(policy_storage, "ARIA_UPLOAD_DIR", tmp_path / "aria")
    monkeypatch.setattr(policy_storage, "WORKFLOW_ROOT", tmp_path / "aria" / "policy_workflow")
    monkeypatch.setattr(evidence_routes, "EVIDENCE_DIR", tmp_path / "evidence")
    monkeypatch.setattr(policy_composition, "policy_authoring_enabled_for", lambda org: org == synthetic_tenant["org_id"])
    from config import settings
    monkeypatch.setattr(settings, "ARIA_POLICY_AUTHORING_ENABLED", True)
    monkeypatch.setattr(settings, "ARIA_POLICY_AUTHORING_ORG_IDS", [synthetic_tenant["org_id"]])

    doc = Document()
    doc.add_paragraph("Subsidiary cover")
    doc.add_paragraph("Contents")
    doc.add_paragraph("Purpose: this policy protects information.")
    out = io.BytesIO(); doc.save(out)

    page = login_as("compliance_manager")
    page.goto(f"{live_app}/aria/policy-imports")
    page.set_input_files("#intake-files", files=[{"name": "subsidiary.docx", "mimeType":
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "buffer": out.getvalue()}])
    page.click('#intake-upload button[type="submit"]')
    page.wait_for_selector("#intake-review:not([hidden])", timeout=15000)
    assert page.locator("#intake-source-text").inner_text().find("Purpose") >= 0
    page.fill("#intake-field-title", "Reviewed Security Policy")
    page.fill("#intake-field-version", "1.0")
    page.click("#intake-register")
    page.wait_for_selector("#intake-composition:not([hidden])", timeout=15000)
    page.select_option("#intake-compose-boundary", "2")
    page.fill("#intake-cover-text", "Group policy cover")
    page.click("#intake-compose-submit")
    page.wait_for_selector("#intake-compose-review:not([hidden])", timeout=15000)
    assert page.locator("#intake-candidate-file").get_attribute("href").startswith("/aria/api/policy-imports/")
    assert not page.console_errors, page.console_errors


def test_new_policy_editor_saves_structured_diagram(login_as, live_app, synthetic_tenant):
    import database
    owner = synthetic_tenant["users"]["compliance_manager"]["user_id"]
    org = synthetic_tenant["org_id"]
    draft_id = uuid.uuid4().hex
    db = database.get_db()
    try:
        db.execute("INSERT INTO aria_policy_drafts "
                   "(id,org_id,owner_user_id,created_by,reserved_doc_id,version_major,version_minor,"
                   "content_kind,body,metadata_json,author_user_ids_json,state) "
                   "VALUES (%s,%s,%s,%s,'DOC-DIAGRAM-UI',1,0,'markdown','# Test Policy',%s,%s,'editing')",
                   (draft_id, org, owner, owner, json.dumps({"title": "Diagram policy"}), json.dumps([owner])))
        db.commit()
    finally:
        db.close()
    page = login_as("compliance_manager")
    page.goto(f"{live_app}/aria/ai-generator?draft={draft_id}")
    page.wait_for_selector('#workflowPanel:not([style*="display:none"])', timeout=15000)
    page.click("#workflowDiagrams summary")
    page.click("#workflowAddDiagram")
    page.fill(".aria-diagram-caption", "Information handoff")
    page.fill(".aria-diagram-alt", "Compliance hands data to the privacy team")
    page.fill(".aria-diagram-nodes", "A: Compliance\nB: Privacy")
    page.fill(".aria-diagram-edges", "A>B")
    page.click("#draftSaveBtn")
    page.wait_for_function("() => document.getElementById('draftSavedIndicator').style.display !== 'none'")
    db = database.get_db()
    try:
        saved = json.loads(db.execute("SELECT metadata_json FROM aria_policy_drafts WHERE id=%s", (draft_id,)).fetchone()[0])
        assert saved["diagrams"][0]["edges"] == [{"from": "A", "to": "B"}]
    finally:
        db.close()
    assert not page.console_errors, page.console_errors
