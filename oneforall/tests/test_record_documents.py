"""
Real, permission gated documents of one DPIA, one business continuity plan and one AI impact assessment
(decision D12 of PLAN_record_document_exports_2026-10-22.md, which closes D11 and gap G8 of PLAN_report_exports_2026-10-21.md).

Until now a DPIA downloaded only as the plain text of its AI draft, and a plan and an assessment could not be exported at all. Each
record now has a route beside it that needs what opens the record (`sentinel.dpia.manage`, `module.bcm.access`, `ai_assessment.view`),
reads the record through the module's own business unit rule (a record of another unit answers 404, like a missing one), writes
the download to the audit trail and answers through the kit's one download response.

The expectations are written out by hand from a seeded file. Each "cannot" (another unit's record, a file or record the reader may not
open, a role without the capability) has a positive twin, so an empty document cannot pass for isolation. The tests of fixed behaviour fail
on the old code: the routes did not exist, the plan's reviews, activations and write routes answered for another unit's plan, and the
plan page named another unit's process in its coverage list.
"""
import dataclasses
import io
import json
import re
from datetime import datetime
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import FastAPI
from pypdf import PdfReader
from starlette.testclient import TestClient

import core.middleware as middleware
import modules.ai_governance.routes as ai_routes
import modules.bcm.routes as bcm_routes
import modules.sentinel.routes as sentinel_routes
from core import rbac
from core.report_kit import theme
from core.report_kit.model import Kpis, Paragraphs
from core.report_kit.render_docx import render_docx
from core.report_kit.render_html import render_html
from core.report_kit.render_pdf import render_pdf
from core.report_kit.render_xlsx import render_xlsx
from modules.ai_governance import assessment_document, assessments
from modules.bcm import data_service as bcm_ds, plan_document
from modules.governance import connected_records, relationships
from modules.governance.entity_scope import linkable_keys
from modules.sentinel import data_service as sentinel_ds, dpia_document
from tests.test_platform_ownership import _insert, world  # noqa: F401

ORIGIN = "https://app.example.test"
NOW = datetime(2026, 10, 22, 8, 0, 0)
PDF = "application/pdf"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

RISKS = ("Unauthorised access to identity documents (Likelihood: Medium, Impact: High)\n"
         "Template kept longer than needed (Likelihood: Low, Impact: High)\n"
         "Inaccurate match refuses a genuine customer (Likelihood: Medium, Impact: Medium)\n"
         "A worry nobody rated")
DRAFT = ("# Data protection impact assessment\n\n## 1. Description\nThe company collects an identity document.\n\n"
         "## 2. Risks\n| Risk | Rating |\n|---|---|\n| Access | High |\n\n## 3. Measures\n- Encrypt documents\n- Delete templates\n")
PLAN = ("# Data centre failover\n\n## 1. Purpose and scope\nRestores settlement and billing.\n\n## 2. Roles\n"
        "| Role | Holder |\n|---|---|\n| Incident commander | Head of Operations |\n\n## 3. Recovery\n1. Declare the incident\n2. Fail over\n")


# ── The app, the people and the seeded file ─────────────────────────────────

@pytest.fixture
def web(world, monkeypatch):
    who = {"name": "super"}

    async def current_user(request):
        return world.users.get(who["name"])

    monkeypatch.setattr(middleware, "get_current_user", current_user)
    monkeypatch.setenv("APP_URL", ORIGIN)
    app = FastAPI()
    for module in (sentinel_routes, bcm_routes, ai_routes):
        app.include_router(module.router)
    client = TestClient(app, base_url="http://testserver", raise_server_exceptions=False, follow_redirects=False)

    def call(method, path, name="super", **kwargs):
        who["name"] = name
        return getattr(client, method)(path, **kwargs)

    call.world = world
    return call


def _persona(world, name, role, unit_tag=None):
    unit = world.unit.get(unit_tag) if unit_tag else None
    uid = 100 + len(world.users)
    world.users[name] = {"id": uid, "username": name, "org_id": 1, "business_unit_id": unit, "is_super_admin": 0, "roles": [role]}
    _insert(world.db, "users", id=uid, username=name, email=f"{name}@example.test", full_name=name, password_hash="x", org_id=1,
            business_unit_id=unit)
    return world.users[name]


def _file(world, title, links, unit=None, **cols):
    """One Evidence Vault file and its links: `(module, entity type, record id, extra link columns)`."""
    item = _insert(world.db, "evidence_items", title=title, category=cols.pop("category", "contract"), status=cols.pop("status", "current"),
                   version=cols.pop("version", 1), org_id=1, business_unit_id=world.unit[unit] if unit else None, **cols)
    for module, kind, record, extra in links:
        _insert(world.db, "evidence_links", evidence_id=item, module=module, entity_type=kind, entity_id=record, **extra)
    return item


def _estate(world):
    """Privacy, continuity and AI governance records in units A and B and organization wide, each with something hidden from the other unit.

    Unit A: a DPIA (rated high before and medium after the measures, four identified risks, three of them with an impact) that
    assesses a RoPA entry, a plan that covers two processes the reader may open and one it may not, an assessment of an AI system.
    Unit B holds a record of each kind that unit A must never see named. Every hidden thing carries the one word marker Bravohidden."""
    db, unit = world.db, world.unit
    e = SimpleNamespace()
    e.ropa_a = _insert(db, "sentinel_ropa", ref_number="R-A", processing_name="Customer onboarding and KYC", business_unit_id=unit["A"])
    e.ropa_a2 = _insert(db, "sentinel_ropa", ref_number="R-A2", processing_name="Marketing campaigns", business_unit_id=unit["A"])
    e.ropa_b = _insert(db, "sentinel_ropa", ref_number="R-B", processing_name="Bravohidden", business_unit_id=unit["B"])
    dpia = dict(status="completed", regulation="GDPR", department="Customer Operations", purpose="Confirm who a customer is",
                activity_desc="Collecting an identity document from every customer.", legal_basis="Legal obligation",
                data_categories=json.dumps(["Name", "ID number", "Address"]), special_cats=json.dumps(["Biometric data"]),
                data_subjects="Customers", subject_count="About 2,400,000", retention="7 years after closure", systems="KYC platform",
                processors="Identity provider", intl_transfer="Yes", transfer_dest="South Africa", transfer_mech="SCCs",
                necessity="Checking identity is required by law.", proportionality="Only the required fields are kept.",
                risks=json.dumps([RISKS]), mitigations="Encrypt at rest.\nDelete templates after 30 days.", overall_risk="high",
                residual_risk="medium", dpo_consulted=1, auth_consulted=0, subjects_consulted=0, consult_notes="The DPO reviewed it.",
                dpo_name="Nyasha Moyo", ai_full_dpia=DRAFT)
    e.dpia_a = _insert(db, "sentinel_dpias", ref_number="DPIA-A-1", title="DPIA: Customer onboarding", business_unit_id=unit["A"],
                       ropa_id=e.ropa_a, **dpia)
    e.dpia_org = _insert(db, "sentinel_dpias", ref_number="DPIA-ORG-1", title="DPIA: Visitor register", business_unit_id=None)
    e.dpia_b = _insert(db, "sentinel_dpias", ref_number="DPIA-B-1", title="DPIA: Bravohidden", business_unit_id=unit["B"])
    e.dpia_bare = _insert(db, "sentinel_dpias", ref_number="DPIA-A-2", title="DPIA: Nothing rated yet", business_unit_id=unit["A"])
    db.execute("UPDATE sentinel_ropa SET dpia_id = %s WHERE id = %s", (e.dpia_a, e.ropa_a))
    db.commit()

    # continuity
    plan = dict(plan_type="Disaster recovery", department="Technology", scope="Settlement and billing", owner="Nyasha Moyo", version="2.1",
                status="approved", content=PLAN, review_frequency="annual", last_reviewed="2026-03-01", next_review="2026-09-01")
    e.plan_a = _insert(db, "bcm_plans", title="Data centre failover", description="Restores the Harare data centre.", business_unit_id=unit["A"], **plan)
    e.plan_org = _insert(db, "bcm_plans", title="Organization wide plan", business_unit_id=None, version="1.0", status="draft")
    e.plan_b = _insert(db, "bcm_plans", title="Plan Bravohidden", business_unit_id=unit["B"], version="1.0", status="draft")
    bia_a = _insert(db, "bcm_bia_records", process_name="Mobile money settlement", rto_hours=2, criticality="critical", business_unit_id=unit["A"])
    bia_org = _insert(db, "bcm_bia_records", process_name="Customer care", rto_hours=8, criticality="high", business_unit_id=None)
    bia_none = _insert(db, "bcm_bia_records", process_name="Archive retrieval", rto_hours=None, criticality="low", business_unit_id=None)
    bia_b = _insert(db, "bcm_bia_records", process_name="Bravohidden payroll", rto_hours=1, criticality="high", business_unit_id=unit["B"])
    for bia in (bia_a, bia_org, bia_none, bia_b):
        _insert(db, "bcm_bia_plan_links", bia_id=bia, plan_id=e.plan_a, coverage_type="primary")
    e.bia = {"a": bia_a, "org": bia_org, "none": bia_none, "b": bia_b}
    for score, when in ((58, "2026-01-20 09:00:00"), (82, "2026-09-14 15:10:00")):
        _insert(db, "bcm_plan_reviews", plan_id=e.plan_a, reviewer_name="Reviewer", provider="AI assistant", overall_score=score, created_at=when)
    e.incident = _insert(db, "bcm_incidents", title="Power failure at the data centre", business_unit_id=unit["A"])
    _insert(db, "bcm_plan_activations", plan_id=e.plan_a, action="activated", reason="Power failure", activated_by="Tatenda", incident_id=e.incident,
            created_at="2026-09-02 02:20:00")
    _insert(db, "bcm_plan_activations", plan_id=e.plan_a, action="deactivated", reason="Services restored", activated_by="Tatenda",
            created_at="2026-09-02 05:40:00")
    e.exercise = _insert(db, "bcm_exercises", title="Live failover test", plan_id=e.plan_a, status="completed", business_unit_id=unit["A"], org_id=1)
    _insert(db, "bcm_exercises", title="Exercise Bravohidden", plan_id=e.plan_a, status="planned", business_unit_id=unit["B"], org_id=1)

    # AI governance (through the module's own service, then the stored classification is set so the expectation is fixed)
    e.app_a = _insert(db, "applications", name="Support Assistant", business_unit_id=unit["A"], is_active=1)
    dpo_a, super_ = world.users["dpo_a"], world.users["super"]
    made = assessments.create_assessment(db, dpo_a, {
        "title": "AI customer support assistant", "ai_system_name": "Support Assistant", "department": "Customer Operations", "owner": "Nyasha Moyo",
        "system_description": "A chat assistant on the customer app.", "outputs_decisions": "It answers; an agent decides.", "third_party": True,
        "third_party_details": "Hosted by an outside provider.", "autonomy_level": "human_in_loop", "deployment_env": "cloud", "status": "approved",
        "data_categories": ["Name", "Chat transcripts"], "application_id": e.app_a, "ropa_id": e.ropa_a, "dpia_id": e.dpia_a,
        "mitigation_measures": "Complaints go to an agent.", "residual_classification": "moderate",
        "impacts": [{"dimension_name": "Privacy", "applicable": 1, "likelihood": 3, "impact": 4, "description": "Chats hold account data."},
                    {"dimension_name": "Financial", "applicable": 1, "likelihood": 2, "impact": 3, "description": "A wrong fee answer."},
                    {"dimension_name": "Societal", "applicable": 0, "description": "None expected."}]})
    e.aiia_a = made["assessment"]["id"]
    e.aiia_a2 = assessments.create_assessment(db, dpo_a, {"title": "Second look at the assistant", "application_id": e.app_a})["assessment"]["id"]
    e.aiia_b = assessments.create_assessment(db, world.users["dpo_b"], {"title": "Assessment Bravohidden"})["assessment"]["id"]
    e.aiia_r = assessments.create_assessment(db, super_, {"title": "Cites a hidden record", "business_unit_id": unit["A"], "ropa_id": e.ropa_b})["assessment"]["id"]
    db.execute("UPDATE sentinel_aiia SET overall_classification = 'High' WHERE id = %s", (e.aiia_a,))
    db.commit()
    return e


@pytest.fixture
def estate(world):
    _persona(world, "bcm_a", rbac.BCM_MANAGER, "A")
    _persona(world, "bcm_b", rbac.BCM_MANAGER, "B")
    _persona(world, "owner_none", rbac.RISK_OWNER)
    return _estate(world)


def _vault(world, e):
    """Files in the Vault. Each line says what it is for."""
    dpia, ropa, plan = ("sentinel", "dpia", e.dpia_a), ("sentinel", "ropa", e.ropa_a), ("bcm", "plan", e.plan_a)
    e.ev = SimpleNamespace(
        # on the DPIA, the plan, another entry of unit A and an entry of unit B: it names the visible entry and never the hidden one
        agreement=_file(world, "Processor agreement", [(*dpia, {"review_status": "accepted"}), (*plan, {"review_status": "accepted"}),
                                                       ("sentinel", "ropa", e.ropa_a2, {}), ("sentinel", "ropa", e.ropa_b, {})],
                        version=2, expiry_date="2999-01-01"),
        schedule=_file(world, "Retention schedule", [(*ropa, {})], category="policy"),                       # on the entry the DPIA assesses
        hidden=_file(world, "Bravohidden contract", [(*dpia, {})], unit="B"),                                  # a file of unit B
        replaced=_file(world, "Old agreement", [(*dpia, {})], status="superseded"),
        archived=_file(world, "Archived note", [(*dpia, {})], status="archived"),
        removed=_file(world, "Removed link", [(*dpia, {"deleted_at": "2026-10-01 00:00:00"})]),
        expired=_file(world, "Expired certificate", [(*dpia, {"review_status": "accepted"})], expiry_date="2001-01-01"),
        returned=_file(world, "Returned report", [(*dpia, {"review_status": "returned"})]),
        process=_file(world, "Recovery runbook", [("bcm", "bia", e.bia["a"], {"review_status": "accepted"})]))
    return e.ev


# ── Helpers ─────────────────────────────────────────────────────────────────

def _text(body):
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(body)).pages)


def _everything(report):
    """All of a report as one string, so a name that must not be in it can be searched for."""
    return json.dumps(dataclasses.asdict(report), default=str)


def _kpis(report):
    return {k.label: k for block in report.blocks if isinstance(block, Kpis) for k in block.items}


def _chart(report, title):
    return next(c for c in report.charts() if c.title == title)


def _table(report, title):
    return next(t for t in report.tables() if t.title == title)


def _rows(report, title, key="title"):
    return {r[key]: r for r in _table(report, title).rows}


def _paragraphs(report):
    return {b.title: b.paragraphs for b in report.blocks if isinstance(b, Paragraphs)}


def _links(body):
    out = []
    for page in PdfReader(io.BytesIO(body)).pages:
        for annotation in page.get("/Annots") or []:
            uri = annotation.get_object().get("/A", {}).get("/URI")
            if uri:
                out.append(uri)
    return out


def _path(kind, record_id, query=""):
    return {"dpia": "/sentinel/api/dpias/%d/document", "plan": "/bcm/api/plans/%d/document",
            "assessment": "/ai-governance/api/assessments/%d/document"}[kind] % record_id + query


def _paths(e):
    return {"dpia": _path("dpia", e.dpia_a), "plan": _path("plan", e.plan_a), "assessment": _path("assessment", e.aiia_a)}


# ── The routes ──────────────────────────────────────────────────────────────

def test_each_document_downloads_as_a_pdf_with_a_safe_name_for_someone_who_may_open_the_record(web, estate):
    stems = {"dpia": ("dpo_a", "ThemisIQ_privacy_DPIA-A-1"), "plan": ("owner_a", f"ThemisIQ_continuity_plan_{estate.plan_a}_Data_centre_failover"),
             "assessment": ("dpo_a", r"ThemisIQ_ai_governance_AIIA-\d{8}-[A-Z0-9]{5}")}
    for kind, path in _paths(estate).items():
        persona, stem = stems[kind]
        response = web("get", path, persona)
        assert response.status_code == 200, (kind, response.text[:200])
        assert response.headers["content-type"].startswith(PDF)
        assert re.fullmatch(rf'attachment; filename="{stem}_\d{{8}}_\d{{4}}\.pdf"', response.headers["content-disposition"]), kind
        assert response.headers["x-content-type-options"] == "nosniff" and response.headers["cache-control"] == "no-store"
        assert response.content.startswith(b"%PDF") and len(PdfReader(io.BytesIO(response.content)).pages) >= 3


def test_the_excel_version_opens_and_holds_the_tables_of_the_document(web, estate):
    for kind, path in _paths(estate).items():
        response = web("get", path + "?format=xlsx", "super")
        assert response.status_code == 200 and response.headers["content-type"].startswith(XLSX), kind
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        assert {"Cover", "Evidence", "Connected"} <= set(book.sheetnames), (kind, book.sheetnames)
        assert web("get", path + "?format=csv", "super").status_code == 400, "only pdf and xlsx are served"


@pytest.mark.parametrize("kind, mine, other, shared, nobody", [
    ("dpia", "dpo_a", "dpo_b", "dpia_org", "dpo_no_unit"), ("plan", "owner_a", "owner_b", "plan_org", "owner_none"),
    ("assessment", "dpo_a", "dpo_b", None, "dpo_no_unit")])
def test_a_record_in_another_unit_answers_like_a_missing_one_and_its_twin_opens(web, estate, kind, mine, other, shared, nobody):
    record = {"dpia": estate.dpia_a, "plan": estate.plan_a, "assessment": estate.aiia_a}[kind]
    hidden, missing = web("get", _path(kind, record), other), web("get", _path(kind, 987654), other)
    assert (hidden.status_code, hidden.content) == (missing.status_code, missing.content) and hidden.status_code == 404
    assert web("get", _path(kind, record), mine).status_code == 200, "the same person's own unit opens"
    assert web("get", _path(kind, record), "super").status_code == 200
    assert web("get", _path(kind, record), nobody).status_code == 404, "a person with no unit sees only organization wide records"
    if shared:
        assert web("get", _path(kind, getattr(estate, shared)), other).status_code == 200, "an organization wide record is in every unit"
        assert web("get", _path(kind, getattr(estate, shared)), nobody).status_code == 200, "and for a person with no unit"


def test_the_capability_that_opens_the_record_gates_the_document(web, estate):
    paths = _paths(estate)
    # a risk owner reaches the plan and the assessment pages, not the DPIAs; a data protection officer reaches the DPIAs and the
    # assessments, not the continuity module; an employee reaches only the continuity module
    for persona, kind, status in (("owner_a", "dpia", 403), ("employee_a", "dpia", 403), ("dpo_a", "plan", 403), ("employee_a", "assessment", 403),
                                  ("employee_a", "plan", 200), ("owner_a", "assessment", 200), ("dpo_a", "dpia", 200)):
        assert web("get", paths[kind], persona).status_code == status, (persona, kind)
    for kind, path in paths.items():
        assert web("get", path, None).status_code == 303, "signed out is sent to log in"


def test_every_download_is_audited_and_a_refusal_is_not(web, estate):
    paths = _paths(estate)
    web("get", paths["dpia"], "dpo_a")
    web("get", paths["plan"] + "?format=xlsx", "owner_a")
    web("get", paths["assessment"], "dpo_a")
    web("get", paths["dpia"], "dpo_b")
    web("get", paths["dpia"], "owner_a")
    rows = [dict(r) for r in web.world.db.execute(
        "SELECT username, module, action, entity_type, entity_id, details FROM audit_log "
        "WHERE action IN ('dpia_exported', 'plan_exported', 'ai_assessment_exported') ORDER BY id").fetchall()]
    assert [(r["username"], r["module"], r["action"], r["entity_type"], r["entity_id"]) for r in rows] == [
        ("dpo_a", "sentinel", "dpia_exported", "dpia", estate.dpia_a), ("owner_a", "bcm", "plan_exported", "plan", estate.plan_a),
        ("dpo_a", "ai_governance", "ai_assessment_exported", "ai_assessment", estate.aiia_a)], "one entry per download, none for a refusal"
    assert rows[0]["details"].startswith("ref=DPIA-A-1; format=pdf; tier=operational"), "the reference, never the title"
    assert rows[1]["details"].startswith("version=2.1; format=xlsx; tier=operational")
    assert "AI customer support assistant" not in rows[2]["details"] and "format=pdf" in rows[2]["details"]


def test_the_download_name_is_ascii_whatever_the_title_says(web, estate):
    """A hand built header with a smart quote, a dash or a non Latin letter is an HTTP 500: the name goes through the kit."""
    title = "Plan \N{LEFT DOUBLE QUOTATION MARK}B\N{RIGHT DOUBLE QUOTATION MARK} \N{EN DASH} Zimb\N{LATIN SMALL LETTER A WITH ACUTE}bwe \N{CJK UNIFIED IDEOGRAPH-5317}\N{CJK UNIFIED IDEOGRAPH-4EAC}"
    web.world.db.execute("UPDATE bcm_plans SET title = %s WHERE id = %s", (title, estate.plan_a))
    web.world.db.commit()
    response = web("get", _paths(estate)["plan"], "owner_a")
    assert response.status_code == 200
    disposition = response.headers["content-disposition"]
    assert disposition.isascii() and re.fullmatch(r'attachment; filename="ThemisIQ_continuity_plan_\d+_[A-Za-z0-9._-]+\.pdf"', disposition)


# ── What each document states ───────────────────────────────────────────────

def test_the_dpia_states_what_is_recorded_and_the_charts_come_from_the_same_rows(world, estate):
    d = dpia_document.build(world.users["dpo_a"], sentinel_ds.get_dpia(estate.dpia_a), NOW)
    assert d.meta.tier == theme.OPERATIONAL and d.meta.reference == "DPIA-A-1" and "Privacy" in d.meta.module_label
    kpis = _kpis(d)
    assert (kpis["Status"].value, kpis["Overall risk"].value, kpis["Residual risk"].value) == ("Completed", "High", "Medium")
    assert kpis["DPO consulted"].value == "Yes" and kpis["Authority consulted"].value == "Not recorded"
    assert kpis["Data subjects"].value == "About 2,400,000"
    before = _chart(d, "Risk level before and after the measures")
    assert before.data == [("Before the measures: High", 3), ("After the measures: Medium", 2)]
    assert "from High to Medium" in before.takeaway
    risks = _table(d, "Identified risks").rows
    assert [(r["risk"], r["likelihood"], r["impact"]) for r in risks] == [
        ("Unauthorised access to identity documents", "Medium", "High"), ("Template kept longer than needed", "Low", "High"),
        ("Inaccurate match refuses a genuine customer", "Medium", "Medium"), ("A worry nobody rated", None, None)]
    impact = _chart(d, "Identified risks by impact")
    assert impact.data == [("High", 2), ("Medium", 1), ("Not rated", 1)], "the chart adds up to the four rows of the table"
    assert "2 of 3 rated risks have a high impact" in impact.takeaway and "1 risk has no impact recorded" in impact.takeaway
    assert _chart(d, "Personal data in scope by kind").data == [("Ordinary data categories", 3), ("Special category data", 1)]
    facts = {r["k"]: r["v"] for r in _table(d, "What is processed, why and how").rows}
    assert facts["Retention"] == "7 years after closure" and facts["Transfer destination"] == "South Africa"
    paragraphs = _paragraphs(d)
    assert paragraphs["Measures and controls"] == ["Encrypt at rest.", "Delete templates after 30 days."]
    assert "The data protection officer is recorded as consulted. Data protection officer: Nyasha Moyo." in paragraphs["Who was consulted"]
    draft = [b for b in d.blocks if isinstance(b, Paragraphs) and b.label]
    assert draft and "does not record that a person reviewed it" in draft[0].label, "the AI text is marked, and not claimed to be unreviewed"
    assert paragraphs["1. Description"] == ["The company collects an identity document."]
    assert any(t.title == "" and t.rows == [{"c0": "Access", "c1": "High"}] for t in d.tables()), "a table of the draft is a table"


def test_the_dpia_shows_its_owner_and_due_date_when_the_record_holds_them(world, estate):
    """The form records an assessor or owner and a due date. A tile is shown only for what is stored, so a record without them
    shows no gap the form never asked about, and a due date in the past is flagged unless the assessment is already completed."""
    record = sentinel_ds.get_dpia(estate.dpia_a)
    assert "Due date" not in _kpis(dpia_document.build(world.users["dpo_a"], record, NOW)), "nothing stored, nothing shown"
    open_late = {**record, "status": "in_progress", "owner": "Nyasha Moyo", "due_date": "2026-10-01"}
    kpis = _kpis(dpia_document.build(world.users["dpo_a"], open_late, NOW))
    assert kpis["Assessor or owner"].value == "Nyasha Moyo"
    assert (kpis["Due date"].value, kpis["Due date"].note, kpis["Due date"].tone) == ("1 Oct 2026", "past the due date", "bad")
    done = _kpis(dpia_document.build(world.users["dpo_a"], {**open_late, "status": "completed"}, NOW))
    assert (done["Due date"].note, done["Due date"].tone) == ("", "neutral"), "a completed assessment is not late"
    ahead = _kpis(dpia_document.build(world.users["dpo_a"], {**open_late, "due_date": "2026-12-31"}, NOW))
    assert (ahead["Due date"].value, ahead["Due date"].note) == ("31 Dec 2026", "")


def test_a_dpia_with_no_rating_reads_n_a_and_never_the_columns_default(world, estate):
    """`risk_level` is a column that defaults to medium for every row. A person who rated nothing has rated nothing."""
    record = sentinel_ds.get_dpia(estate.dpia_bare)
    assert record["risk_level"] is None, "the column no longer defaults to a rating nobody gave"
    d = dpia_document.build(world.users["dpo_a"], record, NOW)
    kpis = _kpis(d)
    assert (kpis["Overall risk"].value, kpis["Residual risk"].value, kpis["Data subjects"].value) == ("n/a", "n/a", "n/a")
    assert kpis["Overall risk"].note == "not rated"
    chart = _chart(d, "Risk level before and after the measures")
    assert chart.is_empty() and "need a rating" in chart.takeaway
    facts = {r["k"]: r["v"] for r in _table(d, "What is processed, why and how").rows}
    assert facts["Purpose"] == "Not recorded" and facts["Retention"] == "Not recorded", "a gap in the record is visible, never blank"
    assert "Transfer destination" not in facts, "a transfer's details are shown only when there is a transfer"
    assert "Controller" not in facts, "the form does not ask for a controller, so the row is not shown as a gap"


def test_the_assessment_states_the_scores_and_counts_what_is_not_scored(world, estate):
    shown = assessments.get_assessment(world.db, world.users["dpo_a"], estate.aiia_a)
    a = assessment_document.build(world.users["dpo_a"], shown, NOW)
    assert a.meta.reference == f"{shown['ref_number']}, version 1" and a.meta.tier == theme.OPERATIONAL
    kpis = _kpis(a)
    assert (kpis["Status"].value, kpis["Overall classification"].value, kpis["Residual classification"].value) == ("Approved", "High", "Moderate")
    assert (kpis["Autonomy"].value, kpis["Deployment"].value, kpis["Third party AI"].value) == ("Human in loop", "Cloud", "Yes")
    assert kpis["Overall classification"].note == "from the scored impacts"
    assert _chart(a, "Impact score by dimension").data == [("Privacy", 12), ("Financial", 6)]
    assert _chart(a, "Impact dimensions by state").data == [("Scored", 2), ("Applies, not yet scored", 5), ("Does not apply", 1)]
    assert _chart(a, "Impact dimensions by state").takeaway == "2 of 8 dimensions are scored. 5 still need a score."
    one_left = {**shown, "impacts": [r if r["dimension_name"] == "Security" else {**r, "likelihood": r["likelihood"] or 1, "impact": r["impact"] or 1}
                                     for r in shown["impacts"]]}
    singular = _chart(assessment_document.build(world.users["dpo_a"], one_left, NOW), "Impact dimensions by state").takeaway
    assert singular == "6 of 8 dimensions are scored. 1 still needs a score."
    rows = _rows(a, "Likelihood and impact, each from 1 to 5", "dim")
    assert (rows["Privacy"]["l"], rows["Privacy"]["i"], rows["Privacy"]["s"]) == (3, 4, 12)
    assert (rows["Operational"]["l"], rows["Operational"]["s"], rows["Operational"]["applies"]) == (None, None, "Yes"), "not scored is n/a, never zero"
    assert rows["Societal"]["applies"] == "No"
    system = {r["k"]: r["v"] for r in _table(a, "What it is and where it runs").rows}
    assert system["Registered system"] == "Support Assistant" and system["Record of processing"] == "R-A Customer onboarding and KYC"
    assert system["Data protection impact assessment"] == "DPIA-A-1 DPIA: Customer onboarding"


def test_a_privacy_record_the_reader_may_not_open_is_not_named_in_the_assessment(world, estate):
    for persona, named in (("dpo_a", False), ("super", True)):
        shown = assessments.get_assessment(world.db, world.users[persona], estate.aiia_r)
        a = assessment_document.build(world.users[persona], shown, NOW)
        assert ("Bravohidden" in _everything(a)) is named, persona
        row = {r["k"]: r["v"] for r in _table(a, "What it is and where it runs").rows}["Record of processing"]
        assert row == ("R-B Bravohidden" if named else "Not recorded"), persona


def test_the_plan_lists_only_the_processes_the_reader_may_open_and_counts_what_has_no_objective(world, estate):
    plan = bcm_ds.get_plan(estate.plan_a)      # the route has already let this person open the plan; the coverage is read through their scope
    for persona, names in (("owner_a", ["Mobile money settlement", "Customer care", "Archive retrieval"]),
                           ("owner_none", ["Customer care", "Archive retrieval"]),
                           ("super", ["Bravohidden payroll", "Mobile money settlement", "Customer care", "Archive retrieval"])):
        report = plan_document.build(world.users[persona], plan, NOW)
        assert [r["process"] for r in _table(report, "Processes this plan covers").rows] == names, persona
        assert ("Bravohidden payroll" in _everything(report)) is (persona == "super"), persona
    report = plan_document.build(world.users["owner_a"], plan, NOW)
    rto = _chart(report, "Recovery time objective of each process the plan covers")
    assert rto.data == [("Mobile money settlement", 2), ("Customer care", 8)]
    assert "tightest objective is Mobile money settlement: 2 hours" in rto.takeaway and "1 covered process has no recovery time objective" in rto.takeaway
    score = _chart(report, "AI review score of the plan over time")
    assert score.data == [("20 Jan 2026", 58), ("14 Sep 2026", 82)] and "latest score is 82 out of 100" in score.takeaway


def test_the_plan_states_its_review_position_and_history_as_recorded(world, estate):
    report = plan_document.build(world.users["owner_a"], bcm_ds.get_plan(estate.plan_a), NOW)
    kpis = _kpis(report)
    assert (kpis["Status"].value, kpis["Version"].value, kpis["Process or unit owner"].value) == ("Approved", "2.1", "Nyasha Moyo")
    assert (kpis["Next review"].value, kpis["Next review"].note) == ("1 Sep 2026", "51 days overdue"), "1 September to 22 October"
    assert kpis["In force"].value == "Not activated" and report.meta.reference == f"Plan {estate.plan_a}, version 2.1"
    history = [(r["action"], r["reason"], r["by"]) for r in _table(report, "Activation history").rows]
    assert history == [("Deactivated", "Services restored", "Tatenda"), ("Activated", "Power failure", "Tatenda")]
    assert _paragraphs(report)["3. Recovery"] == ["1. Declare the incident", "2. Fail over"]
    bare = plan_document.build(world.users["owner_a"], bcm_ds.get_plan(estate.plan_org), NOW)
    assert _kpis(bare)["Next review"].value == "n/a" and _kpis(bare)["Next review"].note == "no date is set"
    assert _chart(bare, "Recovery time objective of each process the plan covers").is_empty()


# ── Linking: the Evidence Vault, the records connected, the live record ─────

def test_the_documents_in_the_vault_are_listed_with_their_state_and_open_the_file_in_the_vault(world, estate, monkeypatch):
    monkeypatch.setenv("APP_URL", ORIGIN)
    ev = _vault(world, estate)
    d = dpia_document.build(world.users["dpo_a"], sentinel_ds.get_dpia(estate.dpia_a), NOW)
    rows = _rows(d, "Documents in the Evidence Vault")
    assert set(rows) == {"Processor agreement", "Retention schedule", "Old agreement", "Expired certificate", "Returned report"}, \
        "not a file of another unit, an archived file or a removed link; a replaced version is listed as replaced, as the connected page lists it"
    assert rows["Processor agreement"]["state"] == "Accepted, expires 1 Jan 2999" and rows["Processor agreement"]["kind"] == "Contract, version 2"
    assert rows["Processor agreement"]["on"] == "This record"
    assert rows["Retention schedule"]["on"] == "Customer onboarding and KYC (Privacy: Processing record)", "the files of the entry the DPIA assesses"
    assert rows["Retention schedule"]["state"] == "Waiting for a reviewer" and rows["Retention schedule"]["kind"] == "Policy, version 1"
    assert rows["Expired certificate"]["state"] == "Expired" and rows["Returned report"]["state"] == "Returned for changes"
    assert rows["Old agreement"]["state"] == "Replaced by a newer version"
    assert rows["Processor agreement"]["url"] == f"{ORIGIN}/evidence/?open=item:{ev.agreement}", "the row opens the file in the Vault"
    # the file links the records to each other: it names the other record the reader may open, never the one they may not
    assert rows["Processor agreement"]["also"] == "Marketing campaigns (Privacy: Processing record)", \
        "the plan is not theirs to open; the entry of unit B is hidden"
    assert rows["Retention schedule"]["also"] == "No other record"
    assert "Bravohidden" not in _everything(d)
    wide = _rows(dpia_document.build(world.users["super"], sentinel_ds.get_dpia(estate.dpia_a), NOW), "Documents in the Evidence Vault")
    assert "Bravohidden contract" in wide, "the same file opens for a person whose scope holds it"
    assert wide["Processor agreement"]["also"] == ("Data centre failover (Resilience: Continuity plan); Marketing campaigns (Privacy: Processing record); "
                                                   "Bravohidden (Privacy: Processing record)")


def test_a_document_lists_what_the_connected_page_lists_to_the_same_person(world, estate):
    """The connected page is the read of record for the files and connections of a record; the document must not compute them again."""
    from modules.governance import entity360
    _vault(world, estate)
    admin = world.users["super"]
    relationships.create_link(world.db, admin, ("sentinel", "dpia", estate.dpia_a), ("sentinel", "breach", world.rec[("breach", "A")]), "related")
    world.db.commit()
    for persona in ("dpo_a", "dpo_a1", "dpo_b", "super"):
        user = world.users[persona]
        view = entity360.get_connected(world.db, user, "sentinel", "dpia", estate.dpia_a)
        if view is None:
            continue                  # a person who may not open the record has no page and no document
        d = dpia_document.build(user, sentinel_ds.get_dpia(estate.dpia_a), NOW)
        own = {r["title"] for r in _table(d, "Documents in the Evidence Vault").rows if r["on"] == "This record"}
        assert own == {i["title"] for i in view["proof"]["items"]}, persona
        listed = {r["title"] for r in _table(d, "Records connected to this one").rows}
        assert {c["title"] for section in view["connections"] for c in section["items"]} <= listed, persona
        assert view["proof"]["items"], persona
    assert entity360.get_connected(world.db, world.users["dpo_b"], "sentinel", "dpia", estate.dpia_a) is None, "a hidden record has no page"


def test_a_source_that_could_not_be_read_is_said_so_and_never_shown_as_empty(world, estate, monkeypatch):
    from modules.governance import entity360
    _vault(world, estate)

    def broken(*args, **kwargs):
        raise RuntimeError("the vault is down")

    monkeypatch.setattr(entity360, "_proof", broken)
    d = dpia_document.build(world.users["dpo_a"], sentinel_ds.get_dpia(estate.dpia_a), NOW)
    table = _table(d, "Documents in the Evidence Vault")
    assert table.rows == [] and "could not be read just now, which is not the same as none" in table.empty_text
    assert "could not be read just now" in table.note


def test_a_plan_lists_the_files_of_the_processes_it_covers_and_a_file_of_the_plan(world, estate, monkeypatch):
    monkeypatch.setenv("APP_URL", ORIGIN)
    ev = _vault(world, estate)
    rows = _rows(plan_document.build(world.users["owner_a"], bcm_ds.get_plan(estate.plan_a), NOW), "Documents in the Evidence Vault")
    assert set(rows) == {"Processor agreement", "Recovery runbook"}
    assert rows["Processor agreement"]["on"] == "This record"
    assert rows["Recovery runbook"]["on"] == "Mobile money settlement (Resilience: Business impact analysis)"
    assert rows["Recovery runbook"]["url"] == f"{ORIGIN}/evidence/?open=item:{ev.process}"


def test_the_records_connected_are_the_links_the_references_and_the_records_that_cite_it(world, estate, monkeypatch):
    monkeypatch.setenv("APP_URL", ORIGIN)
    db, admin = world.db, world.users["super"]
    dpia = ("sentinel", "dpia", estate.dpia_a)
    for breach in (world.rec[("breach", "A")], world.rec[("breach", "B")]):
        relationships.create_link(db, admin, dpia, ("sentinel", "breach", breach), "related")
    relationships.create_link(db, admin, dpia, ("governance", "application", estate.app_a), "assesses")
    relationships.create_link(db, admin, dpia, ("sentinel", "breach", world.rec[("breach", "A1")]), "related", provenance="system_candidate")
    disputed = relationships.create_link(db, admin, dpia, ("erm", "risk", world.rec[("erm", "A")]), "related")["link_id"]
    relationships.set_link_status(db, admin, disputed, "rejected")
    db.commit()
    record = sentinel_ds.get_dpia(estate.dpia_a)
    d = dpia_document.build(world.users["dpo_a"], record, NOW)
    rows = _rows(d, "Records connected to this one")
    assert list(rows)[0] == "DPIA: Customer onboarding" and rows["DPIA: Customer onboarding"]["how"] == "The live record this document was made from"
    assert rows["DPIA: Customer onboarding"]["url"] == f"{ORIGIN}/sentinel/?open=dpia:{estate.dpia_a}"
    assert rows["Customer onboarding and KYC"]["how"] == "Assesses" and rows["Customer onboarding and KYC"]["where"] == "Privacy: Processing record"
    assert rows["needle breach A"]["how"] == "Related to" and rows["needle breach A"]["where"] == "Privacy: Breach"
    assert rows["Support Assistant"]["how"] == "Assesses" and rows["Support Assistant"]["where"] == "AI Governance: Application"
    cited = rows["AI customer support assistant"]
    assert (cited["how"], cited["where"], cited["url"]) == ("Is cited by", "AI Governance: AI impact assessment",
                                                            f"{ORIGIN}/ai-governance/assessments/{estate.aiia_a}")
    assert rows["needle breach A1"]["how"] == "Related to (proposed, not confirmed)", "a connection nobody confirmed is marked, as on the connected page"
    assert "needle breach B" not in rows, "a record in another unit is hidden"
    assert "needle breach B" not in _everything(d)
    assert "needle breach B" in _rows(dpia_document.build(admin, record, NOW), "Records connected to this one"), "the same link opens for a wider scope"
    wide = _rows(dpia_document.build(admin, record, NOW), "Records connected to this one")
    assert wide["ERM risk A"]["how"] == "Related to (disputed: somebody declined this connection)"


def test_the_plan_names_the_exercises_that_tested_it_and_the_incident_it_was_activated_for(world, estate):
    plan = bcm_ds.get_plan(estate.plan_a)
    rows = _rows(plan_document.build(world.users["owner_a"], plan, NOW), "Records connected to this one")
    assert (rows["Live failover test"]["how"], rows["Live failover test"]["where"]) == ("Is tested by (completed)", "Resilience: Exercise")
    assert rows["Power failure at the data centre"]["how"] == "Was activated for"
    assert "Mobile money settlement" not in rows, "the covered processes are in the coverage table, not listed twice"
    assert "Exercise Bravohidden" not in _everything(plan_document.build(world.users["owner_a"], plan, NOW)), "an exercise of another unit is hidden"
    assert "Exercise Bravohidden" in _rows(plan_document.build(world.users["super"], plan, NOW), "Records connected to this one")


def test_the_assessment_names_its_system_the_records_it_cites_and_other_assessments_of_the_system(world, estate):
    shown = assessments.get_assessment(world.db, world.users["dpo_a"], estate.aiia_a)
    rows = _rows(assessment_document.build(world.users["dpo_a"], shown, NOW), "Records connected to this one")
    assert rows["Support Assistant"]["how"] == "Assesses"
    assert rows["Customer onboarding and KYC"]["how"] == "Cites" and rows["DPIA: Customer onboarding"]["how"] == "Cites"
    assert rows["Second look at the assistant"]["how"] == "Assesses the same AI system"
    assert "Assessment Bravohidden" not in rows


def test_a_link_is_drawn_only_when_an_application_address_is_configured(web, estate, monkeypatch):
    _vault(web.world, estate)
    path = _paths(estate)["dpia"]
    drawn = _links(web("get", path, "dpo_a").content)
    assert f"{ORIGIN}/sentinel/?open=dpia:{estate.dpia_a}" in drawn and any(u.startswith(f"{ORIGIN}/evidence/?open=item:") for u in drawn)
    assert all(u.startswith(ORIGIN + "/") for u in drawn), "every link is the configured address, never the request's host"
    monkeypatch.delenv("APP_URL")
    plain = web("get", path, "dpo_a")
    assert plain.status_code == 200 and _links(plain.content) == [], "no address, no link"
    assert "No application address (APP_URL) is configured" in " ".join(_text(plain.content).split())
    monkeypatch.setenv("APP_URL", "http://insecure.example.test")
    assert _links(web("get", path, "dpo_a").content) == [], "an http address is not trusted for a link"


def test_every_kind_the_platform_can_link_says_where_it_lives_and_what_it_is():
    for key in linkable_keys():
        assert connected_records.where(key), key
    assert connected_records.where(("sentinel", "dpia")) == "Privacy impact assessment", "the kind already says Privacy: the module is not said twice"
    assert connected_records.where(("sentinel", "aiia")) == "AI Governance: AI impact assessment", "an assessment lives in AI Governance, not Privacy"
    assert connected_records.where(("bcm", "plan")) == "Resilience: Continuity plan"
    assert connected_records.where(("evidence", "item")) == "Evidence Vault: Item", "a kind the page has no word for gets its own name"


# ── The plan routes apply the business unit rule (these failed before) ──────

def test_the_plan_routes_answer_404_for_another_units_plan_and_change_nothing(web, estate):
    db = web.world.db
    path = f"/bcm/api/plans/{estate.plan_a}"
    before = dict(db.execute("SELECT * FROM bcm_plans WHERE id = %s", (estate.plan_a,)).fetchone())
    for method, suffix, kwargs in (("get", "", {}), ("get", "/reviews", {}), ("get", "/activations", {}), ("put", "", {"json": {"title": "Hijacked"}}),
                                   ("delete", "", {}), ("post", "/activate", {"json": {"reason": "x"}}), ("post", "/deactivate", {"json": {"reason": "x"}}),
                                   ("post", "/reviews", {})):
        response = web(method, path + suffix, "bcm_b", **kwargs)
        assert response.status_code == 404, (method, suffix, response.status_code)
    assert dict(db.execute("SELECT * FROM bcm_plans WHERE id = %s", (estate.plan_a,)).fetchone()) == before, "the plan was not changed or deleted"
    assert db.execute("SELECT COUNT(*) FROM bcm_plan_activations WHERE plan_id = %s", (estate.plan_a,)).fetchone()[0] == 2, "no activation was written"
    for suffix in ("", "/reviews", "/activations"):      # the positive twins: the same calls for a person whose unit holds the plan
        assert web("get", path + suffix, "bcm_a").status_code == 200, suffix
    assert web("put", path, "bcm_a", json={"title": "Data centre failover"}).status_code == 200


def test_the_list_of_active_plans_holds_only_the_plans_the_reader_may_open(web, estate):
    for plan in (estate.plan_a, estate.plan_b, estate.plan_org):
        web.world.db.execute("UPDATE bcm_plans SET is_active_plan = 1 WHERE id = %s", (plan,))
    web.world.db.commit()
    titles = lambda who: sorted(p["title"] for p in web("get", "/bcm/api/plans/active", who).json())
    assert titles("owner_a") == ["Data centre failover", "Organization wide plan"]
    assert titles("owner_b") == ["Organization wide plan", "Plan Bravohidden"]
    assert titles("super") == ["Data centre failover", "Organization wide plan", "Plan Bravohidden"]


def test_the_plan_page_does_not_name_a_process_of_another_unit_in_its_coverage(web, estate):
    names = lambda who: sorted(link["process_name"] for link in web("get", f"/bcm/api/plans/{estate.plan_a}", who).json()["bia_links"])
    assert names("owner_a") == ["Archive retrieval", "Customer care", "Mobile money settlement"]
    assert names("super") == ["Archive retrieval", "Bravohidden payroll", "Customer care", "Mobile money settlement"]


# ── Every format draws every document ───────────────────────────────────────

def test_every_format_renders_every_document_and_the_documents_add_no_dash(world, estate, monkeypatch):
    monkeypatch.setenv("APP_URL", ORIGIN)
    _vault(world, estate)
    admin = world.users["super"]
    reports = [dpia_document.build(admin, sentinel_ds.get_dpia(estate.dpia_a), NOW),
               plan_document.build(admin, bcm_ds.get_plan(estate.plan_a), NOW),
               assessment_document.build(admin, assessments.get_assessment(world.db, admin, estate.aiia_a), NOW)]
    for report in reports:
        pdf = render_pdf(report)
        assert pdf.startswith(b"%PDF") and render_docx(report)[:2] == b"PK" and render_xlsx(report)[:2] == b"PK"
        html = render_html(report)
        assert html.count('<a href="https://app.example.test/') >= 2, "the print view carries the links too"
        dashes = ("\N{EM DASH}", "\N{EN DASH}")
        assert not any(dash in text for dash in dashes for text in (_text(pdf), html)), "no em or en dash"


def test_a_record_with_almost_nothing_in_it_still_renders_in_every_format_and_says_what_is_missing(world, estate):
    """A new DPIA, a plan that only has a description and an assessment nobody has scored are the usual first drafts."""
    admin = world.users["super"]
    only_description = bcm_ds.create_plan({"title": "Described only", "description": "Restore the service in the evening."})
    reports = [dpia_document.build(admin, sentinel_ds.get_dpia(estate.dpia_bare), NOW),
               plan_document.build(admin, bcm_ds.get_plan(only_description), NOW),
               assessment_document.build(admin, assessments.get_assessment(world.db, admin, estate.aiia_a2), NOW)]
    for report in reports:
        assert render_pdf(report).startswith(b"%PDF") and render_docx(report)[:2] == b"PK" and render_xlsx(report)[:2] == b"PK"
        assert render_html(report).startswith("<!doctype html>")
    dpia, plan, assessment = reports
    assert _paragraphs(dpia)["Necessity"] == ["Not recorded."] and _paragraphs(dpia)["Measures and controls"] == ["Not recorded."]
    assert _table(dpia, "Documents in the Evidence Vault").rows == [] and "No document in the Evidence Vault" in _table(dpia, "Documents in the Evidence Vault").empty_text
    assert _kpis(assessment)["Overall classification"].value == "n/a" and _kpis(assessment)["Overall classification"].note == "nothing is scored yet"
    assert _chart(assessment, "Impact score by dimension").is_empty()
    # a plan written as a description alone: the text is shown once, as the plan, and not again as a description
    facts = {r["k"]: r["v"] for r in _table(plan, "The plan").rows}
    assert "Description" not in facts and _paragraphs(plan)["Plan text"] == ["Restore the service in the evening."]
