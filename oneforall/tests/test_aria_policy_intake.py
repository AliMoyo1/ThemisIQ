"""Existing-policy intake: provenance, version integrity and scope."""
import io
import json
import hashlib

import pytest
from docx import Document

from modules.aria import policy_intake as intake
from modules.aria import policy_composition as composition
from modules.aria import policy_pdf_rendition as pdf_rendition
from modules.aria import policy_references


def _setup(db, monkeypatch, tmp_path):
    from modules.aria import routes as aria_routes
    from modules.evidence import routes as evidence_routes
    monkeypatch.setattr(aria_routes, "ARIA_UPLOAD_DIR", tmp_path / "aria")
    monkeypatch.setattr(evidence_routes, "EVIDENCE_DIR", tmp_path / "evidence")
    for org in (10, 11):
        db.execute("INSERT INTO organizations (id,name,slug) VALUES (%s,%s,%s)",
                   (org, f"Org {org}", f"org{org}"))
    for bu in (110, 120):
        db.execute("INSERT INTO business_units (id,name,is_active) VALUES (%s,%s,1)", (bu, f"BU {bu}"))
    for uid, org, bu, role in ((101, 10, 110, "compliance_manager"),
                               (102, 10, 110, "policy_approver"), (103, 10, 120, "compliance_manager"),
                               (104, 11, 110, "compliance_manager")):
        db.execute("INSERT INTO users (id,username,email,full_name,password_hash,org_id,business_unit_id) "
                   "VALUES (%s,%s,%s,%s,'x',%s,%s)",
                   (uid, f"user{uid}", f"user{uid}@example.invalid", f"User {uid}", org, bu))
        db.execute("INSERT INTO user_roles (user_id,role_key) VALUES (%s,%s)", (uid, role))
    db.execute("INSERT INTO frameworks (id,name,is_active) VALUES (701,'ISO 27001',1)")
    db.execute("INSERT INTO controls (id,framework_id,ref,name) VALUES (801,701,'A.5.1','Policies for information security')")
    db.commit()


def _actor(db, uid):
    user = dict(db.execute("SELECT id,username,full_name,org_id,business_unit_id,is_super_admin "
                           "FROM users WHERE id=%s", (uid,)).fetchone())
    user["roles"] = [r[0] for r in db.execute("SELECT role_key FROM user_roles WHERE user_id=%s", (uid,)).fetchall()]
    return user


def _docx():
    doc = Document()
    doc.core_properties.title = "Information Security Policy"
    doc.add_paragraph("Version: 2.1")
    doc.add_paragraph("Owner: Compliance Team")
    doc.add_paragraph("A.5.1 requires policies for information security.")
    out = io.BytesIO(); doc.save(out)
    return out.getvalue()


def _table_policy_docx(explicit_owner=""):
    doc = Document()
    doc.add_paragraph("DATA RETENTION POLICY")
    control = doc.add_table(rows=3, cols=4)
    for row, values in zip(control.rows, (
        ("Document Classification:", "Company Confidential", "Version N°:", "3.0"),
        ("Implementation Date:", "2026/03/19", "Approval Date:", "19/03/2026"),
        ("Document Custodian:", "19/03/2027", "", ""),
    )):
        for cell, value in zip(row.cells, values):
            cell.text = value
    history = doc.add_table(rows=3, cols=4)
    for row, values in zip(history.rows, (
        ("Version Number", "Author(s)", "Changes", "Date"),
        ("2.0", "Earlier Author", "Earlier revision", "2025/04/20"),
        ("3.0", "Latest Author", "Current revision", "2026/03/19"),
    )):
        for cell, value in zip(row.cells, values):
            cell.text = value
    approval = doc.add_table(rows=2, cols=4)
    for row, values in zip(approval.rows, (
        ("Authorized by", "Full Names", "Signature", "Date"),
        ("COO", "External Approver", "", "19-03-2026"),
    )):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.add_paragraph("Data Owner: Determining retention requirements")
    if explicit_owner:
        doc.add_paragraph(f"Document Owner: {explicit_owner}")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_table_policy_suggests_anchored_metadata_without_confusing_responsibility():
    from modules.aria.policy_parser import extract

    parsed = extract(_table_policy_docx(), ".docx", "ISMS-OMNI-P015_v3.0.docx")
    metadata = parsed["metadata"]
    assert parsed["parser"] == "local-docx-2"
    assert {key: metadata[key]["value"] for key in (
        "title", "version", "effective_date", "approval_date", "approver", "owner"
    )} == {
        "title": "Data Retention Policy", "version": "3.0", "effective_date": "2026-03-19",
        "approval_date": "2026-03-19", "approver": "External Approver", "owner": "Latest Author",
    }
    assert metadata["owner"]["kind"] == "inferred"
    assert "confirm" in metadata["owner"]["note"].lower()
    assert "review_date" not in metadata
    explicit = extract(_table_policy_docx("Policy Team"), ".docx", "ISMS-OMNI-P015_v3.0.docx")["metadata"]
    assert explicit["owner"]["value"] == "Policy Team"
    assert explicit["owner"]["kind"] == "stated"


def test_refresh_staged_import_updates_suggestions_without_changing_source_or_scope(test_db, monkeypatch, tmp_path):
    _setup(test_db, monkeypatch, tmp_path)
    actor = _actor(test_db, 101)
    content = _table_policy_docx()
    staged = intake.stage_import(test_db, actor, content, "ISMS-OMNI-P015_v3.0.docx")
    old = {"parser": "local-docx-1", "metadata": {"owner": {"value": "Determining retention requirements",
                                                  "anchor": "block 42", "kind": "stated"}},
           "spans": [], "text_status": "extracted"}
    test_db.execute("UPDATE aria_policy_imports SET extracted_json=%s WHERE id=%s",
                    (json.dumps(old), staged["id"]))
    test_db.commit()
    for uid in (103, 104):
        with pytest.raises(intake.IntakeError) as hidden:
            intake.refresh_suggestions(test_db, _actor(test_db, uid), staged["id"])
        assert hidden.value.status == 404
    refreshed = intake.refresh_suggestions(test_db, actor, staged["id"])
    assert refreshed["extracted"]["metadata"]["owner"]["value"] == "Latest Author"
    assert refreshed["extracted"]["metadata"]["version"]["value"] == "3.0"
    assert refreshed["file_hash"] == hashlib.sha256(content).hexdigest()
    assert intake.stage_import(test_db, actor, content, "ISMS-OMNI-P015_v3.0.docx")["extracted"]["parser"] == "local-docx-2"
    assert test_db.execute("SELECT COUNT(*) FROM aria_policy_imports").fetchone()[0] == 1
    intake.register_import(test_db, actor, staged["id"], {}, [])
    with pytest.raises(intake.IntakeError) as registered:
        intake.refresh_suggestions(test_db, actor, staged["id"])
    assert registered.value.status == 409


def test_import_requires_review_and_independent_external_attestation(test_db, monkeypatch, tmp_path):
    _setup(test_db, monkeypatch, tmp_path)
    author = _actor(test_db, 101)
    submitted = _docx()
    staged = intake.stage_import(test_db, author, submitted, "policy.docx")
    shown = intake.detail(test_db, author, staged["id"])
    assert shown["extracted"]["metadata"]["title"]["anchor"] == "document properties"
    assert shown["links"][0]["ref"] == "A.5.1"
    assert shown["links"][0]["state"] == "candidate"
    assert test_db.execute("SELECT COUNT(*) FROM aria_documents").fetchone()[0] == 0
    result = intake.register_import(test_db, author, staged["id"], {}, [shown["links"][0]["id"]])
    assert result["created"] is True
    assert intake.register_import(test_db, author, staged["id"], {}, [])["created"] is False
    doc = dict(test_db.execute("SELECT * FROM aria_documents WHERE id=%s", (result["document_id"],)).fetchone())
    assert doc["status"] == "Draft"
    assert doc["framework"] == "ISO 27001"
    assert doc["current_policy_version_id"]
    assert test_db.execute("SELECT state,origin FROM aria_policy_versions WHERE document_id=%s", (doc["id"],)).fetchone()["origin"] == "uploaded"
    evidence = test_db.execute("SELECT file_hash,aria_policy_version_id FROM evidence_items WHERE id=%s", (result["evidence_id"],)).fetchone()
    import hashlib
    assert evidence["file_hash"] == hashlib.sha256(submitted).hexdigest()
    assert evidence["aria_policy_version_id"] == doc["current_policy_version_id"]
    with pytest.raises(intake.IntakeError) as same_author:
        intake.verify_external(test_db, author, staged["id"], approver="Board", approved_at="2026-10-01", proof="Signed page 1")
    assert same_author.value.status == 403
    reviewer = _actor(test_db, 102)
    assert intake.verify_external(test_db, reviewer, staged["id"], approver="Board", approved_at="2026-10-01", proof="Signed page 1")["external_status"] == "verified"
    assert test_db.execute("SELECT status FROM aria_documents WHERE id=%s", (doc["id"],)).fetchone()[0] == "Approved"
    assert test_db.execute("SELECT COUNT(*) FROM aria_document_approvals").fetchone()[0] == 0


def test_external_attestation_rejects_stale_current_version(test_db, monkeypatch, tmp_path):
    _setup(test_db, monkeypatch, tmp_path)
    author = _actor(test_db, 101)
    staged = intake.stage_import(test_db, author, _docx(), "policy.docx")
    registered = intake.register_import(test_db, author, staged["id"], {}, [])
    test_db.execute("UPDATE aria_documents SET current_policy_version_id=NULL WHERE id=%s",
                    (registered["document_id"],))
    test_db.commit()

    with pytest.raises(intake.IntakeError, match="current policy changed") as stale:
        intake.verify_external(test_db, _actor(test_db, 102), staged["id"],
                               approver="Board", approved_at="2026-10-01", proof="Signed page 1")
    assert stale.value.status == 409
    assert test_db.execute("SELECT external_status FROM aria_policy_imports WHERE id=%s",
                           (staged["id"],)).fetchone()[0] == "unverified"
    assert test_db.execute("SELECT status FROM aria_documents WHERE id=%s",
                           (registered["document_id"],)).fetchone()[0] == "Draft"


def test_import_hides_foreign_org_and_unit_and_rejects_scope_widening(test_db, monkeypatch, tmp_path):
    _setup(test_db, monkeypatch, tmp_path)
    author = _actor(test_db, 101)
    staged = intake.stage_import(test_db, author, _docx(), "policy.docx")
    for uid in (103, 104):
        with pytest.raises(intake.IntakeError) as hidden:
            intake.detail(test_db, _actor(test_db, uid), staged["id"])
        assert hidden.value.status == 404
    with pytest.raises(intake.IntakeError) as widened:
        intake.stage_import(test_db, author, _docx(), "policy.docx", business_unit_id=120)
    assert widened.value.status == 400


def test_pdf_without_text_is_honestly_unrated(test_db, monkeypatch, tmp_path):
    from pypdf import PdfWriter
    _setup(test_db, monkeypatch, tmp_path)
    writer = PdfWriter(); writer.add_blank_page(width=612, height=792)
    out = io.BytesIO(); writer.write(out)
    result = intake.stage_import(test_db, _actor(test_db, 101), out.getvalue(), "scan.pdf")
    assert result["extracted"]["text_status"] == "no_extractable_text"
    assert intake.detail(test_db, _actor(test_db, 101), result["id"])["links"] == []
    with pytest.raises(intake.IntakeError, match="Confirm the source version"):
        intake.register_import(test_db, _actor(test_db, 101), result["id"], {"title": "Scanned policy", "doc_type": "Policy"}, [])
    with pytest.raises(intake.IntakeError, match="Choose the document type"):
        intake.register_import(test_db, _actor(test_db, 101), result["id"],
                               {"title": "Scanned policy", "version": "1.0"}, [])
    assert intake.register_import(test_db, _actor(test_db, 101), result["id"],
                                  {"title": "Scanned policy", "version": "1.0", "doc_type": "Policy"}, [])["created"]


def test_extractor_timeout_fails_closed_without_forwarding_service_secrets(monkeypatch, tmp_path):
    import subprocess
    monkeypatch.setenv("SECRET_KEY", "never-pass-this-to-a-document-parser")
    def timeout(args, **kwargs):
        assert "SECRET_KEY" not in kwargs["env"]
        assert kwargs["timeout"] == intake.EXTRACT_TIMEOUT
        raise subprocess.TimeoutExpired(args, kwargs["timeout"])
    monkeypatch.setattr(intake.subprocess, "run", timeout)
    with pytest.raises(intake.IntakeError, match="timed out"):
        intake._extract_isolated(tmp_path / "source.pdf", ".pdf", "source.pdf")


def test_legacy_metadata_action_uses_isolated_parser(monkeypatch, tmp_path):
    from modules.aria import routes as aria_routes
    monkeypatch.setattr(aria_routes, "ARIA_UPLOAD_DIR", tmp_path / "aria")
    result = aria_routes._extract_legacy_upload_sync(_docx(), ".docx", "policy.docx")
    assert result["title"] == "Information Security Policy"
    assert result["provenance"]["title"] == "document properties"
    assert not list((tmp_path / "aria").glob("policy-parse-*"))


def test_composed_candidate_preserves_original_and_uses_existing_approval_flow(test_db, monkeypatch, tmp_path):
    import base64
    from PIL import Image
    from docx.oxml import OxmlElement
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.enum.section import WD_ORIENT
    from modules.aria import policy_storage as storage, policy_preview as preview
    _setup(test_db, monkeypatch, tmp_path)
    monkeypatch.setattr(storage, "ARIA_UPLOAD_DIR", tmp_path / "aria")
    monkeypatch.setattr(storage, "WORKFLOW_ROOT", tmp_path / "aria" / "policy_workflow")
    monkeypatch.setattr(composition, "policy_authoring_enabled_for", lambda org: org == 10)
    source_doc = Document()
    source_doc.add_paragraph("Old subsidiary cover")
    source_doc.add_paragraph("Contents")
    source_doc.add_paragraph("Purpose: A.5.1 protects company information.")
    picture = tmp_path / "diagram.png"
    Image.new("RGB", (40, 40), "#47627b").save(picture)
    source_doc.add_picture(str(picture))
    source_doc.add_table(rows=1, cols=2).cell(0, 0).text = "Control owner"
    linked = source_doc.add_paragraph()
    relation = linked.part.relate_to("https://example.invalid/approved-record", RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", relation)
    linked._p.append(hyperlink)
    source_doc.add_paragraph("End of policy")
    source_doc.sections[0].orientation = WD_ORIENT.LANDSCAPE
    original = io.BytesIO(); source_doc.save(original)
    author = _actor(test_db, 101)
    staged = intake.stage_import(test_db, author, original.getvalue(), "subsidiary.docx")
    intake.register_import(test_db, author, staged["id"], {"title": "Security Policy", "version": "1.0", "doc_type": "Policy"}, [])
    blocks = composition.outline(test_db, author, staged["id"])
    assert blocks[2]["preview"].startswith("Purpose")
    candidate = composition.create_candidate(test_db, author, staged["id"],
        {"mode": "replace_cover", "boundary": 2, "exclude_blocks": [], "clear_headers": True,
         "cover_text": "For the group compliance team",
         "cover_logo_base64": base64.b64encode(picture.read_bytes()).decode("ascii")})
    draft = test_db.execute("SELECT * FROM aria_policy_drafts WHERE id=%s", (candidate["draft_id"],)).fetchone()
    output = storage.resolve_stored_path(draft["uploaded_source_path"])
    result = Document(output)
    text = "\n".join(p.text for p in result.paragraphs)
    assert "Old subsidiary cover" not in text
    assert "For the group compliance team" in text
    assert "Purpose: A.5.1" in text
    assert len(result.inline_shapes) == 2
    assert result.tables[0].cell(0, 0).text == "Control owner"
    assert result.sections[0].orientation == WD_ORIENT.LANDSCAPE
    assert any(rel.target_ref == "https://example.invalid/approved-record" for rel in result.part.rels.values())
    assert candidate["manifest"]["cover_logo_sha256"]
    assert intake._source_path(intake.get_import(test_db, author, staged["id"])).read_bytes() == original.getvalue()
    assert draft["content_kind"] == "uploaded_docx"
    monkeypatch.setattr(preview, "submit_conversion_job", lambda data: "job")
    monkeypatch.setattr(preview, "poll_conversion_result", lambda job: b"%PDF-1.4\npreview")
    monkeypatch.setattr(preview, "cleanup_job", lambda job: None)
    from modules.aria import policy_workflow_service as workflow
    built = workflow.build_draft(test_db, author, draft["id"], None, draft["lock_version"])
    assert built["state"] == "ready"
    assert storage.resolve_stored_path(built["source_path"]).read_bytes() == output.read_bytes()
    confirmed = workflow.confirm_draft(test_db, author, draft["id"], built["build_id"], built["lock_version"])
    new_version = test_db.execute("SELECT state,origin,renderer_manifest_json FROM aria_policy_versions WHERE id=%s", (confirmed["version_id"],)).fetchone()
    assert new_version["state"] == "draft" and new_version["origin"] == "authored"
    assert '"replace_cover"' in new_version["renderer_manifest_json"]
    assert test_db.execute("SELECT COUNT(*) FROM aria_policy_versions WHERE document_id=%s", (confirmed["document_id"],)).fetchone()[0] == 2


def test_replacement_body_embeds_reviewed_figure_with_caption_and_alt_text(tmp_path):
    import base64
    from PIL import Image

    source = tmp_path / "source.docx"
    candidate = tmp_path / "candidate.docx"
    document = Document()
    document.add_paragraph("Approved cover")
    document.add_paragraph("Old body")
    document.save(source)
    picture = io.BytesIO()
    Image.new("RGB", (24, 24), "#426781").save(picture, format="PNG")
    figure = {"image_base64": base64.b64encode(picture.getvalue()).decode("ascii"),
              "caption": "Approved system boundary", "alt_text": "Three verified systems"}
    selection = {"mode": "replace_body", "boundary": 1, "replacement_body": "# New reviewed body",
                 "figure": figure}
    manifest = composition._render(source, candidate, selection,
                                   {"title": "Security Policy", "version": "1.0", "effective_date": ""})
    output = Document(candidate)
    assert len(output.inline_shapes) == 1
    assert "Approved system boundary" in [paragraph.text for paragraph in output.paragraphs]
    assert output.inline_shapes[0]._inline.docPr.get("descr") == "Three verified systems"
    assert manifest["figure"]["image_sha256"]
    with pytest.raises(intake.IntakeError, match="caption and an accessible description"):
        composition._render(source, tmp_path / "invalid.docx",
                            {**selection, "figure": {**figure, "alt_text": ""}},
                            {"title": "Security Policy", "version": "1.0", "effective_date": ""})


def test_pdf_cover_rendition_requires_independent_review_and_retains_pages(test_db, monkeypatch, tmp_path):
    from pypdf import PdfWriter, PdfReader
    from modules.aria import policy_storage as storage
    from modules.evidence import routes as evidence_routes
    _setup(test_db, monkeypatch, tmp_path)
    monkeypatch.setattr(storage, "ARIA_UPLOAD_DIR", tmp_path / "aria")
    monkeypatch.setattr(storage, "WORKFLOW_ROOT", tmp_path / "aria" / "policy_workflow")
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_blank_page(width=612, height=792)
    writer.add_blank_page(width=612, height=792)
    out = io.BytesIO(); writer.write(out)
    author = _actor(test_db, 101)
    staged = intake.stage_import(test_db, author, out.getvalue(), "approved.pdf")
    intake.register_import(test_db, author, staged["id"],
                           {"title": "Security Policy", "version": "1.0", "doc_type": "Policy"}, [])
    assert pdf_rendition.page_count(test_db, author, staged["id"])["pages"] == 3
    rendition = pdf_rendition.create(test_db, author, staged["id"], 1)
    with pytest.raises(intake.IntakeError) as same_author:
        pdf_rendition.decide(test_db, author, staged["id"], rendition["rendition_id"], True, "Compared pages")
    assert same_author.value.status == 403
    reviewer = _actor(test_db, 102)
    reviewed = pdf_rendition.decide(test_db, reviewer, staged["id"], rendition["rendition_id"], True,
                                    "Retained pages compared with the original")
    assert reviewed["state"] == "approved" and reviewed["evidence_id"]
    row, file, _ = pdf_rendition.get_rendition(test_db, reviewer, staged["id"], rendition["rendition_id"])
    assert len(PdfReader(file).pages) == 3
    assert row["output_sha256"] == storage.sha256_file(evidence_routes.EVIDENCE_DIR /
        test_db.execute("SELECT file_path FROM evidence_items WHERE id=%s", (reviewed["evidence_id"],)).fetchone()[0])


def test_pdf_page_fingerprint_includes_embedded_diagram_bytes():
    from PIL import Image
    from pypdf import PdfReader
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    def pdf_with_image(color):
        image = Image.new("RGB", (30, 30), color)
        stream = io.BytesIO()
        pdf = canvas.Canvas(stream)
        pdf.drawImage(ImageReader(image), 50, 50, 100, 100)
        pdf.showPage(); pdf.save()
        return PdfReader(io.BytesIO(stream.getvalue()))

    assert pdf_rendition._pages_hash(pdf_with_image("red"), 0) != pdf_rendition._pages_hash(
        pdf_with_image("blue"), 0)


def test_legacy_revision_keeps_previous_vault_bytes_and_scope(test_db, monkeypatch, tmp_path):
    from modules.aria import routes as aria_routes
    _setup(test_db, monkeypatch, tmp_path)
    monkeypatch.setenv("EVIDENCE_DIR", str(tmp_path / "evidence"))
    root = aria_routes.ARIA_UPLOAD_DIR
    root.mkdir(parents=True, exist_ok=True)
    first = _docx()
    (root / "first.docx").write_bytes(first)
    (root / "second.docx").write_bytes(first + b"changed")
    test_db.execute("INSERT INTO aria_documents (id,doc_id,framework,title,org_id,business_unit_id,status) "
                    "VALUES (901,'DOC-0901','Unmapped','Security Policy',10,110,'Draft')")
    test_db.commit()
    initial = aria_routes._sync_aria_to_vault(test_db, 901, "Security Policy", "Unmapped", "", "Policy", "1.0",
                                               "first.docx", "first.docx", len(first), 101)
    assert aria_routes._sync_aria_to_vault(test_db, 901, "Security Policy", "Unmapped", "", "Policy", "1.0",
                                            "first.docx", "first.docx", len(first), 101) == initial
    newer = aria_routes._sync_aria_to_vault(test_db, 901, "Security Policy", "Unmapped", "", "Policy", "1.1",
                                             "second.docx", "second.docx", len(first) + 7, 101)
    test_db.commit()
    old = test_db.execute("SELECT status,file_path,version,org_id,business_unit_id FROM evidence_items WHERE id=%s", (initial,)).fetchone()
    new = test_db.execute("SELECT status,parent_id,version,org_id,business_unit_id FROM evidence_items WHERE id=%s", (newer,)).fetchone()
    assert old["status"] == "superseded" and old["version"] == 1
    assert (tmp_path / "evidence" / old["file_path"]).read_bytes() == first
    assert new["status"] == "current" and new["parent_id"] == initial and new["version"] == 2
    assert (new["org_id"], new["business_unit_id"]) == (10, 110)


def test_exact_duplicate_reuses_visible_import_and_manual_match_needs_review(test_db, monkeypatch, tmp_path):
    _setup(test_db, monkeypatch, tmp_path)
    author = _actor(test_db, 101)
    source = _docx()
    first = intake.stage_import(test_db, author, source, "policy.docx")
    again = intake.stage_import(test_db, author, source, "copy.docx")
    assert again["id"] == first["id"] and again["duplicate_existing"] is True
    assert intake.search_controls(test_db, author, "A.5.1")[0]["id"] == 801
    anchor = first["extracted"]["spans"][0]["anchor"]
    added = intake.add_manual_link(test_db, author, first["id"], 801, anchor, "addresses")
    manual = next(l for l in added["links"] if l["anchor"] == anchor)
    assert manual["match_method"] == "manual" and manual["state"] == "candidate"
    result = intake.register_import(test_db, author, first["id"], {}, [manual["id"]])
    assert intake.stage_import(test_db, author, source, "third.docx")["id"] == first["id"]
    assert test_db.execute("SELECT COUNT(*) FROM evidence_items WHERE id=%s", (result["evidence_id"],)).fetchone()[0] == 1
    other_unit = _actor(test_db, 103)
    staged_other = intake.stage_import(test_db, other_unit, source, "copy.docx")
    assert staged_other["id"] != first["id"] and not staged_other.get("duplicate_existing")


def test_annual_upload_updates_one_policy_only_after_independent_verification(test_db, monkeypatch, tmp_path):
    _setup(test_db, monkeypatch, tmp_path)
    author, reviewer = _actor(test_db, 101), _actor(test_db, 102)
    first = intake.stage_import(test_db, author, _docx(), "security-policy.docx")
    registered = intake.register_import(test_db, author, first["id"], {}, [])
    intake.verify_external(test_db, reviewer, first["id"],
                           approver="Policy board", approved_at="2026-10-01", proof="Signed original")
    current_before = test_db.execute("SELECT current_policy_version_id FROM aria_documents WHERE id=%s",
                                     (registered["document_id"],)).fetchone()[0]
    update = Document()
    update.core_properties.title = "Information Security Policy"
    update.add_paragraph("Version: 3.0")
    update.add_paragraph("Annual update with a changed data-handling clause")
    data = io.BytesIO(); update.save(data)
    staged = intake.stage_import(test_db, author, data.getvalue(), "security-policy-annual.docx")
    candidates = intake.version_candidates(test_db, author, staged["id"], "Information Security")
    assert [item["id"] for item in candidates] == [registered["document_id"]]
    with pytest.raises(intake.IntakeError) as hidden:
        intake.register_import(test_db, _actor(test_db, 103), staged["id"],
                               {"version": "3.0"}, [], registered["document_id"])
    assert hidden.value.status == 404
    second = intake.register_import(test_db, author, staged["id"], {}, [], registered["document_id"])
    assert second["pending_replacement"] and second["document_id"] == registered["document_id"]
    assert test_db.execute("SELECT current_policy_version_id FROM aria_documents WHERE id=%s",
                           (registered["document_id"],)).fetchone()[0] == current_before
    assert test_db.execute("SELECT status FROM evidence_items WHERE id=%s",
                           (second["evidence_id"],)).fetchone()[0] == "draft"
    intake.verify_external(test_db, reviewer, staged["id"],
                           approver="Policy board", approved_at="2026-10-15", proof="Signed annual review")
    current = test_db.execute("SELECT current_policy_version_id,version FROM aria_documents WHERE id=%s",
                              (registered["document_id"],)).fetchone()
    assert current["current_policy_version_id"] != current_before and current["version"] == "3.0"
    assert test_db.execute("SELECT status FROM evidence_items WHERE id=%s",
                           (registered["evidence_id"],)).fetchone()[0] == "superseded"
    assert test_db.execute("SELECT status FROM evidence_items WHERE id=%s",
                           (second["evidence_id"],)).fetchone()[0] == "current"
    assert test_db.execute("SELECT COUNT(*) FROM aria_documents WHERE id=%s",
                           (registered["document_id"],)).fetchone()[0] == 1


def test_policy_reference_hides_unverified_and_stale_clause_links(test_db, monkeypatch, tmp_path):
    _setup(test_db, monkeypatch, tmp_path)
    author = _actor(test_db, 101)
    staged = intake.stage_import(test_db, author, _docx(), "policy.docx")
    link_id = intake.detail(test_db, author, staged["id"])["links"][0]["id"]
    registered = intake.register_import(test_db, author, staged["id"], {}, [link_id])
    assert policy_references.resolve_policy_references(test_db, author, ("aria", "control", 801)) == []
    intake.verify_external(test_db, _actor(test_db, 102), staged["id"],
                           approver="Board", approved_at="2026-10-01", proof="Signed page")
    refs = policy_references.resolve_policy_references(test_db, author, ("aria", "control", 801))
    assert len(refs) == 1 and refs[0]["approval"] == "external_verified"
    assert refs[0]["evidence_id"] == registered["evidence_id"]
    from modules.governance.connected_records import section
    blocks = section(test_db, author, ("aria", "control"), 801)
    policy_table = next((block for block in blocks if getattr(block, "title", "") == "Current policy references"), None)
    assert policy_table is not None and policy_table.rows[0]["version_id"] == refs[0]["version_id"]
    assert policy_references.resolve_policy_references(test_db, _actor(test_db, 103), ("aria", "control", 801)) == []
    test_db.execute("UPDATE aria_documents SET current_policy_version_id=NULL WHERE id=%s", (registered["document_id"],))
    test_db.commit()
    assert policy_references.resolve_policy_references(test_db, author, ("aria", "control", 801)) == []


def test_future_external_approval_is_not_accepted_and_invalid_logo_fails(test_db, monkeypatch, tmp_path):
    _setup(test_db, monkeypatch, tmp_path)
    author = _actor(test_db, 101)
    staged = intake.stage_import(test_db, author, _docx(), "policy.docx")
    intake.register_import(test_db, author, staged["id"], {}, [])
    with pytest.raises(intake.IntakeError, match="future"):
        intake.verify_external(test_db, _actor(test_db, 102), staged["id"],
                               approver="Board", approved_at="2999-01-01", proof="Signed page")
    from modules.aria import policy_storage as storage
    monkeypatch.setattr(storage, "ARIA_UPLOAD_DIR", tmp_path / "aria")
    monkeypatch.setattr(storage, "WORKFLOW_ROOT", tmp_path / "aria" / "policy_workflow")
    monkeypatch.setattr(composition, "policy_authoring_enabled_for", lambda org: True)
    with pytest.raises(intake.IntakeError, match="Cover logo"):
        composition.create_candidate(test_db, author, staged["id"],
            {"mode": "replace_cover", "boundary": 1, "cover_logo_base64": "bm90LWFuLWltYWdl"})
