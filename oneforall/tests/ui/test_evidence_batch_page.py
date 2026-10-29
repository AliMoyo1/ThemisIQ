"""The staged browser queue uses an explicit filing scope and reports each result."""

from playwright.sync_api import expect
from docx import Document
import io


def test_batch_upload_requires_review_and_files_one_document(login_as, live_app):
    page = login_as("compliance_manager")
    page.goto(f"{live_app}/evidence/batch/")
    expect(page.get_by_role("heading", name="Batch upload")).to_be_visible()
    expect(page.locator("#ebSummary")).to_contain_text("Choose files")
    page.locator("#ebFiles").set_input_files({
        "name": "Synthetic_Continuity_Test.txt", "mimeType": "text/plain",
        "buffer": b"Synthetic evidence for a disposable browser test."
    })
    expect(page.locator("#ebQueue .eb-row")).to_have_count(1)
    expect(page.locator("#ebUpload")).to_be_disabled()
    page.locator("#ebQueue select").first.select_option(index=1)
    page.get_by_label("Checked for filing").check()
    expect(page.locator("#ebUpload")).to_be_enabled()
    page.locator("#ebUpload").click()
    expect(page.locator("#ebQueue .eb-state")).to_contain_text("Filed in the Vault")


def test_duplicate_in_batch_does_not_stop_later_files(login_as, live_app):
    page = login_as("compliance_manager")
    page.goto(f"{live_app}/evidence/batch/")
    expect(page.locator("#ebSummary")).to_contain_text("Choose files")
    page.locator("#ebFiles").set_input_files([
        {"name": "First.txt", "mimeType": "text/plain", "buffer": b"same synthetic batch bytes"},
        {"name": "Duplicate.txt", "mimeType": "text/plain", "buffer": b"same synthetic batch bytes"},
        {"name": "Third.txt", "mimeType": "text/plain", "buffer": b"third synthetic batch bytes"},
    ])
    expect(page.locator("#ebQueue .eb-row")).to_have_count(3)
    for row in page.locator("#ebQueue .eb-row").all():
        row.locator("select").first.select_option(index=1)
    for row in page.locator("#ebQueue .eb-row").all():
        row.get_by_label("Checked for filing").check()
    page.locator("#ebUpload").click()
    expect(page.locator("#ebQueue .eb-state[data-state=uploaded]")).to_have_count(2)
    expect(page.locator("#ebQueue .eb-state[data-state=duplicate]")).to_have_count(1)


def test_batch_extracts_reviewable_source_anchored_docx_suggestions(login_as, live_app):
    document = Document()
    document.core_properties.title = "Data Retention Policy"
    document.add_paragraph("Version: 3.0")
    document.add_paragraph("Review date: 2027-06-01")
    content = io.BytesIO(); document.save(content)
    page = login_as("compliance_manager")
    page.goto(f"{live_app}/evidence/batch/")
    expect(page.locator("#ebSummary")).to_contain_text("Choose files")
    page.locator("#ebFiles").set_input_files({
        "name": "old_template.docx", "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "buffer": content.getvalue()
    })
    page.get_by_role("button", name="Extract suggestions").click()
    expect(page.locator("#ebQueue .eb-row input[type=text]").first).to_have_value("Data Retention Policy")
    expect(page.locator("#ebQueue .eb-row input[type=date]")).to_have_value("2027-06-01")
    expect(page.locator("#ebQueue .eb-note")).to_contain_text("version: 3.0")
    expect(page.locator("#ebQueue .eb-note")).to_contain_text("block 1")
    expect(page.locator("#ebUpload")).to_be_disabled()


def test_batch_rejects_invalid_pdf_without_claiming_extraction(login_as, live_app):
    page = login_as("compliance_manager")
    page.goto(f"{live_app}/evidence/batch/")
    expect(page.locator("#ebSummary")).to_contain_text("Choose files")
    page.locator("#ebFiles").set_input_files({
        "name": "fake.pdf", "mimeType": "application/pdf", "buffer": b"not a PDF"
    })
    page.get_by_role("button", name="Extract suggestions").click()
    expect(page.locator("#ebQueue .eb-state")).to_contain_text("PDF file signature is invalid")
    expect(page.locator("#ebQueue input[type=text]").first).to_have_value("fake")
    expect(page.locator("#ebUpload")).to_be_disabled()


def test_batch_new_version_requires_current_record_check_and_keeps_scope(login_as, live_app):
    page = login_as("compliance_manager")
    page.goto(f"{live_app}/evidence/batch/")
    expect(page.locator("#ebSummary")).to_contain_text("Choose files")
    page.locator("#ebFiles").set_input_files({
        "name": "Original.txt", "mimeType": "text/plain", "buffer": b"Original synthetic policy version"
    })
    first = page.locator("#ebQueue .eb-row").first
    first.get_by_label("Filing scope").select_option(index=1)
    first.get_by_label("Checked for filing").check()
    with page.expect_response(lambda response: response.url.endswith("/evidence/api/items") and
                              response.request.method == "POST") as uploaded:
        page.locator("#ebUpload").click()
    original_id = uploaded.value.json()["id"]
    expect(first.locator(".eb-state")).to_contain_text("Filed in the Vault")

    page.locator("#ebFiles").set_input_files({
        "name": "Updated.txt", "mimeType": "text/plain", "buffer": b"Updated synthetic policy version"
    })
    second = page.locator("#ebQueue .eb-row").last
    second.get_by_label("Filing scope").select_option(index=1)
    second.get_by_label("File as").select_option("replace")
    second.get_by_label("Current Vault document ID").fill(str(original_id))
    second.get_by_label("Checked for filing").check()
    expect(page.locator("#ebUpload")).to_be_disabled()
    second.get_by_role("button", name="Check current version").click()
    expect(second.locator(".eb-state")).to_contain_text("Checked:")
    expect(second.locator(".eb-compare")).to_contain_text("Original.txt → Updated.txt")
    expect(second.locator(".eb-compare")).to_contain_text("compares metadata only")
    expect(page.locator("#ebUpload")).to_be_disabled()
    second.get_by_label("Checked for filing").check()
    with page.expect_response(lambda response: response.url.endswith("/evidence/api/items") and
                              response.request.method == "POST") as revised:
        page.locator("#ebUpload").click()
    assert revised.value.json()["version"] == 2
    expect(second.locator(".eb-state")).to_contain_text("New version filed")
